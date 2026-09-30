"""Word (.docx) styling and formatting helpers."""

from typing import Optional, Sequence
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

# Toptal Professional Palette
COLOR_NAVY_HEX = "0F2137"
COLOR_PRIMARY_BLUE_HEX = "1B365D"
COLOR_ACCENT_BLUE_HEX = "2563EB"
COLOR_LIGHT_BG_HEX = "F1F5F9"
COLOR_WARNING_BG_HEX = "FEF3C7"
COLOR_BORDER_HEX = "CBD5E1"
COLOR_TEXT_MUTED_HEX = "64748B"

ACTION_TAG_REGEX = re.compile(
    r'(\[(?:ACT|ACT-REQ)(?:-[A-Za-z0-9_]+)?(?::[^\n\]]*)?\])',
    re.IGNORECASE
)


def set_cell_background(cell, hex_color: str):
    """Set the background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    for shd in tcPr.findall(qn('w:shd')):
        tcPr.remove(shd)
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>'
    tcPr.append(parse_xml(shading_xml))


def set_cell_margins(cell, top: int = 120, bottom: int = 120, left: int = 160, right: int = 160):
    """Set inner padding / margins for a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m_name, m_val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m_name}')
        node.set(qn('w:w'), str(m_val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def add_section_heading(doc: docx.Document, text: str, level: int = 1):
    """Add a styled section heading."""
    heading = doc.add_heading(level=level)
    run = heading.add_run(text)
    run.font.name = "Arial"
    run.bold = True
    if level == 1:
        run.font.size = Pt(16)
        run.font.color.rgb = RGBColor(15, 33, 55)  # Navy
        heading.paragraph_format.space_before = Pt(16)
        heading.paragraph_format.space_after = Pt(6)
    elif level == 2:
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(27, 54, 93)  # Slate Navy
        heading.paragraph_format.space_before = Pt(12)
        heading.paragraph_format.space_after = Pt(4)
    else:
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(37, 99, 235)  # Accent Blue
        heading.paragraph_format.space_before = Pt(8)
        heading.paragraph_format.space_after = Pt(2)
    return heading


def add_callout_box(
    doc: docx.Document,
    text: str,
    title: Optional[str] = None,
    bg_color: str = COLOR_LIGHT_BG_HEX,
    border_color: str = COLOR_ACCENT_BLUE_HEX
):
    """Add a bordered, styled callout container to the document."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)

    # Set left thick border
    tcPr = cell._tc.get_or_add_tcPr()
    borders_xml = f"""
    <w:tcBorders {nsdecls("w")}>
        <w:top w:val="none"/>
        <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>
        <w:bottom w:val="none"/>
        <w:right w:val="none"/>
    </w:tcBorders>
    """
    tcPr.append(parse_xml(borders_xml))

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)

    if title:
        run_title = p.add_run(f"{title}\n")
        run_title.font.name = "Arial"
        run_title.font.size = Pt(10.5)
        run_title.bold = True
        run_title.font.color.rgb = RGBColor(15, 33, 55)

    run_text = p.add_run(text)
    run_text.font.name = "Arial"
    run_text.font.size = Pt(9.5)
    run_text.font.color.rgb = RGBColor(30, 41, 59)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def format_cell_text_and_highlight(cell, text: str, is_warning: bool = False):
    """Format cell text and apply WD_COLOR_INDEX.YELLOW highlight to every [ACT-...] tag run."""
    cell.text = ""
    p = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)

    parts = ACTION_TAG_REGEX.split(text or "")
    for part in parts:
        if not part:
            continue
        run = p.add_run(part)
        run.font.name = "Arial"
        run.font.size = Pt(9)
        if ACTION_TAG_REGEX.match(part):
            run.bold = True
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
            run.font.color.rgb = RGBColor(15, 23, 42)
        else:
            run.font.color.rgb = RGBColor(15, 23, 42)


def sanitize_filename(name: str) -> str:
    """Sanitize project name for safe filename creation."""
    s = re.sub(r'[^a-zA-Z0-9_\- ]+', '', name or "").strip()
    return re.sub(r'\s+', '_', s)


def style_table(
    table: docx.table.Table,
    col_widths: Optional[Sequence[float]] = None
):

    # Format header row
    header_row = table.rows[0]
    for idx, cell in enumerate(header_row.cells):
        set_cell_background(cell, COLOR_NAVY_HEX)
        set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.name = "Arial"
                run.font.size = Pt(9.5)
                run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)

    # Format body rows with zebra striping
    for row_idx, row in enumerate(table.rows[1:], start=1):
        bg = COLOR_LIGHT_BG_HEX if row_idx % 2 == 1 else "FFFFFF"
        for cell in row.cells:
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(9)
                    if ACTION_TAG_REGEX.search(run.text):
                        run.bold = True
                        run.font.highlight_color = WD_COLOR_INDEX.YELLOW
                    if not run.font.color.rgb:
                        run.font.color.rgb = RGBColor(15, 23, 42)

    # Apply column widths if specified (in inches)
    if col_widths:
        for row in table.rows:
            for idx, width in enumerate(col_widths):
                if idx < len(row.cells):
                    row.cells[idx].width = Inches(width)

"""Presentation writer for Talent Team Onboarding Deck (DECK-12, DECK-13, Appendix K.1, K.2)."""

import os
from pathlib import Path
from typing import Optional, List, Dict, Any
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

from src.config import config, DECK_TEMPLATE_PATH
from src.generators.formatting import sanitize_filename
from src.generators.onboarding_deck.builder import (
    DeckModel,
    TableRowData,
    KeyFactRow,
)
from src.generators.onboarding_deck.styles import (
    SLIDE_WIDTH,
    SLIDE_HEIGHT,
    COLOR_DEEP_NAVY,
    COLOR_SLATE,
    COLOR_ACCENT_BLUE,
    COLOR_MUTED_GRAY,
    COLOR_LIGHT_GRAY,
    COLOR_BORDER_GRAY,
    COLOR_WHITE,
    COLOR_RED_TEXT,
    COLOR_RED_BG,
    COLOR_AMBER_TEXT,
    COLOR_AMBER_BG,
    COLOR_GREEN_TEXT,
    COLOR_GREEN_BG,
    FONT_TITLE,
    FONT_BODY,
    FONT_SIZE_TITLE,
    FONT_SIZE_KICKER,
    FONT_SIZE_CARD_HEADER,
    FONT_SIZE_CARD_SUBHEADER,
    FONT_SIZE_BODY,
    FONT_SIZE_TABLE_HEADER,
    FONT_SIZE_TABLE_BODY,
    FONT_SIZE_SOURCES,
)


def _set_cell_text(
    cell: Any,
    text: str,
    bold: bool = False,
    color: RGBColor = COLOR_SLATE,
    font_size: Pt = FONT_SIZE_TABLE_BODY,
    is_placeholder: bool = False,
    italic: bool = False,
    font_name: str = FONT_BODY,
    align: PP_ALIGN = PP_ALIGN.LEFT,
) -> None:
    """Set formatted text inside a table cell."""
    cell.text = ""
    p = cell.text_frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font_name
    run.font.size = font_size
    run.font.bold = bold
    if is_placeholder:
        run.font.italic = True
        run.font.color.rgb = COLOR_ACCENT_BLUE
    else:
        run.font.italic = italic
        run.font.color.rgb = color


def _set_cell_background(cell: Any, color: RGBColor) -> None:
    """Set fill color for a table cell."""
    cell.fill.solid()
    cell.fill.fore_color.rgb = color


def _add_card_shape(
    slide: Any,
    left: Inches,
    top: Inches,
    width: Inches,
    height: Inches,
) -> Any:
    """Add a styled white rounded rectangle card with border."""
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = COLOR_WHITE
    card.line.color.rgb = COLOR_BORDER_GRAY
    card.line.width = Pt(1)
    return card


def _render_table(
    slide: Any,
    headers: List[str],
    rows: List[TableRowData],
    col_widths: List[Inches],
    left: Inches,
    top: Inches,
    width: Inches,
    height: Inches,
) -> Any:
    """Render a styled table on a slide."""
    num_rows = len(rows) + 1
    num_cols = len(headers)
    table_shape = slide.shapes.add_table(num_rows, num_cols, left, top, width, height)
    table = table_shape.table

    for idx, w in enumerate(col_widths):
        table.columns[idx].width = w

    # Render header row (204ECF background, white bold Proxima Nova 12pt)
    for c_idx, h_text in enumerate(headers):
        cell = table.cell(0, c_idx)
        _set_cell_background(cell, COLOR_ACCENT_BLUE)
        _set_cell_text(
            cell,
            h_text,
            bold=True,
            color=COLOR_WHITE,
            font_name=FONT_TITLE,
            font_size=FONT_SIZE_TABLE_HEADER,
        )

    # Render data rows (alternating F8FAFC and FFFFFF, borders E2E8F0)
    for r_idx, row_data in enumerate(rows, start=1):
        bg_col = COLOR_LIGHT_GRAY if (r_idx % 2 == 1) else COLOR_WHITE

        if row_data.is_overflow:
            cell = table.cell(r_idx, 0)
            _set_cell_background(cell, bg_col)
            _set_cell_text(
                cell,
                row_data.cells[0],
                bold=True,
                italic=True,
                color=COLOR_MUTED_GRAY,
                font_name=FONT_BODY,
                font_size=FONT_SIZE_TABLE_BODY,
            )
            for c_idx in range(1, num_cols):
                c = table.cell(r_idx, c_idx)
                _set_cell_background(c, bg_col)
                _set_cell_text(c, "", font_name=FONT_BODY, font_size=FONT_SIZE_TABLE_BODY)
            continue

        for c_idx, cell_text in enumerate(row_data.cells):
            if c_idx >= num_cols:
                break
            cell = table.cell(r_idx, c_idx)
            is_ph = row_data.is_placeholder[c_idx] if c_idx < len(row_data.is_placeholder) else False
            is_first_col = (c_idx == 0)
            text_col = COLOR_DEEP_NAVY if is_first_col else COLOR_SLATE
            bold_flag = is_first_col

            # RAG fill on Rating column
            if headers[c_idx] == "Rating" and cell_text == "High":
                _set_cell_background(cell, COLOR_RED_BG)
                _set_cell_text(cell, cell_text, bold=True, color=COLOR_RED_TEXT, font_size=FONT_SIZE_TABLE_BODY)
            elif headers[c_idx] == "Rating" and cell_text == "Medium":
                _set_cell_background(cell, COLOR_AMBER_BG)
                _set_cell_text(cell, cell_text, bold=True, color=COLOR_AMBER_TEXT, font_size=FONT_SIZE_TABLE_BODY)
            elif headers[c_idx] == "Rating" and cell_text == "Low":
                _set_cell_background(cell, COLOR_GREEN_BG)
                _set_cell_text(cell, cell_text, bold=True, color=COLOR_GREEN_TEXT, font_size=FONT_SIZE_TABLE_BODY)
            else:
                _set_cell_background(cell, bg_col)
                _set_cell_text(
                    cell,
                    cell_text,
                    bold=bold_flag,
                    is_placeholder=is_ph,
                    color=text_col,
                    font_size=FONT_SIZE_TABLE_BODY,
                )

    return table_shape


def write_onboarding_deck(
    model: DeckModel,
    target_path: Path,
    template_path: Optional[Path] = None,
) -> Path:
    """Write OnboardingDeck presentation and trace manifest from DeckModel (DECK-12, DECK-13)."""
    resolved_template = template_path or config.deck_template_path or DECK_TEMPLATE_PATH
    if not resolved_template.exists():
        raise FileNotFoundError(f"Deck template missing: '{resolved_template}'")

    prs = pptx.Presentation(str(resolved_template))

    # Clear template slides and drop relationships (DECK-12)
    slides = prs.slides
    sldIdLst = slides._sldIdLst
    slide_rIds = [elem.rId for elem in list(sldIdLst)]
    for elem in list(sldIdLst):
        sldIdLst.remove(elem)
    for rId in slide_rIds:
        if rId in prs.part.rels:
            prs.part.drop_rel(rId)

    # Locate CUSTOM_1 and CUSTOM_16 layouts
    slide_master = prs.slide_masters[0]
    layout_c1 = None
    layout_c16 = None
    for l in slide_master.slide_layouts:
        if l.name == "CUSTOM_1":
            layout_c1 = l
        elif l.name == "CUSTOM_16":
            layout_c16 = l

    if layout_c1 is None:
        raise ValueError("Required slide layout 'CUSTOM_1' not found in template.")
    if layout_c16 is None:
        raise ValueError("Required slide layout 'CUSTOM_16' not found in template.")

    # -------------------------------------------------------------
    # SLIDE 1: Cover (CUSTOM_1)
    # -------------------------------------------------------------
    s1 = prs.slides.add_slide(layout_c1)
    if len(s1.placeholders) > 0:
        s1.placeholders[0].text = model.cover_slide.title
    if len(s1.placeholders) > 1:
        s1.placeholders[1].text = model.cover_slide.subtitle

    s1.notes_slide.notes_text_frame.text = model.cover_slide.speaker_notes

    # -------------------------------------------------------------
    # SLIDE 2: Project Charter (CUSTOM_16)
    # -------------------------------------------------------------
    s2 = prs.slides.add_slide(layout_c16)
    if len(s2.placeholders) > 0:
        s2.placeholders[0].text = model.charter_slide.title
    if len(s2.placeholders) > 1:
        s2.placeholders[1].text = model.charter_slide.kicker

    # Card 1: Key facts (x=0.83, y=1.70, w=3.71, h=4.58)
    _add_card_shape(s2, Inches(0.83), Inches(1.70), Inches(3.71), Inches(4.58))
    box1 = s2.shapes.add_textbox(Inches(0.95), Inches(1.80), Inches(3.47), Inches(4.38))
    tf1 = box1.text_frame
    tf1.word_wrap = True
    p1 = tf1.paragraphs[0]
    p1.text = "Key facts"
    p1.font.name = FONT_TITLE
    p1.font.size = FONT_SIZE_CARD_HEADER
    p1.font.bold = True
    p1.font.color.rgb = COLOR_DEEP_NAVY

    # Add 2-column key facts table inside card
    kf_table_shape = s2.shapes.add_table(7, 2, Inches(0.95), Inches(2.25), Inches(3.47), Inches(3.8))
    kf_table = kf_table_shape.table
    kf_table.columns[0].width = Inches(1.4)
    kf_table.columns[1].width = Inches(2.07)
    for r_i, kf in enumerate(model.charter_slide.key_facts):
        c0 = kf_table.cell(r_i, 0)
        _set_cell_text(c0, kf.label, bold=True, color=COLOR_DEEP_NAVY, font_size=Pt(9.5))
        c1 = kf_table.cell(r_i, 1)
        _set_cell_text(c1, kf.value, is_placeholder=kf.is_placeholder, color=COLOR_SLATE, font_size=Pt(9.5))

    # Card 2: Purpose and delivery model (x=4.81, y=1.70, w=3.71, h=4.58)
    _add_card_shape(s2, Inches(4.81), Inches(1.70), Inches(3.71), Inches(4.58))
    box2 = s2.shapes.add_textbox(Inches(4.93), Inches(1.80), Inches(3.47), Inches(4.38))
    tf2 = box2.text_frame
    tf2.word_wrap = True

    p2 = tf2.paragraphs[0]
    p2.text = "Purpose and delivery model"
    p2.font.name = FONT_TITLE
    p2.font.size = FONT_SIZE_CARD_HEADER
    p2.font.bold = True
    p2.font.color.rgb = COLOR_DEEP_NAVY

    # Purpose
    p_p_h = tf2.add_paragraph()
    p_p_h.text = "Project Purpose"
    p_p_h.font.name = FONT_TITLE
    p_p_h.font.size = FONT_SIZE_CARD_SUBHEADER
    p_p_h.font.bold = True
    p_p_h.font.color.rgb = COLOR_ACCENT_BLUE

    p_p_t = tf2.add_paragraph()
    p_p_t.text = model.charter_slide.purpose
    p_p_t.font.name = FONT_BODY
    p_p_t.font.size = FONT_SIZE_BODY
    p_p_t.font.color.rgb = COLOR_SLATE

    # Delivery Model
    p_m_h = tf2.add_paragraph()
    p_m_h.text = "Delivery Model & Governance"
    p_m_h.font.name = FONT_TITLE
    p_m_h.font.size = FONT_SIZE_CARD_SUBHEADER
    p_m_h.font.bold = True
    p_m_h.font.color.rgb = COLOR_ACCENT_BLUE

    p_m_t = tf2.add_paragraph()
    p_m_t.text = model.charter_slide.delivery_model
    p_m_t.font.name = FONT_BODY
    p_m_t.font.size = FONT_SIZE_BODY
    p_m_t.font.color.rgb = COLOR_SLATE

    # Escalation Path
    p_e_h = tf2.add_paragraph()
    p_e_h.text = "Escalation Path"
    p_e_h.font.name = FONT_TITLE
    p_e_h.font.size = FONT_SIZE_CARD_SUBHEADER
    p_e_h.font.bold = True
    p_e_h.font.color.rgb = COLOR_ACCENT_BLUE

    p_e_t = tf2.add_paragraph()
    p_e_t.text = model.charter_slide.escalation_path
    p_e_t.font.name = FONT_BODY
    p_e_t.font.size = FONT_SIZE_BODY
    p_e_t.font.color.rgb = COLOR_SLATE

    # Card 3: Phases and boundaries (x=8.79, y=1.70, w=3.71, h=4.58)
    _add_card_shape(s2, Inches(8.79), Inches(1.70), Inches(3.71), Inches(4.58))
    box3 = s2.shapes.add_textbox(Inches(8.91), Inches(1.80), Inches(3.47), Inches(4.38))
    tf3 = box3.text_frame
    tf3.word_wrap = True

    p3 = tf3.paragraphs[0]
    p3.text = "Phases and boundaries"
    p3.font.name = FONT_TITLE
    p3.font.size = FONT_SIZE_CARD_HEADER
    p3.font.bold = True
    p3.font.color.rgb = COLOR_DEEP_NAVY

    p_ph_h = tf3.add_paragraph()
    p_ph_h.text = "Delivery Phases"
    p_ph_h.font.name = FONT_TITLE
    p_ph_h.font.size = FONT_SIZE_CARD_SUBHEADER
    p_ph_h.font.bold = True
    p_ph_h.font.color.rgb = COLOR_ACCENT_BLUE

    for ph in model.charter_slide.phases:
        p_b = tf3.add_paragraph()
        p_b.text = f"• {ph}"
        p_b.font.name = FONT_BODY
        p_b.font.size = FONT_SIZE_BODY
        p_b.font.color.rgb = COLOR_SLATE

    p_ex_h = tf3.add_paragraph()
    p_ex_h.text = "Out of Scope"
    p_ex_h.font.name = FONT_TITLE
    p_ex_h.font.size = FONT_SIZE_CARD_SUBHEADER
    p_ex_h.font.bold = True
    p_ex_h.font.color.rgb = COLOR_ACCENT_BLUE

    for exc in model.charter_slide.exclusions:
        p_b = tf3.add_paragraph()
        p_b.text = f"• {exc}"
        p_b.font.name = FONT_BODY
        p_b.font.size = FONT_SIZE_BODY
        p_b.font.color.rgb = COLOR_SLATE

    if model.charter_slide.overflow_exclusions > 0:
        p_ov = tf3.add_paragraph()
        p_ov.text = f"• +{model.charter_slide.overflow_exclusions} more (see Startup Kit)"
        p_ov.font.name = FONT_BODY
        p_ov.font.size = Pt(10)
        p_ov.font.italic = True
        p_ov.font.color.rgb = COLOR_MUTED_GRAY

    s2.notes_slide.notes_text_frame.text = model.charter_slide.speaker_notes

    # -------------------------------------------------------------
    # SLIDE 3: Workstreams, Milestones, Deliverables and Dates (CUSTOM_16)
    # -------------------------------------------------------------
    s3 = prs.slides.add_slide(layout_c16)
    if len(s3.placeholders) > 0:
        s3.placeholders[0].text = model.schedule_slide.title
    if len(s3.placeholders) > 1:
        s3.placeholders[1].text = model.schedule_slide.kicker

    sched_widths = [Inches(2.2), Inches(3.2), Inches(2.2), Inches(4.07)]
    _render_table(
        s3,
        model.schedule_slide.headers,
        model.schedule_slide.rows,
        sched_widths,
        Inches(0.83),
        Inches(1.70),
        Inches(11.67),
        Inches(4.58),
    )
    s3.notes_slide.notes_text_frame.text = model.schedule_slide.speaker_notes

    # -------------------------------------------------------------
    # SLIDE 4: Acceptance Criteria (CUSTOM_16)
    # -------------------------------------------------------------
    s4 = prs.slides.add_slide(layout_c16)
    if len(s4.placeholders) > 0:
        s4.placeholders[0].text = model.acceptance_slide.title
    if len(s4.placeholders) > 1:
        s4.placeholders[1].text = model.acceptance_slide.kicker

    # Left Card: How acceptance works (x=0.83, y=1.70, w=3.71, h=4.58)
    _add_card_shape(s4, Inches(0.83), Inches(1.70), Inches(3.71), Inches(4.58))
    box4 = s4.shapes.add_textbox(Inches(0.95), Inches(1.80), Inches(3.47), Inches(4.38))
    tf4 = box4.text_frame
    tf4.word_wrap = True

    p4_h = tf4.paragraphs[0]
    p4_h.text = "How acceptance works"
    p4_h.font.name = FONT_TITLE
    p4_h.font.size = FONT_SIZE_CARD_HEADER
    p4_h.font.bold = True
    p4_h.font.color.rgb = COLOR_DEEP_NAVY

    p4_sub1 = tf4.add_paragraph()
    p4_sub1.text = "Milestone Acceptance Process"
    p4_sub1.font.name = FONT_TITLE
    p4_sub1.font.size = FONT_SIZE_CARD_SUBHEADER
    p4_sub1.font.bold = True
    p4_sub1.font.color.rgb = COLOR_ACCENT_BLUE

    for idx, st in enumerate(model.acceptance_slide.steps, start=1):
        p_st = tf4.add_paragraph()
        p_st.text = f"{idx}. {st}"
        p_st.font.name = FONT_BODY
        p_st.font.size = Pt(9.5)
        p_st.font.color.rgb = COLOR_SLATE

    p4_sub2 = tf4.add_paragraph()
    p4_sub2.text = "Governance Parameters"
    p4_sub2.font.name = FONT_TITLE
    p4_sub2.font.size = FONT_SIZE_CARD_SUBHEADER
    p4_sub2.font.bold = True
    p4_sub2.font.color.rgb = COLOR_ACCENT_BLUE

    p_rw = tf4.add_paragraph()
    r_rw1 = p_rw.add_run()
    r_rw1.text = "Review Window: "
    r_rw1.font.bold = True
    r_rw1.font.name = FONT_BODY
    r_rw1.font.size = FONT_SIZE_BODY
    r_rw1.font.color.rgb = COLOR_DEEP_NAVY
    r_rw2 = p_rw.add_run()
    r_rw2.text = model.acceptance_slide.review_window
    r_rw2.font.name = FONT_BODY
    r_rw2.font.size = FONT_SIZE_BODY
    r_rw2.font.color.rgb = COLOR_SLATE

    p_ap = tf4.add_paragraph()
    r_ap1 = p_ap.add_run()
    r_ap1.text = "Client Approver: "
    r_ap1.font.bold = True
    r_ap1.font.name = FONT_BODY
    r_ap1.font.size = FONT_SIZE_BODY
    r_ap1.font.color.rgb = COLOR_DEEP_NAVY
    r_ap2 = p_ap.add_run()
    r_ap2.text = model.acceptance_slide.client_approver
    r_ap2.font.name = FONT_BODY
    r_ap2.font.size = FONT_SIZE_BODY
    if model.acceptance_slide.is_approver_placeholder:
        r_ap2.font.italic = True
        r_ap2.font.color.rgb = COLOR_ACCENT_BLUE
    else:
        r_ap2.font.color.rgb = COLOR_SLATE

    # Right Table: Deliverables (x=4.81, y=1.70, w=7.69, h=4.58)
    acc_widths = [Inches(1.1), Inches(5.39), Inches(1.2)]
    _render_table(
        s4,
        model.acceptance_slide.headers,
        model.acceptance_slide.rows,
        acc_widths,
        Inches(4.81),
        Inches(1.70),
        Inches(7.69),
        Inches(4.58),
    )
    s4.notes_slide.notes_text_frame.text = model.acceptance_slide.speaker_notes

    # -------------------------------------------------------------
    # SLIDE 5: High-Risk Items (CUSTOM_16)
    # -------------------------------------------------------------
    s5 = prs.slides.add_slide(layout_c16)
    if len(s5.placeholders) > 0:
        s5.placeholders[0].text = model.risks_slide.title
    if len(s5.placeholders) > 1:
        s5.placeholders[1].text = model.risks_slide.kicker

    risk_widths = [Inches(1.5), Inches(4.27), Inches(1.1), Inches(1.6), Inches(2.2), Inches(1.0)]
    _render_table(
        s5,
        model.risks_slide.headers,
        model.risks_slide.rows,
        risk_widths,
        Inches(0.83),
        Inches(1.70),
        Inches(11.67),
        Inches(4.58),
    )
    s5.notes_slide.notes_text_frame.text = model.risks_slide.speaker_notes

    # -------------------------------------------------------------
    # SLIDE 6: Client Collaboration (CUSTOM_16)
    # -------------------------------------------------------------
    s6 = prs.slides.add_slide(layout_c16)
    if len(s6.placeholders) > 0:
        s6.placeholders[0].text = model.collaboration_slide.title
    if len(s6.placeholders) > 1:
        s6.placeholders[1].text = model.collaboration_slide.kicker

    # Card 1: Client roles and approvers (x=0.83, y=1.70, w=3.71, h=4.58)
    _add_card_shape(s6, Inches(0.83), Inches(1.70), Inches(3.71), Inches(4.58))
    box6_1 = s6.shapes.add_textbox(Inches(0.95), Inches(1.80), Inches(3.47), Inches(4.38))
    tf6_1 = box6_1.text_frame
    tf6_1.word_wrap = True

    p6_1 = tf6_1.paragraphs[0]
    p6_1.text = "Client roles and approvers"
    p6_1.font.name = FONT_TITLE
    p6_1.font.size = FONT_SIZE_CARD_HEADER
    p6_1.font.bold = True
    p6_1.font.color.rgb = COLOR_DEEP_NAVY

    for cr in model.collaboration_slide.client_roles:
        p_b = tf6_1.add_paragraph()
        p_b.text = f"• {cr}"
        p_b.font.name = FONT_BODY
        p_b.font.size = Pt(9.5)
        p_b.font.color.rgb = COLOR_SLATE

    if model.collaboration_slide.client_roles_overflow > 0:
        p_ov = tf6_1.add_paragraph()
        p_ov.text = f"• +{model.collaboration_slide.client_roles_overflow} more (see Startup Kit)"
        p_ov.font.name = FONT_BODY
        p_ov.font.size = Pt(9.5)
        p_ov.font.italic = True
        p_ov.font.color.rgb = COLOR_MUTED_GRAY

    # Card 2: Working rhythm (x=4.81, y=1.70, w=3.71, h=4.58)
    _add_card_shape(s6, Inches(4.81), Inches(1.70), Inches(3.71), Inches(4.58))
    box6_2 = s6.shapes.add_textbox(Inches(4.93), Inches(1.80), Inches(3.47), Inches(4.38))
    tf6_2 = box6_2.text_frame
    tf6_2.word_wrap = True

    p6_2 = tf6_2.paragraphs[0]
    p6_2.text = "Working rhythm"
    p6_2.font.name = FONT_TITLE
    p6_2.font.size = FONT_SIZE_CARD_HEADER
    p6_2.font.bold = True
    p6_2.font.color.rgb = COLOR_DEEP_NAVY

    for wr in model.collaboration_slide.working_rhythm:
        p_b = tf6_2.add_paragraph()
        p_b.text = f"• {wr}"
        p_b.font.name = FONT_BODY
        p_b.font.size = Pt(9.5)
        p_b.font.color.rgb = COLOR_SLATE

    if model.collaboration_slide.working_rhythm_overflow > 0:
        p_ov = tf6_2.add_paragraph()
        p_ov.text = f"• +{model.collaboration_slide.working_rhythm_overflow} more (see Startup Kit)"
        p_ov.font.name = FONT_BODY
        p_ov.font.size = Pt(9.5)
        p_ov.font.italic = True
        p_ov.font.color.rgb = COLOR_MUTED_GRAY

    # Card 3: What we need from the client (x=8.79, y=1.70, w=3.71, h=4.58)
    _add_card_shape(s6, Inches(8.79), Inches(1.70), Inches(3.71), Inches(4.58))
    box6_3 = s6.shapes.add_textbox(Inches(8.91), Inches(1.80), Inches(3.47), Inches(4.38))
    tf6_3 = box6_3.text_frame
    tf6_3.word_wrap = True

    p6_3 = tf6_3.paragraphs[0]
    p6_3.text = "What we need from the client"
    p6_3.font.name = FONT_TITLE
    p6_3.font.size = FONT_SIZE_CARD_HEADER
    p6_3.font.bold = True
    p6_3.font.color.rgb = COLOR_DEEP_NAVY

    for pr in model.collaboration_slide.prerequisites:
        p_b = tf6_3.add_paragraph()
        p_b.text = f"• {pr}"
        p_b.font.name = FONT_BODY
        p_b.font.size = Pt(9.5)
        p_b.font.color.rgb = COLOR_SLATE

    if model.collaboration_slide.prerequisites_overflow > 0:
        p_ov = tf6_3.add_paragraph()
        p_ov.text = f"• +{model.collaboration_slide.prerequisites_overflow} more (see Project Delivery Workbook · Project Schedule)"
        p_ov.font.name = FONT_BODY
        p_ov.font.size = Pt(9.5)
        p_ov.font.italic = True
        p_ov.font.color.rgb = COLOR_MUTED_GRAY

    s6.notes_slide.notes_text_frame.text = model.collaboration_slide.speaker_notes

    # Save presentation
    target_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(target_path))

    # Save manifest
    manifest_path = target_path.with_name(f"{target_path.stem}.trace.json")
    model.manifest.write_json(manifest_path)

    return target_path

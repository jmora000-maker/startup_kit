"""Openpyxl Excel workbook renderer for the Project Delivery Workbook (v3 spec)."""

import re
from datetime import date
from pathlib import Path
from typing import List, Optional, Any, Sequence, Dict, Tuple
import openpyxl
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

from src.version import __version__
from src.generators.pmo_workbook.rows import (
    WorkbookModel,
    ScheduleRow,
    WBSRow,
    RAIDRow,
)
from src.generators.pmo_workbook.styles import (
    COLOR_NAVY_HEX,
    COLOR_PRIMARY_BLUE_HEX,
    COLOR_ACCENT_BLUE_HEX,
    COLOR_LIGHT_BG_HEX,
    COLOR_WARNING_BG_HEX,
    COLOR_WHITE_HEX,
    COLOR_LIGHT_RED_HEX,
    COLOR_LIGHT_GREEN_HEX,
    COLOR_TIMELINE_SPAN_HEX,
    FONT_TITLE,
    FONT_SUBTITLE,
    FONT_MUTED,
    FONT_HEADER,
    FONT_BODY,
    FONT_BODY_BOLD,
    FONT_BODY_MUTED,
    FONT_WBS_L1,
    FONT_TIMELINE_HEADER,
    FILL_NAVY,
    FILL_PRIMARY_BLUE,
    FILL_ACCENT_BLUE,
    FILL_LIGHT_BG,
    FILL_WARNING,
    FILL_LIGHT_RED,
    FILL_LIGHT_GREEN,
    FILL_TIMELINE_SPAN,
    ALIGN_HEADER,
    ALIGN_HEADER_TIMELINE,
    ALIGN_LEFT,
    ALIGN_CENTER,
    ALIGN_RIGHT,
    BORDER_ALL_THIN,
)

SCHEDULE_COL_WIDTHS = [
    9,   # A: WBS Code
    12,  # B: Row Type
    26,  # C: Workstream
    12,  # D: Milestone ID
    40,  # E: Milestone
    45,  # F: Milestone Scope
    22,  # G: Owner
    13,  # H: Planned Start
    13,  # I: Planned Finish
    13,  # J: Internal Buffer Date
    13,  # K: External Commitment Date
    24,  # L: Date Basis
    11,  # M: Duration (working days)
    11,  # N: Days to Finish
    13,  # O: Status
    12,  # P: Health
    12,  # Q: Predecessor
    35,  # R: Client Prerequisites
    35,  # S: Critical Path Assumptions
    20,  # T: Linked Deliverables
    18,  # U: Linked RAID IDs
    26,  # V: Source
    40,  # W: Notes
]

WBS_COL_WIDTHS = [
    9,   # A: WBS Code
    7,   # B: Level
    14,  # C: Element Type
    45,  # D: Name
    26,  # E: Workstream
    12,  # F: Milestone ID
    12,  # G: Deliverable ID
    12,  # H: Source ID
    22,  # I: Owner
    13,  # J: Planned Start
    13,  # K: Planned Finish
    13,  # L: Milestone Finish
    13,  # M: Status
    10,  # N: % Complete
    14,  # O: Cadence
    18,  # P: Predecessors
    40,  # Q: Acceptance / Completion Criteria
    18,  # R: Linked RAID IDs
    26,  # S: Source
    22,  # T: Mapping Basis
    40,  # U: Notes
]

RAID_COL_WIDTHS = [
    10,  # A: RAID ID
    12,  # B: Type
    45,  # C: Description
    35,  # D: Contract Reference (v3 A7)
    22,  # E: Category
    26,  # F: Workstream
    16,  # G: Linked Milestone
    16,  # H: Linked WBS Code
    22,  # I: Owner
    12,  # J: Probability
    12,  # K: Impact
    12,  # L: Rating (v3 A11)
    12,  # M: Score (P x I)
    35,  # N: Trigger / Early Warning
    35,  # O: Mitigation / Response
    13,  # P: Due Date
    13,  # Q: Status
    13,  # R: Date Raised
    13,  # S: Last Updated
    20,  # T: Linked Decision
    26,  # U: Linked Dependency / Assumption
    26,  # V: Source
    12,  # W: Source ID
    40,  # X: Notes
]


def _clean_cell_str(val: Any) -> Any:
    """Strip illegal Excel control characters from string values."""
    if isinstance(val, str):
        return ILLEGAL_CHARACTERS_RE.sub("", val)
    return val


def _apply_row_styling(ws: Worksheet, row_idx: int, font: Font, fill: Optional[PatternFill] = None, align: Optional[Alignment] = None, max_col: int = 24):
    """Apply styling across row columns."""
    for col in range(1, max_col + 1):
        cell = ws.cell(row=row_idx, column=col)
        cell.font = font
        cell.border = BORDER_ALL_THIN
        if fill:
            cell.fill = fill
        if align:
            cell.alignment = align


def _add_title_block(ws: Worksheet, model: WorkbookModel, tab_name: str, has_schedule_note: bool = False):
    """Add standard metadata title block rows 1 to 4."""
    # Row 1: Title
    ws["A1"] = f"{model.project_name.upper()} - {tab_name.upper()}"
    ws["A1"].font = FONT_TITLE

    # Row 2: Subtitle
    ws["A2"] = (
        f"Client: {model.client_name} | "
        f"Contract: {model.contract_type} | "
        f"Talent PM: {model.talent_pm} | "
        f"Delivery Manager: {model.delivery_manager}"
    )
    ws["A2"].font = FONT_SUBTITLE

    # Row 3: Meta
    start_str = model.start_date.isoformat() if model.start_date else "TBD"
    gen_str = model.generation_date.isoformat() if model.generation_date else date.today().isoformat()
    ws["A3"] = f"Start Date: {start_str} ({model.start_date_basis}) | Generated {gen_str} from the project baseline"
    ws["A3"].font = FONT_MUTED

    # Row 4: Schedule note if applicable
    if has_schedule_note:
        ws["A4"] = "Planned dates are estimated from the SOW week ranges and the Start Date. External Commitment Date is shown only where the SOW states a date."
        ws["A4"].font = FONT_MUTED
    else:
        ws["A4"] = ""


def _write_lists_sheet(wb: Workbook, model: WorkbookModel):
    """Create hidden _Lists sheet with validation lists and named ranges (v3 A12)."""
    ws = wb.create_sheet(title="_Lists")
    ws.sheet_state = "hidden"

    # 1. Distinct phase workstreams (v3 A1: phase workstreams plus Multiple phases and Cross-phase)
    phase_workstreams: List[str] = []
    for s in model.schedule_rows:
        if s.row_type == "Workstream" and s.workstream:
            if s.workstream not in phase_workstreams:
                phase_workstreams.append(s.workstream)

    for opt in ["Multiple phases", "Cross-phase"]:
        if opt not in phase_workstreams:
            phase_workstreams.append(opt)

    # 2. Distinct Milestone IDs (v3 A12)
    milestone_ids: List[str] = []
    for s in model.schedule_rows:
        if s.row_type == "Milestone" and s.milestone_id:
            if s.milestone_id not in milestone_ids:
                milestone_ids.append(s.milestone_id)

    # 3. Other validation domains
    types = ["Risk", "Assumption", "Issue", "Dependency"]
    categories = [
        "Delivery Risk", "Technical Dependency", "Commercial Assumption",
        "Contract Clarification", "Client Prerequisite", "Scope Gap", "Governance", "Open Question"
    ]
    probabilities = ["High", "Medium", "Low"]
    impacts = ["High", "Medium", "Low"]
    raid_statuses = ["Open", "In Progress", "Monitoring", "Escalated", "Closed"]
    schedule_statuses = ["Not Started", "In Progress", "Complete", "On Hold", "Blocked"]
    cadences = ["Daily", "Weekly", "Biweekly", "Monthly", "Per Milestone", "One-time", "As needed"]

    cols_data = [
        ("List_Workstreams", phase_workstreams),
        ("List_Milestone_IDs", milestone_ids),
        ("List_RAID_Types", types),
        ("List_RAID_Categories", categories),
        ("List_Probabilities", probabilities),
        ("List_Impacts", impacts),
        ("List_RAID_Statuses", raid_statuses),
        ("List_Schedule_Statuses", schedule_statuses),
        ("List_Cadences", cadences),
    ]

    for col_idx, (name, items) in enumerate(cols_data, start=1):
        col_letter = get_column_letter(col_idx)
        for row_idx, val in enumerate(items, start=1):
            ws.cell(row=row_idx, column=col_idx, value=val)
        num_rows = len(items) if items else 1
        formula = f"'_Lists'!${col_letter}$1:${col_letter}${num_rows}"
        wb.defined_names.add(DefinedName(name=name, attr_text=formula))


def _write_schedule_sheet(wb: Workbook, model: WorkbookModel):
    """Render Project Schedule worksheet."""
    ws = wb.active
    ws.title = "Project Schedule"
    ws.views.sheetView[0].showGridLines = True

    _add_title_block(ws, model, "PROJECT SCHEDULE", has_schedule_note=True)

    headers = [
        "WBS Code", "Row Type", "Workstream", "Milestone ID", "Milestone",
        "Milestone Scope", "Owner", "Planned Start", "Planned Finish",
        "Internal Buffer Date", "External Commitment Date", "Date Basis",
        "Duration (working days)", "Days to Finish", "Status", "Health",
        "Predecessor", "Client Prerequisites", "Critical Path Assumptions",
        "Linked Deliverables", "Linked RAID IDs", "Source", "Notes"
    ]

    # Row 5: Column headers
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=5, column=col_idx, value=h)
        cell.font = FONT_HEADER
        cell.fill = FILL_NAVY
        cell.alignment = ALIGN_HEADER
        cell.border = BORDER_ALL_THIN

    # Timeline headers starting at col 24 (X)
    timeline_start_col = 24
    for idx, w_date in enumerate(model.timeline_weeks):
        col_idx = timeline_start_col + idx
        cell = ws.cell(row=5, column=col_idx, value=w_date)
        cell.font = FONT_TIMELINE_HEADER
        cell.fill = FILL_PRIMARY_BLUE
        cell.alignment = ALIGN_HEADER_TIMELINE
        cell.number_format = "YYYY-MM-DD"
        cell.border = BORDER_ALL_THIN

    # Find workstream child row ranges
    ws_child_ranges: Dict[str, Tuple[int, int]] = {}
    current_ws_code: Optional[str] = None
    current_ws_start_row = 0

    cur_row = 6
    for s in model.schedule_rows:
        if s.row_type == "Workstream":
            if current_ws_code and cur_row - 1 >= current_ws_start_row:
                ws_child_ranges[current_ws_code] = (current_ws_start_row, cur_row - 1)
            current_ws_code = s.wbs_code
            current_ws_start_row = cur_row + 1
        cur_row += 1

    if current_ws_code and cur_row - 1 >= current_ws_start_row:
        ws_child_ranges[current_ws_code] = (current_ws_start_row, cur_row - 1)

    # Populate Schedule Data Rows
    row_idx = 6
    for s in model.schedule_rows:
        r = row_idx

        # Col A: WBS Code
        ws.cell(row=r, column=1, value=_clean_cell_str(s.wbs_code))

        # Col B: Row Type
        ws.cell(row=r, column=2, value=_clean_cell_str(s.row_type))

        # Col C: Workstream
        ws.cell(row=r, column=3, value=_clean_cell_str(s.workstream))

        # Col D: Milestone ID
        ws.cell(row=r, column=4, value=_clean_cell_str(s.milestone_id))

        # Col E: Milestone
        ws.cell(row=r, column=5, value=_clean_cell_str(s.name))

        # Col F: Milestone Scope
        ws.cell(row=r, column=6, value=_clean_cell_str(s.scope))

        # Col G: Owner
        ws.cell(row=r, column=7, value=_clean_cell_str(s.owner))

        # Col H: Planned Start
        cell_h = ws.cell(row=r, column=8)
        if s.row_type == "Workstream" and s.wbs_code in ws_child_ranges:
            sr, er = ws_child_ranges[s.wbs_code]
            cell_h.value = f"=MIN(H{sr}:H{er})"
        else:
            cell_h.value = s.planned_start
        cell_h.number_format = "YYYY-MM-DD"

        # Col I: Planned Finish
        cell_i = ws.cell(row=r, column=9)
        if s.row_type == "Workstream" and s.wbs_code in ws_child_ranges:
            sr, er = ws_child_ranges[s.wbs_code]
            cell_i.value = f"=MAX(I{sr}:I{er})"
        else:
            cell_i.value = s.planned_finish
        cell_i.number_format = "YYYY-MM-DD"

        # Col J: Internal Buffer Date
        cell_j = ws.cell(row=r, column=10, value=s.internal_buffer_date)
        cell_j.number_format = "YYYY-MM-DD"

        # Col K: External Commitment Date
        cell_k = ws.cell(row=r, column=11, value=s.external_date)
        cell_k.number_format = "YYYY-MM-DD"

        # Col L: Date Basis
        ws.cell(row=r, column=12, value=_clean_cell_str(s.date_basis))

        # Col M: Duration (working days)
        cell_m = ws.cell(row=r, column=13, value=f'=IF(OR(H{r}="",I{r}=""),"",NETWORKDAYS(H{r},I{r}))')

        # Col N: Days to Finish
        cell_n = ws.cell(row=r, column=14, value=f'=IF(I{r}="","",I{r}-TODAY())')

        # Col O: Status
        ws.cell(row=r, column=15, value=_clean_cell_str(s.status))

        # Col P: Health Formula
        cell_p = ws.cell(row=r, column=16, value=f'=IF(O{r}="Complete","Complete",IF(AND(K{r}="",I{r}=""),"Date TBC",IF(TODAY()>IF(K{r}<>"",K{r},I{r}),"Overdue",IF(AND(J{r}<>"",TODAY()>J{r}),"In Buffer","On Track"))))')

        # Col Q: Predecessor
        ws.cell(row=r, column=17, value=_clean_cell_str(s.predecessor))

        # Col R: Client Prerequisites
        ws.cell(row=r, column=18, value=_clean_cell_str(s.client_prerequisites))

        # Col S: Critical Path Assumptions
        ws.cell(row=r, column=19, value=_clean_cell_str(s.critical_path_assumptions))

        # Col T: Linked Deliverables
        ws.cell(row=r, column=20, value=_clean_cell_str(s.linked_deliverables))

        # Col U: Linked RAID IDs
        ws.cell(row=r, column=21, value=_clean_cell_str(s.linked_raid_ids))

        # Col V: Source
        ws.cell(row=r, column=22, value=_clean_cell_str(s.source))

        # Col W: Notes
        ws.cell(row=r, column=23, value=_clean_cell_str(s.notes))

        # Formatting row
        if s.row_type == "Workstream":
            _apply_row_styling(ws, r, FONT_WBS_L1, fill=FILL_LIGHT_BG, max_col=23)
        else:
            _apply_row_styling(ws, r, FONT_BODY, max_col=23)

        # Center align codes & dates
        for col_c in (1, 2, 4, 8, 9, 10, 11, 13, 14, 15, 16, 17):
            ws.cell(row=r, column=col_c).alignment = ALIGN_CENTER

        # Outline level
        if s.outline_level > 0:
            ws.row_dimensions[r].outline_level = s.outline_level

        # Timeline cells borders
        for t_idx in range(len(model.timeline_weeks)):
            t_col = timeline_start_col + t_idx
            t_cell = ws.cell(row=r, column=t_col)
            t_cell.border = BORDER_ALL_THIN

        row_idx += 1

    last_row = row_idx - 1

    # Widths
    for col_idx, w in enumerate(SCHEDULE_COL_WIDTHS, start=1):
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = w

    for t_idx in range(len(model.timeline_weeks)):
        col_letter = get_column_letter(timeline_start_col + t_idx)
        ws.column_dimensions[col_letter].width = 4.5

    # Freeze panes at F6
    ws.freeze_panes = "F6"

    # Autofilter A5:W{last_row}
    if last_row >= 5:
        ws.auto_filter.ref = f"A5:W{last_row}"

    # Data Validation
    if last_row >= 6:
        dv_status = DataValidation(type="list", formula1="=List_Schedule_Statuses", allow_blank=True)
        ws.add_data_validation(dv_status)
        dv_status.add(f"O6:O{last_row}")

        dv_ws = DataValidation(type="list", formula1="=List_Workstreams", allow_blank=True)
        ws.add_data_validation(dv_ws)
        dv_ws.add(f"C6:C{last_row}")

        # Health Conditional Formatting
        ws.conditional_formatting.add(
            f"P6:P{last_row}",
            CellIsRule(operator="equal", formula=['"On Track"'], fill=FILL_LIGHT_GREEN, font=Font(color="1E4620"))
        )
        ws.conditional_formatting.add(
            f"P6:P{last_row}",
            CellIsRule(operator="equal", formula=['"Complete"'], fill=FILL_LIGHT_GREEN, font=Font(color="1E4620"))
        )
        ws.conditional_formatting.add(
            f"P6:P{last_row}",
            CellIsRule(operator="equal", formula=['"In Buffer"'], fill=FILL_WARNING, font=Font(color="7D4A00"))
        )
        ws.conditional_formatting.add(
            f"P6:P{last_row}",
            CellIsRule(operator="equal", formula=['"Overdue"'], fill=FILL_LIGHT_RED, font=Font(color="9C0006"))
        )
        ws.conditional_formatting.add(
            f"P6:P{last_row}",
            CellIsRule(operator="equal", formula=['"Date TBC"'], fill=FILL_LIGHT_BG, font=FONT_MUTED)
        )

    # Timeline Conditional Formatting
    if model.timeline_weeks and last_row >= 6:
        first_t_col = get_column_letter(timeline_start_col)
        last_t_col = get_column_letter(timeline_start_col + len(model.timeline_weeks) - 1)
        t_range = f"{first_t_col}6:{last_t_col}{last_row}"

        rule_finish = FormulaRule(
            formula=[f"AND($I6>={first_t_col}$5,$I6<{first_t_col}$5+7)"],
            fill=FILL_NAVY
        )
        rule_buffer = FormulaRule(
            formula=[f"AND($J6<>\"\",$J6>={first_t_col}$5,$J6<{first_t_col}$5+7)"],
            fill=FILL_ACCENT_BLUE
        )
        rule_span = FormulaRule(
            formula=[f"AND($H6<>\"\",$I6<>\"\",{first_t_col}$5>=$H6-WEEKDAY($H6,2)+1,{first_t_col}$5<=$I6)"],
            fill=FILL_TIMELINE_SPAN
        )

        ws.conditional_formatting.add(t_range, rule_finish)
        ws.conditional_formatting.add(t_range, rule_buffer)
        ws.conditional_formatting.add(t_range, rule_span)


def _write_wbs_sheet(wb: Workbook, model: WorkbookModel):
    """Render Work Breakdown Structure worksheet."""
    ws = wb.create_sheet(title="WBS")
    ws.views.sheetView[0].showGridLines = True

    _add_title_block(ws, model, "WORK BREAKDOWN STRUCTURE", has_schedule_note=False)

    headers = [
        "WBS Code", "Level", "Element Type", "Name", "Workstream",
        "Milestone ID", "Deliverable ID", "Source ID", "Owner", "Planned Start",
        "Planned Finish", "Milestone Finish", "Status", "% Complete", "Cadence",
        "Predecessors", "Acceptance / Completion Criteria", "Linked RAID IDs",
        "Source", "Mapping Basis", "Notes"
    ]

    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=5, column=col_idx, value=h)
        cell.font = FONT_HEADER
        cell.fill = FILL_NAVY
        cell.alignment = ALIGN_HEADER
        cell.border = BORDER_ALL_THIN

    row_idx = 6
    for w in model.wbs_rows:
        r = row_idx
        ws.cell(row=r, column=1, value=_clean_cell_str(w.wbs_code))
        ws.cell(row=r, column=2, value=w.level)
        ws.cell(row=r, column=3, value=_clean_cell_str(w.element_type))
        ws.cell(row=r, column=4, value=_clean_cell_str(w.name))
        ws.cell(row=r, column=5, value=_clean_cell_str(w.workstream))
        ws.cell(row=r, column=6, value=_clean_cell_str(w.milestone_id))
        ws.cell(row=r, column=7, value=_clean_cell_str(w.deliverable_id))
        ws.cell(row=r, column=8, value=_clean_cell_str(w.source_id))
        ws.cell(row=r, column=9, value=_clean_cell_str(w.owner))

        # Dates
        cell_j = ws.cell(row=r, column=10, value=w.planned_start)
        cell_k = ws.cell(row=r, column=11, value=w.planned_finish)
        cell_l = ws.cell(row=r, column=12, value=w.milestone_date)
        cell_j.number_format = "YYYY-MM-DD"
        cell_k.number_format = "YYYY-MM-DD"
        cell_l.number_format = "YYYY-MM-DD"

        # Col M: Status
        ws.cell(row=r, column=13, value=_clean_cell_str(w.status))

        # Col N: % Complete
        cell_n = ws.cell(row=r, column=14, value=1.0 if w.status == "Complete" else 0.0)
        cell_n.number_format = "0%"

        ws.cell(row=r, column=15, value=_clean_cell_str(w.cadence))
        ws.cell(row=r, column=16, value=_clean_cell_str(w.predecessors))
        ws.cell(row=r, column=17, value=_clean_cell_str(w.acceptance_criteria))
        ws.cell(row=r, column=18, value=_clean_cell_str(w.linked_raid_ids))
        ws.cell(row=r, column=19, value=_clean_cell_str(w.source))
        ws.cell(row=r, column=20, value=_clean_cell_str(w.mapping_basis))
        ws.cell(row=r, column=21, value=_clean_cell_str(w.notes))

        # Styling
        if w.level == 1:
            _apply_row_styling(ws, r, FONT_WBS_L1, fill=FILL_LIGHT_BG, max_col=21)
        elif w.level == 2:
            _apply_row_styling(ws, r, FONT_BODY_BOLD, max_col=21)
        elif w.level == 3:
            _apply_row_styling(ws, r, FONT_BODY_BOLD, max_col=21)
        else:
            _apply_row_styling(ws, r, FONT_BODY, max_col=21)

        # Center alignment
        for col_c in (1, 2, 3, 6, 7, 8, 10, 11, 12, 13, 14, 15, 16):
            ws.cell(row=r, column=col_c).alignment = ALIGN_CENTER

        # Evidence consistency warning fill on Criteria cell (v3 A9)
        if w.level == 3 and w.deliverable_id and w.deliverable_id in model.flagged_evidence_deliverables:
            ws.cell(row=r, column=17).fill = FILL_WARNING

        # Outline level
        if w.outline_level > 0:
            ws.row_dimensions[r].outline_level = w.outline_level

        row_idx += 1

    last_row = row_idx - 1

    for col_idx, w in enumerate(WBS_COL_WIDTHS, start=1):
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = w

    # Freeze panes at E6
    ws.freeze_panes = "E6"

    if last_row >= 5:
        ws.auto_filter.ref = f"A5:U{last_row}"

    # Data validation
    if last_row >= 6:
        dv_status = DataValidation(type="list", formula1="=List_Schedule_Statuses", allow_blank=True)
        ws.add_data_validation(dv_status)
        dv_status.add(f"M6:M{last_row}")

        dv_cad = DataValidation(type="list", formula1="=List_Cadences", allow_blank=True)
        ws.add_data_validation(dv_cad)
        dv_cad.add(f"O6:O{last_row}")


def _write_raid_sheet(wb: Workbook, model: WorkbookModel):
    """Render RAID Log worksheet (v3 A7, A8, A11, A12)."""
    ws = wb.create_sheet(title="RAID Log")
    ws.views.sheetView[0].showGridLines = True

    _add_title_block(ws, model, "RAID LOG", has_schedule_note=False)

    headers = [
        "RAID ID", "Type", "Description", "Contract Reference", "Category", "Workstream",
        "Linked Milestone", "Linked WBS Code", "Owner", "Probability", "Impact",
        "Rating", "Score (P x I)", "Trigger / Early Warning", "Mitigation / Response",
        "Due Date", "Status", "Date Raised", "Last Updated", "Linked Decision",
        "Linked Dependency / Assumption", "Source", "Source ID", "Notes"
    ]

    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=5, column=col_idx, value=h)
        cell.font = FONT_HEADER
        cell.fill = FILL_NAVY
        cell.alignment = ALIGN_HEADER
        cell.border = BORDER_ALL_THIN

    row_idx = 6

    if not model.raid_rows:
        # Empty placeholder row
        r = 6
        ws.cell(row=r, column=1, value="RAID-01")
        ws.cell(row=r, column=2, value="Risk")
        ws.cell(row=r, column=3, value="No baseline risks identified; populate during Startup Gateway.")
        ws.cell(row=r, column=4, value="")
        ws.cell(row=r, column=5, value="Delivery Risk")
        ws.cell(row=r, column=6, value="Cross-phase")
        ws.cell(row=r, column=7, value="")
        ws.cell(row=r, column=8, value="")
        ws.cell(row=r, column=9, value="Talent PM")
        ws.cell(row=r, column=10, value="Low")
        ws.cell(row=r, column=11, value="Low")
        ws.cell(row=r, column=12, value=f'=IF(OR(J{r}="",K{r}=""),"",IF(OR(AND(J{r}="High",K{r}="High"),AND(J{r}="High",K{r}="Medium"),AND(J{r}="Medium",K{r}="High")),"High",IF(AND(J{r}="Low",K{r}="Low"),"Low","Medium")))')
        ws.cell(row=r, column=13, value=f'=IF(OR(J{r}="",K{r}=""),"",(IF(J{r}="High",3,IF(J{r}="Medium",2,1)))*(IF(K{r}="High",3,IF(K{r}="Medium",2,1))))')
        ws.cell(row=r, column=14, value="")
        ws.cell(row=r, column=15, value="")
        ws.cell(row=r, column=16, value=None)
        ws.cell(row=r, column=17, value="Open")
        ws.cell(row=r, column=18, value=model.generation_date)
        ws.cell(row=r, column=19, value=model.generation_date)
        ws.cell(row=r, column=20, value="")
        ws.cell(row=r, column=21, value="")
        ws.cell(row=r, column=22, value="PM Best Practice")
        ws.cell(row=r, column=23, value="")
        ws.cell(row=r, column=24, value="Placeholder RAID item")
        _apply_row_styling(ws, r, FONT_BODY, max_col=24)
        row_idx = 7
    else:
        for r_row in model.raid_rows:
            r = row_idx
            ws.cell(row=r, column=1, value=_clean_cell_str(r_row.raid_id))
            ws.cell(row=r, column=2, value=_clean_cell_str(r_row.type))
            ws.cell(row=r, column=3, value=_clean_cell_str(r_row.description))
            ws.cell(row=r, column=4, value=_clean_cell_str(r_row.contract_reference))
            ws.cell(row=r, column=5, value=_clean_cell_str(r_row.category))
            ws.cell(row=r, column=6, value=_clean_cell_str(r_row.workstream))
            ws.cell(row=r, column=7, value=_clean_cell_str(r_row.linked_milestone))
            ws.cell(row=r, column=8, value=_clean_cell_str(r_row.linked_wbs_code))
            ws.cell(row=r, column=9, value=_clean_cell_str(r_row.owner))
            ws.cell(row=r, column=10, value=_clean_cell_str(r_row.probability))
            ws.cell(row=r, column=11, value=_clean_cell_str(r_row.impact))

            # Col L (12): Rating (v3 A11)
            if r_row.severity:
                ws.cell(row=r, column=12, value=_clean_cell_str(r_row.severity))
            else:
                ws.cell(row=r, column=12, value=f'=IF(OR(J{r}="",K{r}=""),"",IF(OR(AND(J{r}="High",K{r}="High"),AND(J{r}="High",K{r}="Medium"),AND(J{r}="Medium",K{r}="High")),"High",IF(AND(J{r}="Low",K{r}="Low"),"Low","Medium")))')

            # Col M (13): Score (P x I) Formula
            ws.cell(row=r, column=13, value=f'=IF(OR(J{r}="",K{r}=""),"",(IF(J{r}="High",3,IF(J{r}="Medium",2,1)))*(IF(K{r}="High",3,IF(K{r}="Medium",2,1))))')

            ws.cell(row=r, column=14, value=_clean_cell_str(r_row.trigger_or_early_warning))
            ws.cell(row=r, column=15, value=_clean_cell_str(r_row.mitigation_or_response))

            cell_p = ws.cell(row=r, column=16, value=r_row.due_date)
            cell_p.number_format = "YYYY-MM-DD"

            ws.cell(row=r, column=17, value=_clean_cell_str(r_row.status))

            cell_r = ws.cell(row=r, column=18, value=r_row.date_raised)
            cell_s = ws.cell(row=r, column=19, value=r_row.last_updated)
            cell_r.number_format = "YYYY-MM-DD"
            cell_s.number_format = "YYYY-MM-DD"

            ws.cell(row=r, column=20, value=_clean_cell_str(r_row.linked_decision))
            ws.cell(row=r, column=21, value=_clean_cell_str(r_row.linked_dependency_or_assumption))
            ws.cell(row=r, column=22, value=_clean_cell_str(r_row.source))
            ws.cell(row=r, column=23, value=_clean_cell_str(r_row.source_id))
            ws.cell(row=r, column=24, value=_clean_cell_str(r_row.notes))

            _apply_row_styling(ws, r, FONT_BODY, max_col=24)

            # Center alignment
            for col_c in (1, 2, 5, 7, 8, 10, 11, 12, 13, 16, 17, 18, 19, 23):
                ws.cell(row=r, column=col_c).alignment = ALIGN_CENTER

            row_idx += 1

    last_row = row_idx - 1

    for col_idx, w in enumerate(RAID_COL_WIDTHS, start=1):
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = w

    # Freeze panes at E6
    ws.freeze_panes = "E6"

    if last_row >= 5:
        ws.auto_filter.ref = f"A5:X{last_row}"

    # Data Validations (v3 A7 shifted columns, v3 A12 List_Milestone_IDs)
    if last_row >= 6:
        dv_type = DataValidation(type="list", formula1="=List_RAID_Types", allow_blank=True)
        ws.add_data_validation(dv_type)
        dv_type.add(f"B6:B{last_row}")

        dv_cat = DataValidation(type="list", formula1="=List_RAID_Categories", allow_blank=True)
        ws.add_data_validation(dv_cat)
        dv_cat.add(f"E6:E{last_row}")

        dv_ws = DataValidation(type="list", formula1="=List_Workstreams", allow_blank=True)
        ws.add_data_validation(dv_ws)
        dv_ws.add(f"F6:F{last_row}")

        dv_ms = DataValidation(type="list", formula1="=List_Milestone_IDs", allow_blank=True)
        ws.add_data_validation(dv_ms)
        dv_ms.add(f"G6:G{last_row}")

        dv_prob = DataValidation(type="list", formula1="=List_Probabilities", allow_blank=True)
        ws.add_data_validation(dv_prob)
        dv_prob.add(f"J6:J{last_row}")

        dv_imp = DataValidation(type="list", formula1="=List_Impacts", allow_blank=True)
        ws.add_data_validation(dv_imp)
        dv_imp.add(f"K6:K{last_row}")

        dv_status = DataValidation(type="list", formula1="=List_RAID_Statuses", allow_blank=True)
        ws.add_data_validation(dv_status)
        dv_status.add(f"Q6:Q{last_row}")

        # Rating & Status Conditional Formatting
        ws.conditional_formatting.add(
            f"L6:L{last_row}",
            CellIsRule(operator="equal", formula=['"High"'], fill=FILL_LIGHT_RED, font=Font(color="9C0006", bold=True))
        )
        ws.conditional_formatting.add(
            f"L6:L{last_row}",
            CellIsRule(operator="equal", formula=['"Medium"'], fill=FILL_WARNING, font=Font(color="7D4A00"))
        )
        ws.conditional_formatting.add(
            f"L6:L{last_row}",
            CellIsRule(operator="equal", formula=['"Low"'], fill=FILL_LIGHT_GREEN, font=Font(color="1E4620"))
        )

        ws.conditional_formatting.add(
            f"Q6:Q{last_row}",
            CellIsRule(operator="equal", formula=['"Closed"'], fill=FILL_LIGHT_BG, font=FONT_MUTED)
        )
        ws.conditional_formatting.add(
            f"Q6:Q{last_row}",
            CellIsRule(operator="equal", formula=['"Escalated"'], fill=FILL_LIGHT_RED, font=Font(color="9C0006", bold=True))
        )


def write_workbook(model: WorkbookModel, target_path: Path) -> Path:
    """Render complete WorkbookModel into target_path Excel workbook."""
    wb = Workbook()

    # Document properties
    wb.properties.title = f"{model.project_name} Project Delivery Workbook"
    wb.properties.creator = f"Toptal PMO Generator v{__version__}"

    # Render Sheets
    _write_schedule_sheet(wb, model)
    _write_wbs_sheet(wb, model)
    _write_raid_sheet(wb, model)
    _write_lists_sheet(wb, model)

    target_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(target_path))
    return target_path

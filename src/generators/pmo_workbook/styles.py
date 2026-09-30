"""Workbook styling constants and openpyxl style helpers."""

from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from src.generators.formatting import (
    COLOR_NAVY_HEX,
    COLOR_PRIMARY_BLUE_HEX,
    COLOR_ACCENT_BLUE_HEX,
    COLOR_LIGHT_BG_HEX,
    COLOR_WARNING_BG_HEX,
    COLOR_BORDER_HEX,
    COLOR_TEXT_MUTED_HEX,
    ACTION_TAG_REGEX,
)

# Colors
COLOR_WHITE_HEX = "FFFFFF"
COLOR_LIGHT_RED_HEX = "FEE2E2"
COLOR_LIGHT_GREEN_HEX = "DCFCE7"
COLOR_TIMELINE_SPAN_HEX = "DBEAFE"
COLOR_TEXT_GREY_HEX = "9CA3AF"

# Fonts
FONT_TITLE = Font(name="Arial", size=16, bold=True, color=COLOR_NAVY_HEX)
FONT_SUBTITLE = Font(name="Arial", size=10, bold=True, color=COLOR_NAVY_HEX)
FONT_MUTED = Font(name="Arial", size=9, italic=True, color=COLOR_TEXT_MUTED_HEX)
FONT_HEADER = Font(name="Arial", size=9.5, bold=True, color=COLOR_WHITE_HEX)
FONT_BODY = Font(name="Arial", size=9)
FONT_BODY_BOLD = Font(name="Arial", size=9, bold=True)
FONT_BODY_MUTED = Font(name="Arial", size=9, color=COLOR_TEXT_MUTED_HEX)
FONT_BODY_GREY = Font(name="Arial", size=9, color=COLOR_TEXT_GREY_HEX)
FONT_WBS_L1 = Font(name="Arial", size=9.5, bold=True, color=COLOR_NAVY_HEX)
FONT_TIMELINE_HEADER = Font(name="Arial", size=8.5, bold=True, color=COLOR_WHITE_HEX)

# Fills
FILL_NAVY = PatternFill(start_color=COLOR_NAVY_HEX, end_color=COLOR_NAVY_HEX, fill_type="solid")
FILL_PRIMARY_BLUE = PatternFill(start_color=COLOR_PRIMARY_BLUE_HEX, end_color=COLOR_PRIMARY_BLUE_HEX, fill_type="solid")
FILL_ACCENT_BLUE = PatternFill(start_color=COLOR_ACCENT_BLUE_HEX, end_color=COLOR_ACCENT_BLUE_HEX, fill_type="solid")
FILL_LIGHT_BG = PatternFill(start_color=COLOR_LIGHT_BG_HEX, end_color=COLOR_LIGHT_BG_HEX, fill_type="solid")
FILL_WARNING = PatternFill(start_color=COLOR_WARNING_BG_HEX, end_color=COLOR_WARNING_BG_HEX, fill_type="solid")
FILL_LIGHT_RED = PatternFill(start_color=COLOR_LIGHT_RED_HEX, end_color=COLOR_LIGHT_RED_HEX, fill_type="solid")
FILL_LIGHT_GREEN = PatternFill(start_color=COLOR_LIGHT_GREEN_HEX, end_color=COLOR_LIGHT_GREEN_HEX, fill_type="solid")
FILL_TIMELINE_SPAN = PatternFill(start_color=COLOR_TIMELINE_SPAN_HEX, end_color=COLOR_TIMELINE_SPAN_HEX, fill_type="solid")

# Alignments
ALIGN_HEADER = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_HEADER_TIMELINE = Alignment(horizontal="center", vertical="center", text_rotation=90)
ALIGN_LEFT = Alignment(horizontal="left", vertical="top", wrap_text=True)
ALIGN_CENTER = Alignment(horizontal="center", vertical="top")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="top")

# Borders
BORDER_THIN_SIDE = Side(style="thin", color=COLOR_BORDER_HEX)
BORDER_ALL_THIN = Border(
    left=BORDER_THIN_SIDE,
    right=BORDER_THIN_SIDE,
    top=BORDER_THIN_SIDE,
    bottom=BORDER_THIN_SIDE
)

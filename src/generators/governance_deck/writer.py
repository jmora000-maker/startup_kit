"""PPTX Presentation writer and dynamic shape tree engine for PS Governance Kickoff Deck (HTL-72)."""

import copy
import os
from pathlib import Path
from typing import List, Optional, Tuple
import pptx
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Length, Pt
from pptx.dml.color import RGBColor

from src.config import CLIENT_DECK_TEMPLATE_PATH, GOVERNANCE_DECK_TEMPLATE_PATH
from src.generators.client_deck.fit import (
    estimate_text_height,
    fit_deliverables_for_block,
    select_fitting_font_size,
    truncate_at_clause,
    wrap_lines,
)
from src.generators.client_deck.spec import (
    ClientDeckModel,
    Slide1Model as ClientSlide1Model,
    Slide5Model as ClientSlide5Model,
    Slide6Model as ClientSlide6Model,
    Slide12Model as ClientSlide12Model,
    Slide15Model as ClientSlide15Model,
    RiskItem as ClientRiskItem,
    NextStepItem as ClientNextStepItem,
    WorkstreamBlock as ClientWorkstreamBlock,
)
from src.generators.client_deck.writer import (
    _format_table_cell,
    _replace_shape_text,
    _set_header_paragraph_format,
    _set_paragraph_bullet_format,
    _populate_deliverables_box,
    write_slide5,
    write_slide6,
    write_slide12,
    write_slide15,
)
from src.generators.governance_deck.spec import GovernanceDeckModel


def resolve_governance_deck_template_path(custom_path: Optional[Path] = None) -> Path:
    """Resolve the PS Governance Kickoff Deck template path with fallback resolution (HTL-72)."""
    if custom_path and custom_path.exists():
        return custom_path

    env_path = os.getenv("GOVERNANCE_DECK_TEMPLATE_PATH")
    if env_path:
        p = Path(env_path)
        if p.exists():
            return p

    primary_path = Path("templates/PS_Governance Kickoff_TEMPLATE.pptx")
    if primary_path.exists():
        return primary_path

    fallback_path = Path("templates/PS_Governance_Kickoff_TEMPLATE.pptx")
    if fallback_path.exists():
        return fallback_path

    return primary_path


def resolve_client_deck_template_path(custom_path: Optional[Path] = None) -> Path:
    """Resolve the Client Kickoff Deck template path with fallback resolution (HTL-72)."""
    if custom_path and custom_path.exists():
        return custom_path
    if CLIENT_DECK_TEMPLATE_PATH.exists():
        return CLIENT_DECK_TEMPLATE_PATH
    fallback = Path("templates/Project_Kickoff_Deck_TEMPLATE.pptx")
    if fallback.exists():
        return fallback
    return CLIENT_DECK_TEMPLATE_PATH


def write_slide1(slide, model: GovernanceDeckModel) -> None:
    """Populate Slide 1 [Client Name] and [Project Name] placeholders with 'Governance Kickoff' and fixed subtitle (HTL-72)."""
    candidate_sizes = [28.0, 24.0, 22.0, 20.0, 18.0, 16.0, 14.0, 13.0, 12.0, 11.0, 10.0]
    client_name = model.slide1.client_name or "Client"
    project_name = model.slide1.project_name or "Project Baseline"

    # Composite text for title line-wrapping and height estimation
    composite_title = f"{client_name}\n{project_name} Governance Kickoff"

    # Available geometry on Slide 1:
    # Title box top = 2.966 in (2,712,227 EMUs), Layout divider line top = 4.423 in (4,044,375 EMUs).
    # To guarantee generous positive visual clearance (>= 0.40 in) above the divider line and subtitle (top = 4.542 in),
    # max_height_in is constrained to 1.00 in across candidate sizes.
    title_sz, _ = select_fitting_font_size(
        composite_title,
        width_in=3.30,
        max_height_in=1.00,
        candidate_sizes=candidate_sizes,
        is_bold=True,
    )

    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        text = shape.text_frame.text

        # Title Shape
        if "[Client Name]" in text or "[Project Name]" in text or shape.shape_id == 93:
            tf = shape.text_frame
            tf.word_wrap = True
            tf.margin_top = 0
            tf.margin_bottom = 0
            tf.margin_left = 0
            tf.margin_right = 0

            # Strip normAutofit if present and enforce top anchor
            for bpr in shape._element.xpath(".//a:bodyPr"):
                bpr.set("anchor", "t")
                for autofit in list(bpr.xpath("./a:normAutofit")):
                    bpr.remove(autofit)
                if not bpr.xpath("./a:noAutofit"):
                    from lxml import etree
                    ns_a = "http://schemas.openxmlformats.org/drawingml/2006/main"
                    bpr.append(etree.Element(f"{{{ns_a}}}noAutofit"))

            # Clear paragraphs except first
            for p in list(tf.paragraphs[1:]):
                tf._txBody.remove(p._p)

            p0 = tf.paragraphs[0]
            p0.text = f"{client_name}\n{project_name} Governance Kickoff"
            for r in p0.runs:
                r.font.name = "Proxima Nova"
                r.font.size = Pt(title_sz)
                r.font.bold = True

        # Subtitle Shape
        elif "Onboarding for Project Managers" in text or "Aligning on governance" in text or "Alignment on governance" in text or "Aligning on outcomes" in text or shape.shape_id == 94:
            tf = shape.text_frame
            tf.word_wrap = True
            tf.margin_top = 0
            tf.margin_bottom = 0
            tf.margin_left = 0
            tf.margin_right = 0

            for bpr in shape._element.xpath(".//a:bodyPr"):
                bpr.set("anchor", "t")

            for p in list(tf.paragraphs[1:]):
                tf._txBody.remove(p._p)

            p0 = tf.paragraphs[0]
            p0.text = "Aligning on governance, outcomes and project plan."
            for r in p0.runs:
                r.font.name = "Proxima Nova"
                r.font.size = Pt(11)


def write_slide15_acceptance_criteria(slide, model: GovernanceDeckModel) -> None:
    """Populate Slide 15 Deliverable Acceptance Criteria table with multi-bullet criteria and uniform font fitting (HTL-62/63/64/72)."""
    # 1. Clear any placeholder shapes
    for s in list(slide.shapes):
        slide.shapes._spTree.remove(s._element)

    # 2. Add Kicker
    kicker_box = slide.shapes.add_textbox(457200, 384048, 6858300, 205800)
    tf_k = kicker_box.text_frame
    tf_k.word_wrap = True
    tf_k.margin_top = 0
    tf_k.margin_bottom = 0
    tf_k.margin_left = 0
    tf_k.margin_right = 0
    p_k = tf_k.paragraphs[0]
    p_k.text = "PROJECT SUCCESS AGREEMENT"
    for r in p_k.runs:
        r.font.name = "Proxima Nova"
        r.font.size = Pt(11)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0x04, 0x11, 0x33)

    # 3. Add Title
    title_box = slide.shapes.add_textbox(457200, 640080, 8023800, 332400)
    tf_t = title_box.text_frame
    tf_t.word_wrap = True
    tf_t.margin_top = 0
    tf_t.margin_bottom = 0
    tf_t.margin_left = 0
    tf_t.margin_right = 0
    p_t = tf_t.paragraphs[0]
    p_t.text = "Deliverable acceptance criteria"
    for r in p_t.runs:
        r.font.name = "Proxima Nova"
        r.font.size = Pt(20)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0x04, 0x11, 0x33)

    # 4. Add Subtitle
    sub_box = slide.shapes.add_textbox(457200, 1042416, 8023800, 274200)
    tf_s = sub_box.text_frame
    tf_s.word_wrap = True
    tf_s.margin_top = 0
    tf_s.margin_bottom = 0
    tf_s.margin_left = 0
    tf_s.margin_right = 0
    p_s = tf_s.paragraphs[0]
    p_s.text = "Contractual criteria and verification requirements mapped to milestone gates."
    for r in p_s.runs:
        r.font.name = "Proxima Nova"
        r.font.size = Pt(11)
        r.font.color.rgb = RGBColor(0x47, 0x54, 0x67)

    # 5. Add Table (Left: 0.50 in, Top: 1.52 in, Width: 9.00 in, Height: 2.80 in)
    rows = model.slide15.rows
    num_rows = len(rows)
    table_rows = max(2, num_rows + 1)
    table_shape = slide.shapes.add_table(table_rows, 3, 457200, 1389888, 8229600, 2560320)
    table = table_shape.table

    # Column widths: Col 0: 2.00 in (1,828,800), Col 1: 5.80 in (5,303,520), Col 2: 1.20 in (1,097,280)
    table.columns[0].width = 1828800
    table.columns[1].width = 5303520
    table.columns[2].width = 1097280

    # Header Row (Row 0)
    headers = ["Deliverable", "Acceptance Criteria", "Milestone Gate"]
    for col_idx, h_text in enumerate(headers):
        cell = table.rows[0].cells[col_idx]
        _format_table_cell(cell, h_text, font_size_pt=10.0, bold=True, color_rgb=RGBColor(0x04, 0x11, 0x33))

    if not rows:
        cell0 = table.rows[1].cells[0]
        _format_table_cell(cell0, "[CONFIRMATION REQUIRED]", font_size_pt=8.0)
        return

    # Pass 1: Compute uniform table font size
    cell_candidate_sizes = [8.5, 8.0, 7.5, 7.0]
    chosen_sz = 8.0
    for sz in cell_candidate_sizes:
        fits = True
        for r in rows:
            crit_text = "\n".join(f"• {c}" for c in r.acceptance_criteria) if len(r.acceptance_criteria) > 1 else (r.acceptance_criteria[0] if r.acceptance_criteria else "")
            h_pt = estimate_text_height(crit_text, width_in=5.6, size_pt=sz)
            if h_pt > 36.0:  # max row height in pt for 6 deliverables
                fits = False
                break
        if fits:
            chosen_sz = sz
            break

    # Pass 2: Populate Data Rows
    for r_idx, row_data in enumerate(rows):
        row = table.rows[r_idx + 1]

        # Col 0: Deliverable (ID bold + Name)
        cell0 = row.cells[0]
        tf0 = cell0.text_frame
        tf0.word_wrap = True
        tf0.margin_top = 0
        tf0.margin_bottom = 0
        tf0.margin_left = 0
        tf0.margin_right = 0
        for tcPr in cell0._tc.xpath(".//a:tcPr"):
            tcPr.set("anchor", "t")
            tcPr.set("marT", "18288")
            tcPr.set("marB", "18288")
            tcPr.set("marL", "36576")
            tcPr.set("marR", "36576")

        p0 = tf0.paragraphs[0]
        p0.space_before = Pt(0)
        p0.space_after = Pt(0)
        p0.text = ""
        r_id = p0.add_run()
        r_id.text = row_data.deliverable_id
        r_id.font.name = "Proxima Nova"
        r_id.font.size = Pt(chosen_sz)
        r_id.font.bold = True
        r_id.font.color.rgb = RGBColor(0x04, 0x11, 0x33)

        if row_data.deliverable_name:
            r_name = p0.add_run()
            r_name.text = f"\n{row_data.deliverable_name}"
            r_name.font.name = "Proxima Nova"
            r_name.font.size = Pt(7.5)
            r_name.font.bold = False
            r_name.font.color.rgb = RGBColor(0x47, 0x54, 0x67)

        # Col 1: Acceptance Criteria (Separate countable bullets)
        cell1 = row.cells[1]
        tf1 = cell1.text_frame
        tf1.word_wrap = True
        tf1.margin_top = 0
        tf1.margin_bottom = 0
        tf1.margin_left = 0
        tf1.margin_right = 0
        for tcPr in cell1._tc.xpath(".//a:tcPr"):
            tcPr.set("anchor", "t")
            tcPr.set("marT", "18288")
            tcPr.set("marB", "18288")
            tcPr.set("marL", "36576")
            tcPr.set("marR", "36576")

        if not row_data.acceptance_criteria:
            p = tf1.paragraphs[0]
            p.space_before = Pt(0)
            p.space_after = Pt(0)
            p.text = "[CONFIRMATION REQUIRED]"
            for r in p.runs:
                r.font.name = "Proxima Nova"
                r.font.size = Pt(chosen_sz)
                r.font.color.rgb = RGBColor(0x47, 0x54, 0x67)
        else:
            use_bullets = len(row_data.acceptance_criteria) > 1
            for c_idx, crit in enumerate(row_data.acceptance_criteria):
                p = tf1.paragraphs[0] if c_idx == 0 else tf1.add_paragraph()
                p.space_before = Pt(0)
                p.space_after = Pt(1)
                p.text = f"• {crit}" if use_bullets else crit
                for r in p.runs:
                    r.font.name = "Proxima Nova"
                    r.font.size = Pt(chosen_sz)
                    r.font.color.rgb = RGBColor(0x04, 0x11, 0x33)

        # Col 2: Milestone Gate
        cell2 = row.cells[2]
        _format_table_cell(cell2, row_data.milestone_gate, font_size_pt=chosen_sz, color_rgb=RGBColor(0x04, 0x11, 0x33))


def write_slide16_acceptance_narrative(slide, model: GovernanceDeckModel) -> None:
    """Populate Slide 16 How Milestone Acceptance Works narrative with Proxima Nova typography and font fitting (HTL-72)."""
    # 1. Clear placeholder shapes
    for s in list(slide.shapes):
        slide.shapes._spTree.remove(s._element)

    # 2. Add Kicker
    kicker_box = slide.shapes.add_textbox(457200, 384048, 6858300, 205800)
    tf_k = kicker_box.text_frame
    tf_k.word_wrap = True
    tf_k.margin_top = 0
    tf_k.margin_bottom = 0
    tf_k.margin_left = 0
    tf_k.margin_right = 0
    p_k = tf_k.paragraphs[0]
    p_k.text = "PROJECT SUCCESS AGREEMENT"
    for r in p_k.runs:
        r.font.name = "Proxima Nova"
        r.font.size = Pt(11)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0x04, 0x11, 0x33)

    # 3. Add Title
    title_box = slide.shapes.add_textbox(457200, 640080, 8023800, 332400)
    tf_t = title_box.text_frame
    tf_t.word_wrap = True
    tf_t.margin_top = 0
    tf_t.margin_bottom = 0
    tf_t.margin_left = 0
    tf_t.margin_right = 0
    p_t = tf_t.paragraphs[0]
    p_t.text = "How milestone acceptance works"
    for r in p_t.runs:
        r.font.name = "Proxima Nova"
        r.font.size = Pt(20)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0x04, 0x11, 0x33)

    # 4. Add Subtitle
    sub_box = slide.shapes.add_textbox(457200, 1042416, 8023800, 274200)
    tf_s = sub_box.text_frame
    tf_s.word_wrap = True
    tf_s.margin_top = 0
    tf_s.margin_bottom = 0
    tf_s.margin_left = 0
    tf_s.margin_right = 0
    p_s = tf_s.paragraphs[0]
    p_s.text = "Workflow, review turnarounds, remediation thresholds, and sign-off protocols for the Talent Team."
    for r in p_s.runs:
        r.font.name = "Proxima Nova"
        r.font.size = Pt(11)
        r.font.color.rgb = RGBColor(0x47, 0x54, 0x67)

    # 5. Add Body Narrative Panel
    narrative_text = model.slide16.narrative or "[CONFIRMATION REQUIRED]"
    paragraphs = [p.strip() for p in narrative_text.split("\n") if p.strip()]

    # Font fitting across candidate sizes
    candidate_sizes = [12.0, 11.5, 11.0, 10.5, 10.0, 9.5, 9.0, 8.5]
    chosen_sz = 10.0
    for sz in candidate_sizes:
        total_h = 0.0
        for p in paragraphs:
            lines = wrap_lines(p, width_in=8.8, size_pt=sz)
            h_pt = lines * sz * 1.25 + 8.0
            total_h += h_pt
        if (total_h / 72.0) <= 3.55:  # Available height is ~3.60 in (from 1.45 to 5.05 in)
            chosen_sz = sz
            break

    # Body Text Box (Left: 0.50 in, Top: 1.45 in, Width: 9.00 in, Height: 3.60 in)
    body_box = slide.shapes.add_textbox(457200, 1325880, 8229600, 3291840)
    tf_b = body_box.text_frame
    tf_b.word_wrap = True
    tf_b.margin_top = 0
    tf_b.margin_bottom = 0
    tf_b.margin_left = 0
    tf_b.margin_right = 0
    for bpr in body_box._element.xpath(".//a:bodyPr"):
        bpr.set("anchor", "t")

    for p_idx, p_text in enumerate(paragraphs):
        p_elem = tf_b.paragraphs[0] if p_idx == 0 else tf_b.add_paragraph()
        p_elem.space_before = Pt(0)
        p_elem.space_after = Pt(8)
        p_elem.line_spacing = 1.2

        # Check for lead-in label like "Talent Team:"
        if p_text.startswith("Talent Team:"):
            lead_run = p_elem.add_run()
            lead_run.text = "Talent Team: "
            lead_run.font.name = "Proxima Nova"
            lead_run.font.size = Pt(chosen_sz)
            lead_run.font.bold = True
            lead_run.font.color.rgb = RGBColor(0x04, 0x11, 0x33)

            rest_text = p_text[len("Talent Team:"):].lstrip()
            rest_run = p_elem.add_run()
            rest_run.text = rest_text
            rest_run.font.name = "Proxima Nova"
            rest_run.font.size = Pt(chosen_sz)
            rest_run.font.bold = False
            rest_run.font.color.rgb = RGBColor(0x1D, 0x29, 0x39)
        else:
            run = p_elem.add_run()
            run.text = p_text
            run.font.name = "Proxima Nova"
            run.font.size = Pt(chosen_sz)
            run.font.bold = False
            run.font.color.rgb = RGBColor(0x1D, 0x29, 0x39)


def write_slide17_approach(slide, model: GovernanceDeckModel) -> None:
    """Populate Slide 17 Approach 3-card layout with uniform font fitting and Proxima Nova typography (HTL-72)."""
    # 1. Clear any placeholder shapes
    for s in list(slide.shapes):
        slide.shapes._spTree.remove(s._element)

    # 2. Add Kicker
    kicker_box = slide.shapes.add_textbox(457200, 384048, 6858300, 205800)
    tf_k = kicker_box.text_frame
    tf_k.word_wrap = True
    tf_k.margin_top = 0
    tf_k.margin_bottom = 0
    tf_k.margin_left = 0
    tf_k.margin_right = 0
    p_k = tf_k.paragraphs[0]
    p_k.text = model.slide17.kicker or "PROJECT SUCCESS AGREEMENT"
    for r in p_k.runs:
        r.font.name = "Proxima Nova"
        r.font.size = Pt(11)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0x04, 0x11, 0x33)

    # 3. Add Title
    title_box = slide.shapes.add_textbox(457200, 640080, 8023800, 332400)
    tf_t = title_box.text_frame
    tf_t.word_wrap = True
    tf_t.margin_top = 0
    tf_t.margin_bottom = 0
    tf_t.margin_left = 0
    tf_t.margin_right = 0
    p_t = tf_t.paragraphs[0]
    p_t.text = model.slide17.title or "Approach"
    for r in p_t.runs:
        r.font.name = "Proxima Nova"
        r.font.size = Pt(20)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0x04, 0x11, 0x33)

    # 4. Add Subtitle
    sub_box = slide.shapes.add_textbox(457200, 1042416, 8023800, 274200)
    tf_s = sub_box.text_frame
    tf_s.word_wrap = True
    tf_s.margin_top = 0
    tf_s.margin_bottom = 0
    tf_s.margin_left = 0
    tf_s.margin_right = 0
    p_s = tf_s.paragraphs[0]
    p_s.text = model.slide17.subtitle or "Execution phases, workstream deliverables, and governance model."
    for r in p_s.runs:
        r.font.name = "Proxima Nova"
        r.font.size = Pt(11)
        r.font.color.rgb = RGBColor(0x47, 0x54, 0x67)

    # 5. Add 3 Cards (Left, Middle, Right)
    # Card geometry matching Governance Slide 5/10 (Inches -> EMUs):
    # Width: 2.778 in (2,540,092 EMU), Height: 3.250 in (2,971,800 EMU), Top: 1.675 in (1,531,620 EMU)
    # Left positions: 0.668 in (610,870 EMU), 3.611 in (3,301,898 EMU), 6.554 in (5,992,926 EMU)
    card_w = 2540092
    card_h = 2971800
    top_pos = 1531620
    left_positions = [610870, 3301898, 5992926]

    cards = model.slide17.cards
    if not cards:
        return

    # Pass 1: Compute uniform fitting font size across all cards for bullets
    candidate_sizes = [10.5, 10.0, 9.5, 9.0, 8.5, 8.0]
    avail_w_in = 2.40
    avail_h_in = 2.05

    card_best_sizes = []
    for c in cards[:3]:
        c_sz = 8.0
        for sz in candidate_sizes:
            tot_h = 0.0
            for b in c.bullets:
                lines = wrap_lines(f"• {b}", width_in=avail_w_in, size_pt=sz)
                tot_h += lines * sz * 1.15 + 5.0
            if (tot_h / 72.0) <= avail_h_in:
                c_sz = sz
                break
        card_best_sizes.append(c_sz)

    uniform_sz = min(card_best_sizes) if card_best_sizes else 9.0

    # Pass 2: Render card background shapes and text boxes
    for idx, (c, left_pos) in enumerate(zip(cards[:3], left_positions)):
        # Card container rectangle
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos, top_pos, card_w, card_h)
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(0xF8, 0xF9, 0xFA)
        shape.line.color.rgb = RGBColor(0xEA, 0xEC, 0xF0)
        shape.line.width = Pt(1)

        # Inset text box
        # Inset: left + 0.18 in (164,592 EMU), top + 0.16 in (146,304 EMU), width - 0.36 in (2,210,908 EMU), height - 0.32 in (2,679,192 EMU)
        tb = slide.shapes.add_textbox(left_pos + 164592, top_pos + 146304, 2210908, 2679192)
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_top = 0
        tf.margin_bottom = 0
        tf.margin_left = 0
        tf.margin_right = 0
        for bpr in tb._element.xpath(".//a:bodyPr"):
            bpr.set("anchor", "t")

        # Paragraph 0: Category / Kicker
        p_cat = tf.paragraphs[0]
        p_cat.text = c.category
        p_cat.space_after = Pt(2)
        for r in p_cat.runs:
            r.font.name = "Proxima Nova"
            r.font.size = Pt(8.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0x2E, 0x5A, 0xAC)

        # Paragraph 1: Card Title
        p_title = tf.add_paragraph()
        p_title.text = c.title
        p_title.space_after = Pt(2)
        for r in p_title.runs:
            r.font.name = "Proxima Nova"
            r.font.size = Pt(13)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0x04, 0x11, 0x33)

        # Paragraph 2: Card Lead / Subtitle
        p_lead = tf.add_paragraph()
        p_lead.text = c.lead
        p_lead.space_after = Pt(8)
        for r in p_lead.runs:
            r.font.name = "Proxima Nova"
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(0x47, 0x54, 0x67)

        # Paragraphs 3+: Bullets
        for b in c.bullets:
            p_b = tf.add_paragraph()
            _set_paragraph_bullet_format(p_b, font_size_pt=uniform_sz, text=b)


def write_governance_deck(
    model: GovernanceDeckModel,
    target_path: Path,
    template_path: Optional[Path] = None,
    client_template_path: Optional[Path] = None,
) -> Path:
    """Load template presentation, assemble 19 slides, populate lifted and new slides, and save output file (HTL-72)."""
    src_template = resolve_governance_deck_template_path(template_path)
    if not src_template.exists():
        raise FileNotFoundError(f"Governance Deck template not found at {src_template}")

    src_client_template = resolve_client_deck_template_path(client_template_path)
    if not src_client_template.exists():
        raise FileNotFoundError(f"Client Deck template not found at {src_client_template}")

    prs_gov = pptx.Presentation(str(src_template))
    prs_client = pptx.Presentation(str(src_client_template))

    layout_default = prs_gov.slide_layouts[14]

    # Add Acceptance Criteria (Slide 15 - Phase 2 placeholder)
    s_acc_crit = prs_gov.slides.add_slide(layout_default)
    for s in list(s_acc_crit.shapes):
        s_acc_crit.shapes._spTree.remove(s._element)

    # Add How Acceptance Works (Slide 16 - Phase 2 placeholder)
    s_acc_works = prs_gov.slides.add_slide(layout_default)
    for s in list(s_acc_works.shapes):
        s_acc_works.shapes._spTree.remove(s._element)

    # Add Risk Register (Slide 18) - Copy shapes from Client Deck Slide 12
    s_risk = prs_gov.slides.add_slide(layout_default)
    for s in list(s_risk.shapes):
        s_risk.shapes._spTree.remove(s._element)
    for shape in prs_client.slides[11].shapes:
        s_risk.shapes._spTree.append(copy.deepcopy(shape._element))

    # Copy Slide 3 shapes (Team Roster placeholders) and images from Client Deck onto Gov Slide 12 (index 11)
    s_gov_12 = prs_gov.slides[11]
    for s in list(s_gov_12.shapes):
        s_gov_12.shapes._spTree.remove(s._element)

    # Drop existing image relationships from s_gov_12.part to avoid stale/wrong photo targets
    for rId in list(s_gov_12.part.rels.keys()):
        rel = s_gov_12.part.rels[rId]
        if "image" in rel.reltype:
            s_gov_12.part.drop_rel(rId)

    # Map client slide 3 image relationships into s_gov_12.part
    rId_map = {}
    import io
    for old_rId, rel in prs_client.slides[2].part.rels.items():
        if "image" in rel.reltype:
            client_img_part = rel.target_part
            image_stream = io.BytesIO(client_img_part.blob)
            target_img_part, new_rId = s_gov_12.part.get_or_add_image_part(image_stream)
            rId_map[old_rId] = new_rId

    # Copy shapes and update blip r:embed to faithful client template image references
    ns_r = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    for shape in prs_client.slides[2].shapes:
        elem = copy.deepcopy(shape._element)
        for blip in elem.xpath(".//a:blip"):
            old_embed = blip.get(f"{{{ns_r}}}embed")
            if old_embed and old_embed in rId_map:
                blip.set(f"{{{ns_r}}}embed", rId_map[old_embed])
        s_gov_12.shapes._spTree.append(elem)

    # Copy Slide 5 shapes (Why We're Here) from Client Deck onto Gov Slide 13 (index 12)
    s_gov_13 = prs_gov.slides[12]
    for s in list(s_gov_13.shapes):
        s_gov_13.shapes._spTree.remove(s._element)
    for s in prs_client.slides[4].shapes:
        s_gov_13.shapes._spTree.append(copy.deepcopy(s._element))

    # Copy Slide 6 shapes (What We Will Deliver) from Client Deck onto Gov Slide 14 (index 13)
    s_gov_14 = prs_gov.slides[13]
    for s in list(s_gov_14.shapes):
        s_gov_14.shapes._spTree.remove(s._element)
    for s in prs_client.slides[5].shapes:
        s_gov_14.shapes._spTree.append(copy.deepcopy(s._element))

    # Reorder slides into exact 19-slide sequence:
    # 0..10: Slides 1..11
    # 11: Slide 12 (Team Roster - index 11)
    # 12: Slide 13 (Why We're Here - index 12)
    # 13: Slide 14 (What We Deliver - index 13)
    # 14: Slide 15 (Acceptance Criteria - index 17)
    # 15: Slide 16 (How Acceptance Works - index 18)
    # 16: Slide 17 (Approach - index 14)
    # 17: Slide 18 (Initial Risk Register - index 19)
    # 18: Slide 19 (Next Steps - index 15)
    desired_indices = [0, 1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 17, 18, 14, 19, 15]
    sldIdLst = prs_gov.slides._sldIdLst
    all_sld_ids = list(sldIdLst)

    for child in list(sldIdLst):
        sldIdLst.remove(child)

    for idx in desired_indices:
        sldIdLst.append(all_sld_ids[idx])

    # Convert governance deck models to client models for slide writers
    client_deck_adapter = ClientDeckModel(
        project_name=model.project_name,
        slide1=ClientSlide1Model(client_name=model.slide1.client_name, project_name=model.slide1.project_name),
        slide5=ClientSlide5Model(
            why_are_we_doing_this=model.slide13.why_are_we_doing_this,
            what_problem_are_we_solving=model.slide13.what_problem_are_we_solving,
            what_does_success_look_like=model.slide13.what_does_success_look_like,
        ),
        slide6=ClientSlide6Model(
            workstreams=[
                ClientWorkstreamBlock(
                    title=ws.title,
                    deliverables=list(ws.deliverables),
                    overflow_count=ws.overflow_count,
                )
                for ws in model.slide14.workstreams
            ],
            solution_overview=model.slide14.solution_overview,
            out_of_scope=list(model.slide14.out_of_scope),
        ),
        slide12=ClientSlide12Model(
            risks=[
                ClientRiskItem(
                    id=r.id,
                    description=r.description,
                    prob_impact=r.prob_impact,
                    mitigation=r.mitigation,
                    target_date=r.target_date,
                    severity_score=r.severity_score,
                )
                for r in model.slide18.risks
            ]
        ),
        slide15=ClientSlide15Model(
            next_steps=[
                ClientNextStepItem(
                    owner=st.owner,
                    action=st.action,
                    due=st.due,
                )
                for st in model.slide19.next_steps
            ]
        ),
    )

    # Populate slides:
    # Slide 1 (index 0)
    write_slide1(prs_gov.slides[0], model)

    # Slide 13 (index 12): Why We're Here (Client Slide 5)
    write_slide5(prs_gov.slides[12], client_deck_adapter)

    # Slide 14 (index 13): What We Deliver (Client Slide 6)
    write_slide6(prs_gov.slides[13], client_deck_adapter)

    # Slide 15 (index 14): Acceptance Criteria Table (HTL-72)
    write_slide15_acceptance_criteria(prs_gov.slides[14], model)

    # Slide 16 (index 15): How Milestone Acceptance Works Narrative (HTL-72)
    write_slide16_acceptance_narrative(prs_gov.slides[15], model)

    # Slide 17 (index 16): Approach 3-Card Layout (HTL-72)
    write_slide17_approach(prs_gov.slides[16], model)

    # Slide 18 (index 17): Initial Risk Register (Client Slide 12)
    write_slide12(prs_gov.slides[17], client_deck_adapter)

    # Slide 19 (index 18): Next Steps (Client Slide 15)
    write_slide15(prs_gov.slides[18], client_deck_adapter)

    target_path.parent.mkdir(parents=True, exist_ok=True)
    prs_gov.save(str(target_path))

    # Write companion trace manifest: {project}_PS_Governance_Kickoff.trace.json
    manifest_path = target_path.with_name(f"{target_path.stem}.trace.json")
    model.manifest.write_json(manifest_path)

    return target_path

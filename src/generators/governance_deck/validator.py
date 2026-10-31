"""Validation and structural checks for PS Governance Kickoff Deck (HTL-72)."""

from pathlib import Path
from typing import List
import pptx

from src.generators.governance_deck.spec import GovernanceDeckModel


def check_governance_deck_traceability(
    model: GovernanceDeckModel,
    pptx_path: Path,
) -> List[str]:
    """Validate structural correctness, placeholder eradication, and traceability for Governance Deck (HTL-72)."""
    errors: List[str] = []

    if not pptx_path.exists():
        return [f"Presentation file does not exist: {pptx_path}"]

    try:
        prs = pptx.Presentation(str(pptx_path))
    except Exception as e:
        return [f"Failed to open presentation: {e}"]

    # 1. Slide Count
    if len(prs.slides) != 19:
        errors.append(f"Expected exactly 19 slides in Governance Deck, found {len(prs.slides)}")

    # 2. Slide 1 Validation
    if len(prs.slides) >= 1:
        s1 = prs.slides[0]
        s1_text = " ".join([s.text_frame.text for s in s1.shapes if s.has_text_frame])
        if "Governance Kickoff" not in s1_text:
            errors.append("Slide 1 title is missing 'Governance Kickoff'")
        if "Aligning on governance, outcomes and project plan." not in s1_text:
            errors.append("Slide 1 is missing subtitle 'Aligning on governance, outcomes and project plan.'")
        if "[Client Name]" in s1_text or "[Project Name]" in s1_text:
            errors.append("Slide 1 contains unreplaced [Client Name] or [Project Name] placeholders")

    # 3. Slide 12 (Team Roster) Validation
    if len(prs.slides) >= 12:
        s12 = prs.slides[11]
        s12_text = " ".join([s.text_frame.text for s in s12.shapes if s.has_text_frame])
        if "The people delivering this project" not in s12_text:
            errors.append("Slide 12 title is missing 'The people delivering this project'")
        if "The people delivering your project" in s12_text:
            errors.append("Slide 12 title still contains 'The people delivering your project'")

    # 4. Slide 13 (Why We're Here) Validation
    if len(prs.slides) >= 13:
        s13 = prs.slides[12]
        s13_text = " ".join([s.text_frame.text for s in s13.shapes if s.has_text_frame])
        if "[Pre-filled:" in s13_text:
            errors.append("Slide 13 contains unreplaced '[Pre-filled:' placeholders")

    # 4. Slide 14 (What We Deliver) Validation
    if len(prs.slides) >= 14:
        s14 = prs.slides[13]
        s14_text = " ".join([s.text_frame.text for s in s14.shapes if s.has_text_frame])
        if "[high level summary" in s14_text or "[Workstream or Feature" in s14_text:
            errors.append("Slide 14 contains unreplaced placeholder text")

    # 5. Slide 15 (Acceptance Criteria) Validation
    if len(prs.slides) >= 15:
        s15 = prs.slides[14]
        table_found = any(s.has_table for s in s15.shapes)
        if not table_found:
            errors.append("Slide 15 is missing Deliverable Acceptance Criteria table")
        s15_text = " ".join([s.text_frame.text for s in s15.shapes if s.has_text_frame])
        if "Deliverable acceptance criteria" not in s15_text:
            errors.append("Slide 15 is missing title 'Deliverable acceptance criteria'")

    # 6. Slide 16 (How Milestone Acceptance Works) Validation
    if len(prs.slides) >= 16:
        s16 = prs.slides[15]
        s16_text = " ".join([s.text_frame.text for s in s16.shapes if s.has_text_frame])
        if "How milestone acceptance works" not in s16_text:
            errors.append("Slide 16 is missing title 'How milestone acceptance works'")
        if "[Pre-filled:" in s16_text or "[CONFIRMATION REQUIRED]" in s16_text:
            errors.append("Slide 16 contains unreplaced placeholder text")

    # 7. Slide 17 (Approach) Validation
    if len(prs.slides) >= 17:
        s17 = prs.slides[16]
        s17_text = " ".join([s.text_frame.text for s in s17.shapes if s.has_text_frame])
        if "Approach" not in s17_text:
            errors.append("Slide 17 is missing title 'Approach'")
        if "[Pre-filled:" in s17_text or "[CONFIRMATION REQUIRED]" in s17_text:
            errors.append("Slide 17 contains unreplaced placeholder text")

    # 8. Slide 18 (Risk Register) Validation
    if len(prs.slides) >= 18:
        s18 = prs.slides[17]
        table_found = any(s.has_table for s in s18.shapes)
        if not table_found:
            errors.append("Slide 18 is missing Initial Risk Register table")

    # 9. Slide 19 (Next Steps) Validation
    if len(prs.slides) >= 19:
        s19 = prs.slides[18]
        table_found = any(s.has_table for s in s19.shapes)
        if not table_found:
            errors.append("Slide 19 is missing Next Steps table")

    # 10. Companion Trace Manifest Validation
    manifest_path = pptx_path.with_name(f"{pptx_path.stem}.trace.json")
    if not manifest_path.exists():
        errors.append(f"Trace manifest file missing at {manifest_path}")
    elif len(model.manifest.entries) == 0:
        errors.append("Trace manifest has 0 entries")

    return errors

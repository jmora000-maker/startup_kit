"""Dump a deck's slides, text frames, tables, notes, and slide-part count (review aid for DECK-17, Rev 10 report)."""

import sys
import zipfile
from pathlib import Path

import pptx
from pptx.util import Emu


def dump_deck(deck_path: Path) -> str:
    """Return a readable dump of every slide: layout, title, text frames, tables, and speaker notes."""
    out = []
    prs = pptx.Presentation(str(deck_path))
    for idx, slide in enumerate(prs.slides, start=1):
        title = slide.shapes.title.text if slide.shapes.title is not None else ""
        out.append(f"=== Slide {idx} | layout={slide.slide_layout.name} | title={title!r}")
        for sh in slide.shapes:
            geom = f"({Emu(sh.left).inches:.2f},{Emu(sh.top).inches:.2f} {Emu(sh.width).inches:.2f}x{Emu(sh.height).inches:.2f})"
            if sh.has_text_frame and sh.text_frame.text.strip():
                out.append(f"  [text] {sh.name} {geom}")
                for p in sh.text_frame.paragraphs:
                    out.append(f"      | {p.text}")
            if getattr(sh, "has_table", False) and sh.has_table:
                out.append(f"  [table] {sh.name} {geom}")
                for row in sh.table.rows:
                    out.append("      | " + " || ".join(c.text.replace("\n", " / ") for c in row.cells))
        notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
        out.append("  [notes]")
        for line in notes.splitlines():
            out.append(f"      | {line}")
    with zipfile.ZipFile(deck_path) as z:
        parts = [n for n in z.namelist() if n.startswith("ppt/slides/slide") and n.endswith(".xml")]
    out.append(f"slide parts in zip: {len(parts)}")
    return "\n".join(out)


def main() -> int:
    if len(sys.argv) != 2:
        print("usage: python -m src.tools.dump_deck <deck.pptx>")
        return 2
    sys.stdout.reconfigure(encoding="utf-8")
    print(dump_deck(Path(sys.argv[1])))
    return 0


if __name__ == "__main__":
    sys.exit(main())

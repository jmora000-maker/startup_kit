"""Deterministic parser extracting StartupKitBaseline from an existing Startup Kit Word (.docx) document."""

import re
import logging
from datetime import datetime, date
from pathlib import Path
from typing import List, Optional, Dict, Any, Tuple

import docx
from docx.table import Table, _Cell

from src.core.interfaces import IStartupKitDocxParser
from src.core.models import (
    StartupKitBaseline,
    ProjectStartupCharter,
    Deliverable,
    Milestone,
    RiskAssumption,
    DependencyAssumptionItem,
    DecisionItem,
    WorkPackageSeed,
    SOWStoryItem,
    SOWWorkItem,
    Stakeholder,
    RACIItem,
    CommunicationsPlanItem,
    CommercialGuardrail,
    TalentOnboardingRecord,
    TalentMember,
    ReadinessChecklistItem,
    GateDecision,
    GovernanceContext,
    SourceReference,
    SOWInterpretationSummary,
    ContractAmbiguityItem,
)
from src.config import extract_sow_references, detect_sow_reference_kind
from src.generators.pmo_workbook.mapping import strip_work_package_prefix
from src.scoring.readiness_engine import ReadinessScoringEngine

logger = logging.getLogger(__name__)


def clean_text(text: Optional[str]) -> str:
    """Strip [ACT-XX] badges and normalize whitespace from extracted table cell text."""
    if not text:
        return ""
    # 1. Strip badges ending with (+X.X% Recovery)] even if they contain nested brackets
    cleaned = re.sub(r'\[ACT(?:-REQ)?-[^\]:]+:\s*.*?\(\+\d+(?:\.\d+)?%\s*Recovery\)\]', '', text, flags=re.IGNORECASE)
    # 2. Strip standard action badges [ACT-...]
    cleaned = re.sub(r'\[ACT(?:-REQ)?-[^\]]+\]', '', cleaned, flags=re.IGNORECASE)
    # 3. Strip any trailing/orphaned recovery tags
    cleaned = re.sub(r'\(\+\d+(?:\.\d+)?%\s*Recovery\)?\]?', '', cleaned, flags=re.IGNORECASE)
    return cleaned.strip()


def parse_date_safely(date_str: Optional[str]) -> Optional[date]:
    """Parse date from string with multiple format attempts; returns None on failure/placeholder."""
    if not date_str:
        return None
    cleaned = clean_text(date_str)
    if cleaned.upper() in ("", "NONE", "N/A", "TBD", "TO BE DETERMINED", "[CONFIRMATION REQUIRED]", "UNASSIGNED"):
        return None

    # Strip day names or extra notes like "(Day 5)"
    cleaned = re.sub(r'\(.*?\)', '', cleaned).strip()

    formats = [
        "%Y-%m-%d",
        "%m/%d/%Y",
        "%d/%m/%Y",
        "%B %d, %Y",
        "%b %d, %Y",
        "%d-%b-%Y",
        "%d-%B-%Y",
        "%Y/%m/%d",
    ]
    for fmt in formats:
        try:
            return datetime.strptime(cleaned, fmt).date()
        except ValueError:
            continue

    # Try regex match for YYYY-MM-DD anywhere in string
    match = re.search(r'\b(\d{4})-(\d{2})-(\d{2})\b', cleaned)
    if match:
        try:
            return date(int(match.group(1)), int(match.group(2)), int(match.group(3)))
        except ValueError:
            pass

    return None


class StartupKitDocxParser(IStartupKitDocxParser):
    """Deterministic parser extracting StartupKitBaseline from an existing *_Startup_Kit.docx file."""

    def parse_startup_kit_docx(self, file_path: Path) -> StartupKitBaseline:
        """Parse an existing *_Startup_Kit.docx file into a StartupKitBaseline model."""
        if not file_path.exists():
            raise FileNotFoundError(f"Startup Kit Word document not found at: {file_path}")

        if file_path.suffix.lower() != ".docx":
            raise ValueError(f"Target file must be a .docx file: {file_path}")

        try:
            doc = docx.Document(str(file_path))
        except Exception as e:
            raise ValueError(f"Failed to open Word document {file_path}: {e}") from e

        # Validate basic Startup Kit document signature
        all_text = " ".join(p.text for p in doc.paragraphs)
        tables = doc.tables
        if not tables:
            raise ValueError(f"Ingested Word document does not match expected Startup Kit schema: No tables found in {file_path}")

        # Extract Project Name from subtitle or paragraphs
        project_name = self._extract_project_name(doc)

        # 1. Metadata Header Table
        meta_dict = self._parse_metadata_table(tables)
        if not meta_dict and "STARTUP KIT" not in all_text.upper():
            raise ValueError(f"Ingested Word document does not match expected Startup Kit schema: Missing Startup Kit metadata table in {file_path}")

        project_name = meta_dict.get("project_name") or project_name or "Project"
        client_name = meta_dict.get("client_name") or "[Client Sponsor]"
        governance_tier = meta_dict.get("governance_tier") or "Partnered"
        if governance_tier not in ("Guided", "Partnered", "Elevated"):
            governance_tier = "Partnered"
        contract_type = meta_dict.get("contract_type") or "Time and Materials"
        delivery_manager = meta_dict.get("delivery_manager") or "[UNASSIGNED - TO BE CONFIRMED]"
        talent_pm = meta_dict.get("talent_pm") or "[UNASSIGNED - TO BE CONFIRMED]"
        pmo_lead = meta_dict.get("pmo_lead") or "[UNASSIGNED - TO BE CONFIRMED]"
        sla_met = meta_dict.get("sla_met", True)
        workflow_state = meta_dict.get("workflow_state") or "Ready for G-01 Gate Review"

        src_ref = SourceReference(
            document_name=file_path.name,
            clause_or_slide="Section 4 Baseline",
            confidence_score=1.0
        )

        # 2. G-01 Checklist Table
        readiness_checklist = self._parse_g01_table(tables, pmo_lead, src_ref)
        c_doc = None
        if file_path.parent.exists():
            stem = file_path.stem
            candidate_names = [
                f"{stem.replace('_Startup_Kit', '')}_Startup_Readiness_Checklist.docx",
                f"{stem.replace('_Startup_Kit', '')}_Startup_Readiness_Checklist.doc",
                f"{stem}_Startup_Readiness_Checklist.docx",
                f"{stem}_Startup_Readiness_Checklist.doc",
                f"{project_name.replace(' ', '_')}_Startup_Readiness_Checklist.docx",
                f"{project_name.replace(' ', '_')}_Startup_Readiness_Checklist.doc",
                "Startup_Readiness_Checklist.docx",
                "Startup_Readiness_Checklist.doc",
            ]
            for c_name in candidate_names:
                c_path = file_path.parent / c_name
                if c_path.exists() and c_path != file_path:
                    try:
                        c_doc = docx.Document(str(c_path))
                        if not readiness_checklist:
                            readiness_checklist = self._parse_g01_table(c_doc.tables, pmo_lead, src_ref)
                        if c_doc:
                            break
                    except Exception as ex:
                        logger.debug("Could not parse companion checklist doc %s: %s", c_path, ex)
            if not c_doc:
                for c_path in sorted(file_path.parent.glob("*_Startup_Readiness_Checklist.doc*")):
                    if c_path != file_path:
                        try:
                            c_doc = docx.Document(str(c_path))
                            if not readiness_checklist:
                                readiness_checklist = self._parse_g01_table(c_doc.tables, pmo_lead, src_ref)
                            if c_doc:
                                break
                        except Exception as ex:
                            logger.debug("Could not parse companion checklist doc %s: %s", c_path, ex)

        if not readiness_checklist:
            readiness_checklist = self._build_default_checklist(pmo_lead, delivery_manager, talent_pm, governance_tier, src_ref)

        exceptions_count = len([i for i in readiness_checklist if i.exception_required or i.status == "Exception Required"])

        all_tables = list(tables) + (list(c_doc.tables) if c_doc else [])

        # 3. Open Questions & Clarifications
        open_questions = self._parse_open_questions(doc, all_tables, exceptions_count=exceptions_count, c_doc=c_doc)

        # 4. Project Startup Charter
        charter = self._parse_charter(doc, tables, project_name, client_name, governance_tier, contract_type, delivery_manager, talent_pm, pmo_lead, src_ref)

        # 5. SOW Interpretation Summary & Contract Ambiguities
        sow_interpretation, contract_ambiguities = self._parse_sow_interpretation(all_tables, src_ref)

        # 6. Milestone Delivery Plan
        milestones = self._parse_milestones_table(tables, delivery_manager, src_ref)

        # 6b. Interim Checkpoints Table (KIT-02)
        interim_checkpoints = self._parse_checkpoints_table(tables, src_ref)

        # 7. Scope Decomposition / Backlog Seed
        backlog_seed = self._parse_backlog_table(tables)

        # 8. Deliverables & Acceptance Matrix
        deliverables = self._parse_deliverables_table(tables, src_ref)

        # 9. RAID Items & Decisions
        raid_items, decisions = self._parse_raid_and_decisions_tables(tables, src_ref)

        # 9b. Dependency and Assumption Log Table
        dependencies_assumptions = self._parse_dependencies_assumptions_table(tables, src_ref)
        if not dependencies_assumptions:
            dependencies_assumptions = [
                DependencyAssumptionItem(
                    id=getattr(r, "id", None) or f"DA-{idx+1:02d}",
                    type="Dependency" if r.type == "Dependency" else "Assumption",
                    description=r.description,
                    source_reference=src_ref,
                    owner=r.owner,
                    status=r.status
                )
                for idx, r in enumerate(raid_items) if r.type in ("Assumption", "Dependency")
            ]

        # 10. Communications Plan
        communications_plan = self._parse_communications_table(tables)

        # 11. Stakeholders & RACI Matrix
        stakeholders, raci_matrix = self._parse_stakeholders_and_raci(tables)

        # 12. Commercial Guardrails
        commercial_guardrails = self._parse_commercial_guardrails(doc, tables, contract_type, src_ref, c_doc=c_doc)

        # 13. Talent Onboarding Record & Delivery Roster
        talent_onboarding = self._parse_talent_onboarding(doc, tables, pmo_lead, delivery_manager, talent_pm, src_ref)

        def _is_assigned(val_str: Optional[str]) -> bool:
            if not val_str:
                return False
            v = val_str.strip().upper()
            return not ("UNASSIGNED" in v or "CONFIRMATION" in v or v in ("", "NONE", "N/A", "TBD", "[TBD]"))

        # Check roster for assigned leadership names
        roster_dm = None
        roster_tpm = None
        roster_pmo = None
        if talent_onboarding.delivery_talent_roster:
            for tm in talent_onboarding.delivery_talent_roster:
                if any(w in tm.role.lower() for w in ("delivery manager", "delivery lead", "dm")):
                    if _is_assigned(tm.name):
                        roster_dm = tm.name
                elif any(w in tm.role.lower() for w in ("talent pm", "project manager", "tpm")):
                    if _is_assigned(tm.name):
                        roster_tpm = tm.name
                elif "pmo lead" in tm.role.lower():
                    if _is_assigned(tm.name):
                        roster_pmo = tm.name

        # Check stakeholders for assigned leadership names
        stk_dm = None
        stk_tpm = None
        stk_pmo = None
        if stakeholders:
            for stk in stakeholders:
                if any(w in stk.role.lower() for w in ("delivery manager", "delivery lead", "dm")):
                    if _is_assigned(stk.name):
                        stk_dm = stk.name
                elif any(w in stk.role.lower() for w in ("talent pm", "project manager", "tpm")):
                    if _is_assigned(stk.name):
                        stk_tpm = stk.name
                elif "pmo lead" in stk.role.lower():
                    if _is_assigned(stk.name):
                        stk_pmo = stk.name

        # Synchronize resolved leadership across parsed structures
        resolved_pmo = (
            pmo_lead if _is_assigned(pmo_lead) else None
        ) or (
            talent_onboarding.pmo_lead if _is_assigned(talent_onboarding.pmo_lead) else None
        ) or (
            charter.pmo_lead if _is_assigned(charter.pmo_lead) else None
        ) or roster_pmo or pmo_lead

        resolved_dm = (
            delivery_manager if _is_assigned(delivery_manager) else None
        ) or (
            talent_onboarding.delivery_manager if _is_assigned(talent_onboarding.delivery_manager) else None
        ) or (
            charter.delivery_manager if _is_assigned(charter.delivery_manager) else None
        ) or roster_dm or delivery_manager

        resolved_tpm = (
            talent_pm if _is_assigned(talent_pm) else None
        ) or (
            talent_onboarding.talent_pm if _is_assigned(talent_onboarding.talent_pm) else None
        ) or (
            charter.talent_pm if _is_assigned(charter.talent_pm) else None
        ) or roster_tpm or talent_pm

        charter.pmo_lead = resolved_pmo
        charter.delivery_manager = resolved_dm
        charter.talent_pm = resolved_tpm

        talent_onboarding.pmo_lead = resolved_pmo
        talent_onboarding.delivery_manager = resolved_dm
        talent_onboarding.talent_pm = resolved_tpm

        if talent_onboarding.delivery_talent_roster:
            for tm in talent_onboarding.delivery_talent_roster:
                if any(w in tm.role.lower() for w in ("delivery manager", "delivery lead", "dm")):
                    if _is_assigned(resolved_dm):
                        if not _is_assigned(tm.name):
                            tm.name = resolved_dm
                        if not tm.status or tm.status.lower() in ("pending", "needs alignment", "unassigned", "staffing required"):
                            tm.status = "Confirmed"
                elif any(w in tm.role.lower() for w in ("talent pm", "project manager", "tpm")):
                    if _is_assigned(resolved_tpm):
                        if not _is_assigned(tm.name):
                            tm.name = resolved_tpm
                        if not tm.status or tm.status.lower() in ("pending", "needs alignment", "unassigned", "staffing required"):
                            tm.status = "Confirmed"
                elif _is_assigned(tm.name) and (not tm.status or tm.status.lower() in ("pending", "needs alignment", "unassigned", "staffing required")):
                    tm.status = "Confirmed"

        # Build Governance Context
        gov_context = GovernanceContext(
            project_name=project_name,
            governance_tier=governance_tier,
            contract_type=contract_type,
            client_name=client_name,
            delivery_manager=resolved_dm,
            talent_pm=resolved_tpm,
            pmo_lead=resolved_pmo,
            executive_summary=charter.project_purpose,
            workflow_state=workflow_state,
            sla_met=sla_met
        )

        # Build baseline (scores and gate decision will be recalculated by aggregator)
        gate_decision = GateDecision(
            decision_status="Pending",
            target_mobilize_date=date.today(),
            tier_description=governance_tier,
            readiness_score=0.0,
            workflow_state=workflow_state,
            sla_met=sla_met,
            author_name=resolved_pmo
        )

        # Extract SOW story / work item catalogue (v5 B13, v6 Section 1.1)
        sow_stories_catalogue: List[SOWWorkItem] = []
        seen_catalogue_stories: Set[str] = set()
        for d in deliverables:
            text_block = f"{d.name or ''} {d.description or ''} {d.sow_reference or ''}"
            for s_id in extract_sow_references(text_block):
                if s_id not in seen_catalogue_stories:
                    seen_catalogue_stories.add(s_id)
                    sow_stories_catalogue.append(SOWWorkItem(
                        reference=s_id,
                        reference_kind=detect_sow_reference_kind(s_id) or "Story ID",
                        title="",
                        phase="",
                        owner="Toptal",
                        type="Build",
                        deliverable_id=d.id,
                        source_reference=d.source_reference
                    ))
        for wp in backlog_seed:
            text_block = f"{wp.title or ''} {wp.description or ''} {wp.sow_reference or ''}"
            wp_clean_title = strip_work_package_prefix(wp.title)
            extracted_title = wp_clean_title
            if ":" in wp_clean_title:
                parts = wp_clean_title.split(":", 1)
                if re.match(r'^[A-Za-z0-9_\-]+$', parts[0].strip()):
                    extracted_title = parts[1].strip()
            for s_id in extract_sow_references(text_block):
                if s_id not in seen_catalogue_stories:
                    seen_catalogue_stories.add(s_id)
                    sow_stories_catalogue.append(SOWWorkItem(
                        reference=s_id,
                        reference_kind=detect_sow_reference_kind(s_id) or "Story ID",
                        title=extracted_title,
                        phase="",
                        owner=wp.owner or "Toptal",
                        type="Build",
                        deliverable_id=wp.parent_deliverable_id,
                        source_reference=getattr(wp, "source_reference", None)
                    ))

        baseline = StartupKitBaseline(
            project_name=project_name,
            governance_tier=governance_tier,
            contract_type=contract_type,
            governance_context=gov_context,
            charter=charter,
            sow_interpretation=sow_interpretation,
            deliverables=deliverables,
            milestones=milestones,
            interim_checkpoints=interim_checkpoints,
            backlog_seed=backlog_seed,
            dependencies_assumptions=dependencies_assumptions,
            raid_items=raid_items,
            decisions=decisions,
            communications_plan=communications_plan,
            stakeholders=stakeholders,
            raci_matrix=raci_matrix,
            commercial_guardrails=commercial_guardrails,
            talent_onboarding=talent_onboarding,
            readiness_checklist=readiness_checklist,
            gate_decision=gate_decision,
            open_questions=open_questions,
            contract_ambiguities=contract_ambiguities,
            sow_stories_catalogue=sow_stories_catalogue,
            readiness_score=0.0,
            workflow_state=workflow_state,
            sow_awarded_date=date.today(),
            kit_drafted_date=date.today(),
            sla_met=sla_met,
            author_name=resolved_pmo
        )
        ReadinessScoringEngine.synchronize_checklist_with_artifacts(baseline)
        return baseline

    # -------------------------------------------------------------------------
    # Helper Extraction Methods
    # -------------------------------------------------------------------------

    def _extract_project_name(self, doc: docx.Document) -> str:
        """Extract project name from subtitle run or paragraphs."""
        for p in doc.paragraphs:
            text = p.text.strip()
            if text.startswith("Project Baseline:"):
                return text.replace("Project Baseline:", "").strip()
        return "Project"

    def _parse_metadata_table(self, tables: List[Table]) -> Dict[str, Any]:
        """Extract key metadata properties from the top header table."""
        res: Dict[str, Any] = {}
        for tbl in tables:
            # Metadata table is typically 4 columns and contains 'Project Name' or 'Governance Tier'
            full_tbl_text = " ".join(c.text for row in tbl.rows for c in row.cells).lower()
            if "governance tier" in full_tbl_text and "project name" in full_tbl_text:
                for row in tbl.rows:
                    cells = [c.text.strip() for c in row.cells]
                    if len(cells) >= 4:
                        # Pairs: (cells[0], cells[1]), (cells[2], cells[3])
                        k1, v1 = cells[0].lower(), cells[1]
                        k2, v2 = cells[2].lower(), cells[3]
                        self._map_meta_field(res, k1, v1)
                        self._map_meta_field(res, k2, v2)
                    elif len(cells) >= 2:
                        k1, v1 = cells[0].lower(), cells[1]
                        self._map_meta_field(res, k1, v1)
                break
        return res

    def _map_meta_field(self, res: Dict[str, Any], label: str, value: str) -> None:
        """Map label string to dictionary key."""
        label_clean = label.replace(":", "").strip()
        if "project name" in label_clean:
            res["project_name"] = value
        elif "client sponsor" in label_clean or "client" in label_clean:
            res["client_name"] = value
        elif "governance tier" in label_clean:
            res["governance_tier"] = value
        elif "contract type" in label_clean:
            res["contract_type"] = value
        elif "delivery manager" in label_clean or "delivery lead" in label_clean:
            res["delivery_manager"] = value
        elif "talent pm" in label_clean:
            res["talent_pm"] = value
        elif "pmo" in label_clean:
            res["pmo_lead"] = value
        elif "1-day sla status" in label_clean or "sla status" in label_clean:
            res["sla_met"] = not ("breached" in value.lower())
        elif "workflow state" in label_clean:
            res["workflow_state"] = value

    def _parse_open_questions(
        self,
        doc: docx.Document,
        tables: Optional[List[Table]] = None,
        exceptions_count: int = 0,
        c_doc: Optional[docx.Document] = None
    ) -> List[str]:
        """Extract unresolved open questions from callout boxes, paragraphs, Actionable Questions table, Action Required table, or in-cell badges."""
        questions: List[str] = []
        in_questions_section = False

        all_paragraphs = list(doc.paragraphs) + (list(c_doc.paragraphs) if c_doc else [])

        for p in all_paragraphs:
            text = p.text.strip()
            if not text:
                continue

            if "Unresolved Validation Points" in text or "Clarifications Pending Mobilize" in text or "Open Questions & Clarifications" in text or "Actionable Questions" in text:
                in_questions_section = True
                continue

            # Section headings end questions collection
            if in_questions_section and (text.startswith("Layer ") or text.startswith("Executive Readiness Gateway") or text.startswith("1.") or text.startswith("2.") or "Commercial and Margin Guardrails" in text or "Contract Ambiguities" in text):
                in_questions_section = False
                continue

            if in_questions_section:
                # Split lines
                for line in text.split("\n"):
                    l_clean = line.strip()
                    if (
                        l_clean.startswith("Total Score Recovery Potential:")
                        or l_clean.startswith("Resolving the ")
                        or l_clean.startswith("Action Required:")
                        or l_clean.startswith("• Corrective")
                        or l_clean.startswith("• Action:")
                    ):
                        continue
                    if l_clean.startswith("•") or l_clean.startswith("-") or l_clean.startswith("*"):
                        q_text = re.sub(r'^[•\-\*]\s*', '', l_clean).strip()
                        q_clean = clean_text(q_text)
                        if q_clean and "[RESOLVED]" not in q_clean.upper() and "No open clarification questions" not in q_clean:
                            questions.append(q_clean)
                    elif l_clean and not l_clean.startswith("Action Required:"):
                        q_clean = clean_text(l_clean)
                        if q_clean and "[RESOLVED]" not in q_clean.upper() and "No open clarification questions" not in q_clean:
                            questions.append(q_clean)

        # Check Actionable Questions table if present
        if tables:
            q_tbl = self._find_table_by_header(tables, ["question id", "actionable"])
            if not q_tbl:
                q_tbl = self._find_table_by_header(tables, ["question id", "clarification question"])
            if not q_tbl:
                q_tbl = self._find_table_by_header(tables, ["question id", "question"])
            if q_tbl:
                for row in q_tbl.rows[1:]:
                    cells = [clean_text(c.text) for c in row.cells]
                    if len(cells) >= 2:
                        q_text = cells[1]
                        if q_text and "RESOLVED" not in q_text.upper() and "No open clarification questions" not in q_text:
                            if q_text not in questions:
                                questions.append(q_text)

        # Also check Action Required table if present (legacy support)
        if tables:
            act_tbl = self._find_table_by_header(tables, ["action id", "type", "gate id"])
            if act_tbl:
                for row in act_tbl.rows[1:]:
                    cells = [clean_text(c.text) for c in row.cells]
                    if len(cells) >= 5 and cells[1] == "Open Clarification":
                        finding = cells[4].split("\n• Action:")[0].replace("• Action:", "").strip()
                        if finding and "RESOLVED" not in finding.upper() and finding not in questions:
                            questions.append(finding)

        # Also check in-table cell action badges if questions were not in a separate table
        if tables and not questions:
            for tbl in tables:
                # Exclude header/metadata and checklist tables
                if any(c.text.strip().lower() in ("gate id", "project baseline:", "governance tier") for row in tbl.rows[:1] for c in row.cells):
                    continue
                for row in tbl.rows[1:]:
                    for cell in row.cells:
                        # 1. Badges with recovery tag (handles nested brackets)
                        for m in re.finditer(r'\[ACT-(\d+):\s*(.*?)\s*\(\+\d+(?:\.\d+)?%\s*Recovery\)\]', cell.text, flags=re.IGNORECASE):
                            act_num = int(m.group(1))
                            desc = m.group(2).strip()
                            if act_num > exceptions_count:
                                if ":" in desc:
                                    prefix, rest = desc.split(":", 1)
                                    if any(w in prefix.lower() for w in ("kickoff", "clarify", "confirm", "review", "sponsor", "prerequisites", "validation", "mobilization", "align")):
                                        desc = rest.strip()
                                if desc and "RESOLVED" not in desc.upper() and desc not in questions:
                                    questions.append(desc)

                        # Strip the matched recovery badges before looking for simple badges
                        remaining = re.sub(r'\[ACT(?:-REQ)?-[^\]:]+:\s*.*?\(\+\d+(?:\.\d+)?%\s*Recovery\)\]', '', cell.text, flags=re.IGNORECASE)

                        # 2. Badges without recovery tag
                        for m in re.finditer(r'\[ACT-(\d+):\s*([^\]]+)\]', remaining, flags=re.IGNORECASE):
                            act_num = int(m.group(1))
                            desc = m.group(2).strip()
                            if act_num > exceptions_count:
                                if ":" in desc:
                                    prefix, rest = desc.split(":", 1)
                                    if any(w in prefix.lower() for w in ("kickoff", "clarify", "confirm", "review", "sponsor", "prerequisites", "validation", "mobilization", "align")):
                                        desc = rest.strip()
                                if desc and "RESOLVED" not in desc.upper() and desc not in questions:
                                    questions.append(desc)

        return questions

    def _find_table_by_header(self, tables: List[Table], header_keywords: List[str]) -> Optional[Table]:
        """Locate table whose header row matches given keyword criteria."""
        for tbl in tables:
            if not tbl.rows:
                continue
            header_text = " ".join(c.text.lower() for c in tbl.rows[0].cells)
            if all(kw.lower() in header_text for kw in header_keywords):
                return tbl
        return None

    def _parse_g01_table(self, tables: List[Table], pmo_lead: str, src_ref: SourceReference) -> List[ReadinessChecklistItem]:
        """Parse G-01 Checklist table from Executive Gateway section."""
        items: List[ReadinessChecklistItem] = []
        tbl = self._find_table_by_header(tables, ["gate id", "gate criterion", "status"])
        if not tbl:
            # Try alternate header: "item id", "criterion"
            tbl = self._find_table_by_header(tables, ["criterion", "status", "owner"])

        if not tbl:
            return items

        # Column indices
        header_cells = [c.text.strip().lower() for c in tbl.rows[0].cells]
        id_idx = next((i for i, h in enumerate(header_cells) if "id" in h), 0)
        crit_idx = next((i for i, h in enumerate(header_cells) if "criterion" in h), 1)
        art_idx = next((i for i, h in enumerate(header_cells) if "artifact" in h), 2)
        status_idx = next((i for i, h in enumerate(header_cells) if "status" in h), 3)
        owner_idx = next((i for i, h in enumerate(header_cells) if "owner" in h), 4)
        rev_idx = next((i for i, h in enumerate(header_cells) if "reviewer" in h), 5)
        app_idx = next((i for i, h in enumerate(header_cells) if "approver" in h), 6)
        ev_idx = next((i for i, h in enumerate(header_cells) if "evidence" in h or "exception" in h), 7)

        for row in tbl.rows[1:]:
            cells = [clean_text(c.text) for c in row.cells]
            if len(cells) < 4:
                continue
            item_id = cells[id_idx] if id_idx < len(cells) else f"G01-{len(items)+1:02d}"
            criterion = cells[crit_idx] if crit_idx < len(cells) else ""
            artifact = cells[art_idx] if art_idx < len(cells) else ""
            status = cells[status_idx] if status_idx < len(cells) else "Review Required"
            owner = cells[owner_idx] if owner_idx < len(cells) else pmo_lead
            reviewer = cells[rev_idx] if rev_idx < len(cells) else "Delivery Manager"
            approver = cells[app_idx] if app_idx < len(cells) else pmo_lead
            ev_text = cells[ev_idx] if ev_idx < len(cells) else ""

            # Extract exception info if embedded
            if status in ("Complete", "Approved"):
                exception_required = False
                exception_details = None
            else:
                exception_required = (status == "Exception Required")
                exception_details = None
                if "[EXCEPTION:" in ev_text:
                    m = re.search(r'\[EXCEPTION:\s*(.*?)\]', ev_text)
                    if m:
                        exception_details = m.group(1).strip()
                        exception_required = True
                elif exception_required:
                    exception_details = ev_text

            items.append(ReadinessChecklistItem(
                item_id=item_id,
                gate_criterion=criterion,
                related_section4_artifact=artifact,
                owner=owner,
                reviewer=reviewer,
                approver=approver,
                status=status,
                evidence=ev_text,
                exception_required=exception_required,
                exception_details=exception_details,
                approval_status="Approved" if status in ("Complete", "Approved") else ("Exception Required" if exception_required else "Pending Review")
            ))

        return items

    def _build_default_checklist(
        self,
        pmo_lead: str,
        delivery_manager: str,
        talent_pm: str,
        governance_tier: str,
        src_ref: SourceReference
    ) -> List[ReadinessChecklistItem]:
        """Build standard 15 G-01 checklist items when table is stored externally or not present."""
        checklist_defs = [
            ("G01-01", "Startup Kit created within 1 day of delivery handoff (SLA)", "Project Startup Charter", pmo_lead),
            ("G01-02", "Governance tier assigned and cadence established", "Project Startup Charter", pmo_lead),
            ("G01-03", "Deliverables mapped to owners with explicit acceptance routes", "Deliverables and Acceptance Matrix", talent_pm),
            ("G01-04", "Milestone delivery plan with external dates and internal buffers committed", "Milestone Delivery Plan", delivery_manager),
            ("G01-05", "Risks, Assumptions, Issues, and Dependencies seeded and owned", "RAID Log", talent_pm),
            ("G01-06", "Talent PM and DM briefed on governance cadence and KO decks", "Talent Onboarding Record", pmo_lead),
            ("G01-07", "Client onboarding dependencies and prerequisites mapped", "SOW Interpretation Summary", delivery_manager),
            ("G01-08", "Talent roster staffed with named leads and delivery talent", "Talent Onboarding Record", talent_pm),
            ("G01-09", "Client sponsor and decision escalation path identified", "Stakeholder and Responsibility Model", delivery_manager),
            ("G01-10", "RACI and PMO decision rights matrix baselined", "RACI / Decision Rights Matrix", pmo_lead),
            ("G01-11", "Communications and weekly reporting distribution established", "Communications and Reporting Plan", pmo_lead),
            ("G01-12", "Scope boundaries, assumptions, and non-approved work rules defined", "Scope Decomposition / Backlog Seed", delivery_manager),
            ("G01-13", "Change control procedure and threshold rules baselined", "Project Startup Charter", pmo_lead),
            ("G01-14", "Scope clarity confirmed and open contractual ambiguities resolved", "SOW Interpretation Summary", delivery_manager),
            ("G01-15", "Readiness Checklist verified and Approved for Mobilize", "Startup Readiness Checklist", pmo_lead),
        ]
        items = []
        for item_id, crit, art, owner in checklist_defs:
            items.append(ReadinessChecklistItem(
                item_id=item_id,
                gate_criterion=crit,
                related_section4_artifact=art,
                owner=owner,
                reviewer=delivery_manager if owner != delivery_manager else pmo_lead,
                approver=pmo_lead,
                status="Complete",
                evidence="Section 4 baseline artifact",
                exception_required=False,
                exception_details=None,
                approval_status="Approved"
            ))
        return items

    def _parse_charter(
        self,
        doc: docx.Document,
        tables: List[Table],
        project_name: str,
        client_name: str,
        tier: str,
        contract_type: str,
        dm: str,
        tpm: str,
        pmo: str,
        src_ref: SourceReference
    ) -> ProjectStartupCharter:
        """Parse Charter details from paragraphs, callout boxes, and list items."""
        project_purpose = f"Delivery mobilization and execution of {project_name}."
        delivery_model = f"Toptal {tier} Managed Delivery"
        governance_model = f"PMO {tier} Governance Framework"
        escalation_path = f"Talent PM ({tpm}) -> Delivery Manager ({dm}) -> PMO Lead ({pmo}) -> Director, PMO"
        unresolved_assumptions = "Logged in Startup RAID log"
        delivery_objectives: List[str] = []
        success_criteria: List[str] = []
        exclusions: List[str] = []

        in_charter = False
        in_objectives = False
        in_exclusions = False

        for p in doc.paragraphs:
            t = p.text.strip()
            if not t:
                continue

            if "Project Startup Charter" in t or "Startup Charter & Delivery Purpose" in t:
                in_charter = True
                continue
            elif "SOW Interpretation Summary" in t or "Layer 2" in t:
                in_charter = False
                in_objectives = False
                in_exclusions = False
                continue

            if in_charter:
                if "Project Purpose:" in t:
                    m = re.search(r'Project Purpose:\s*(.*?)(?:\n|$)', t)
                    if m and m.group(1).strip():
                        project_purpose = m.group(1).strip()
                if "Delivery Model:" in t:
                    m = re.search(r'Delivery Model:\s*(.*?)(?:\s*\|\s*Governance Model:|$)', t)
                    if m and m.group(1).strip():
                        delivery_model = m.group(1).strip()
                if "Governance Model:" in t:
                    m = re.search(r'Governance Model:\s*(.*?)(?:\n|$)', t)
                    if m and m.group(1).strip():
                        governance_model = m.group(1).strip()
                if "Escalation Path:" in t:
                    m = re.search(r'Escalation Path:\s*(.*?)(?:\n|$)', t)
                    if m and m.group(1).strip():
                        escalation_path = m.group(1).strip()
                if "Unresolved Assumptions Status:" in t:
                    m = re.search(r'Unresolved Assumptions Status:\s*(.*?)(?:\n|$)', t)
                    if m and m.group(1).strip():
                        unresolved_assumptions = m.group(1).strip()

                if "Delivery Objectives & Success Criteria:" in t:
                    in_objectives = True
                    in_exclusions = False
                    continue
                elif "High-Level Scope Exclusions:" in t:
                    in_exclusions = True
                    in_objectives = False
                    continue

                if in_objectives:
                    clean_line = re.sub(r'^[•\-\*]\s*', '', t).strip()
                    if clean_line:
                        if clean_line.lower().startswith("success criteria:"):
                            success_criteria.append(re.sub(r'^success criteria:\s*', '', clean_line, flags=re.I).strip())
                        else:
                            delivery_objectives.append(clean_line)
                elif in_exclusions:
                    clean_line = re.sub(r'^[•\-\*]\s*', '', t).strip()
                    if clean_line:
                        exclusions.append(clean_line)

        # Check for structured Charter table
        ch_tbl = self._find_table_by_header(tables, ["startup charter dimension", "charter commitment"])
        if ch_tbl:
            for row in ch_tbl.rows[1:]:
                cells = [clean_text(c.text) for c in row.cells]
                if len(cells) >= 2:
                    dim, val = cells[0].lower(), cells[1]
                    if "project purpose" in dim and val:
                        project_purpose = val
                    elif "delivery model" in dim:
                        m_del = re.search(r'Delivery:\s*(.*?)(?:\s*\|\s*Governance:|$)', val)
                        if m_del and m_del.group(1).strip():
                            delivery_model = m_del.group(1).strip()
                        m_gov = re.search(r'Governance:\s*(.*?)(?:$)', val)
                        if m_gov and m_gov.group(1).strip():
                            governance_model = m_gov.group(1).strip()
                    elif "escalation path" in dim and val:
                        escalation_path = val
                    elif "assumptions" in dim and val:
                        unresolved_assumptions = val

        return ProjectStartupCharter(
            project_name=project_name,
            client_name=client_name,
            governance_tier=tier,
            contract_type=contract_type,
            delivery_manager=dm,
            talent_pm=tpm,
            pmo_lead=pmo,
            project_purpose=project_purpose,
            delivery_model=delivery_model,
            governance_model=governance_model,
            escalation_path=escalation_path,
            unresolved_assumptions_status=unresolved_assumptions,
            delivery_objectives=delivery_objectives or [f"Deliver {project_name} within baseline commitments."],
            success_criteria=success_criteria or ["Formal client acceptance and sign-off."],
            high_level_scope=[f"Delivery scope defined in {project_name} SOW baseline."],
            exclusions=exclusions or ["Out-of-scope tasks defined in SOW."],
            source_reference=src_ref
        )

    def _parse_sow_interpretation(
        self,
        tables: List[Table],
        src_ref: SourceReference
    ) -> Tuple[SOWInterpretationSummary, List[ContractAmbiguityItem]]:
        """Parse SOW Interpretation summary table and contract ambiguities table."""
        contracted_delivs: List[str] = []
        exclusions: List[str] = []
        obligations: List[str] = []
        assumptions: List[str] = []
        constraints: List[str] = []
        platforms: List[str] = []
        dependencies: List[str] = []
        approval_expectations: Optional[str] = "Client sign-off required within 5 business days."
        ambiguity_notes: List[str] = []

        tbl = self._find_table_by_header(tables, ["sow interpretation dimension", "contractual summary"])
        if tbl:
            for row in tbl.rows[1:]:
                cells = [clean_text(c.text) for c in row.cells]
                if len(cells) < 2:
                    continue
                dim, val = cells[0].lower(), cells[1]
                items = [re.sub(r'^[•\-]\s*', '', line).strip() for line in val.split("\n") if line.strip()]

                if "deliverables" in dim:
                    contracted_delivs.extend(items)
                elif "exclusions" in dim or "out-of-scope" in dim:
                    exclusions.extend(items)
                elif "obligations" in dim or "prerequisites" in dim:
                    if not items or all(not it.strip() for it in items):
                        obligations.append("[CONFIRMATION REQUIRED]")
                    else:
                        obligations.extend([it for it in items if it.strip()])
                elif "assumptions" in dim or "constraints" in dim:
                    assumptions.extend(items)
                elif "platform" in dim or "environment" in dim:
                    platforms.extend(items)
                elif "dependencies" in dim:
                    dependencies.extend(items)
                elif "approval" in dim or "acceptance" in dim:
                    approval_expectations = val
                elif "ambiguities" in dim:
                    ambiguity_notes.extend(items)

        # Parse contract ambiguities table
        ambiguities: List[ContractAmbiguityItem] = []
        amb_tbl = self._find_table_by_header(tables, ["anomaly id", "category"])
        if not amb_tbl:
            amb_tbl = self._find_table_by_header(tables, ["ambiguity id", "clause / topic", "identified conflict"])
        if amb_tbl:
            for row in amb_tbl.rows[1:]:
                cells = [clean_text(c.text) for c in row.cells]
                if len(cells) >= 5:
                    is_resolved = (
                        "RESOLVED" in cells[4].upper()
                        or (len(cells) >= 6 and "RESOLVED" in cells[5].upper())
                        or "RESOLVED" in cells[0].upper()
                        or "RESOLVED" in cells[1].upper()
                        or "RESOLVED" in cells[2].upper()
                        or "RESOLVED" in cells[3].upper()
                    )
                    amb_status = "Resolved" if is_resolved else "Open"
                    ambiguities.append(ContractAmbiguityItem(
                        anomaly_id=cells[0],
                        category=cells[1],
                        conflicting_clauses=cells[2],
                        risk_impact=cells[3],
                        recommended_clarification=cells[4],
                        status=amb_status
                    ))

        summary = SOWInterpretationSummary(
            contracted_deliverables=contracted_delivs,
            out_of_scope_items=exclusions,
            customer_obligations=obligations,
            assumptions=assumptions,
            constraints=constraints,
            platform_environment_commitments=platforms,
            dependencies=dependencies,
            approval_expectations=approval_expectations,
            ambiguity_notes=ambiguity_notes,
            source_reference=src_ref
        )
        return summary, ambiguities

    def _parse_milestones_table(
        self,
        tables: List[Table],
        default_owner: str,
        src_ref: SourceReference
    ) -> List[Milestone]:
        """Parse Milestone Delivery Plan table."""
        milestones: List[Milestone] = []
        tbl = self._find_table_by_header(tables, ["milestone id", "description", "date"])
        if not tbl:
            tbl = self._find_table_by_header(tables, ["milestone", "date"])
        if not tbl:
            tbl = self._find_table_by_header(tables, ["milestone id", "description"])

        if not tbl:
            return milestones

        header_cells = [c.text.strip().lower() for c in tbl.rows[0].cells]
        id_idx = next((i for i, h in enumerate(header_cells) if "id" in h), 0)
        desc_idx = next((i for i, h in enumerate(header_cells) if "desc" in h or "name" in h or "milestone" in h), 1)
        ext_date_idx = next((i for i, h in enumerate(header_cells) if "target" in h or "external" in h), 2)
        buf_date_idx = next((i for i, h in enumerate(header_cells) if "buffer" in h or "internal" in h), 3)
        owner_idx = next((i for i, h in enumerate(header_cells) if "owner" in h), 4)
        dep_idx = next((i for i, h in enumerate(header_cells) if "dep" in h), 5)

        for row in tbl.rows[1:]:
            cells = [clean_text(c.text) for c in row.cells]
            if len(cells) < 2:
                continue
            m_id = cells[id_idx] if id_idx < len(cells) else f"M{len(milestones)+1:02d}"
            m_desc = cells[desc_idx] if desc_idx < len(cells) else ""
            ext_date_str = cells[ext_date_idx] if ext_date_idx < len(cells) else None
            buf_date_str = cells[buf_date_idx] if buf_date_idx < len(cells) else None
            owner = cells[owner_idx] if owner_idx < len(cells) else default_owner
            dep_str = cells[dep_idx] if dep_idx < len(cells) else ""

            milestones.append(Milestone(
                id=m_id,
                description=m_desc,
                external_date=parse_date_safely(ext_date_str),
                internal_buffer_date=parse_date_safely(buf_date_str),
                owner=owner or default_owner,
                key_dependencies=[d.strip() for d in dep_str.split(",") if d.strip()] if dep_str else [],
                source_reference=src_ref
            ))

        return milestones

    def _parse_checkpoints_table(
        self,
        tables: List[Table],
        src_ref: SourceReference
    ) -> List[Milestone]:
        """Parse Interim Checkpoints table (KIT-02)."""
        checkpoints: List[Milestone] = []
        tbl = self._find_table_by_header(tables, ["checkpoint id", "phase", "description"])
        if not tbl:
            tbl = self._find_table_by_header(tables, ["checkpoint", "description"])
        if not tbl:
            return checkpoints

        header_cells = [c.text.strip().lower() for c in tbl.rows[0].cells]
        id_idx = next((i for i, h in enumerate(header_cells) if "id" in h or "checkpoint" in h), 0)
        phase_idx = next((i for i, h in enumerate(header_cells) if "phase" in h), -1)
        desc_idx = next((i for i, h in enumerate(header_cells) if "desc" in h), 2)
        if desc_idx >= len(header_cells):
            desc_idx = 1

        for row in tbl.rows[1:]:
            cells = [clean_text(c.text) for c in row.cells]
            if len(cells) >= 2:
                cp_id = cells[id_idx] if id_idx < len(cells) else f"CP-{len(checkpoints)+1:02d}"
                cp_phase = cells[phase_idx] if (phase_idx != -1 and phase_idx < len(cells)) else None
                cp_desc = cells[desc_idx] if desc_idx < len(cells) else ""
                checkpoints.append(Milestone(
                    id=cp_id,
                    description=cp_desc,
                    phase=cp_phase,
                    external_date=None,
                    internal_buffer_date=None,
                    owner="Talent PM",
                    source_reference=src_ref
                ))
        return checkpoints

    def _parse_backlog_table(self, tables: List[Table]) -> List[WorkPackageSeed]:
        """Parse Scope Decomposition / Backlog Seed table."""
        backlog: List[WorkPackageSeed] = []
        tbl = self._find_table_by_header(tables, ["wp id", "parent deliv", "work package title"])
        if not tbl:
            tbl = self._find_table_by_header(tables, ["work package", "parent", "status"])

        if not tbl:
            return backlog

        header_cells = [c.text.strip().lower() for c in tbl.rows[0].cells] if tbl.rows else []
        id_idx = next((i for i, h in enumerate(header_cells) if "wp id" in h or "id" in h), 0)
        parent_idx = next((i for i, h in enumerate(header_cells) if "parent" in h), 1)
        title_idx = next((i for i, h in enumerate(header_cells) if "title" in h or "work package" in h or "desc" in h), 2)
        seq_idx = next((i for i, h in enumerate(header_cells) if "seq" in h), 3)
        owner_idx = next((i for i, h in enumerate(header_cells) if "owner" in h), 4)
        sow_idx = next((i for i, h in enumerate(header_cells) if "sow" in h or "ref" in h or "story" in h or "stories" in h), 5)
        status_idx = next((i for i, h in enumerate(header_cells) if "status" in h), 6)

        for row in tbl.rows[1:]:
            cells = [clean_text(c.text) for c in row.cells]
            if len(cells) >= 3:
                wp_id = cells[id_idx] if id_idx < len(cells) else f"WP-{len(backlog)+1:02d}"
                parent_id = cells[parent_idx] if parent_idx < len(cells) else ""
                title = cells[title_idx] if title_idx < len(cells) else ""
                seq_val = cells[seq_idx] if seq_idx < len(cells) else "1"
                seq = int(seq_val) if seq_val.isdigit() else 1
                owner = cells[owner_idx] if owner_idx < len(cells) else "[UNASSIGNED - TO BE CONFIRMED]"
                sow_ref = cells[sow_idx] if (sow_idx < len(cells) and sow_idx != -1) else None
                status = cells[status_idx] if (status_idx < len(cells) and status_idx != -1) else "Draft"

                backlog.append(WorkPackageSeed(
                    id=wp_id,
                    parent_deliverable_id=parent_id,
                    title=title,
                    description=title,
                    preliminary_sequence=seq,
                    owner=owner,
                    sow_reference=sow_ref,
                    status=status
                ))

        return backlog

    def _parse_deliverables_table(
        self,
        tables: List[Table],
        src_ref: SourceReference
    ) -> List[Deliverable]:
        """Parse Deliverables and Acceptance Matrix table."""
        deliverables: List[Deliverable] = []
        tbl = self._find_table_by_header(tables, ["deliverable name", "internal owner", "acceptance criteria"])
        if not tbl:
            tbl = self._find_table_by_header(tables, ["deliverable", "criteria", "approver"])

        if not tbl:
            return deliverables

        header_cells = [c.text.strip().lower() for c in tbl.rows[0].cells]
        id_idx = next((i for i, h in enumerate(header_cells) if "id" in h and "parent" not in h), 0)
        desc_idx = next((i for i, h in enumerate(header_cells) if "name" in h or "desc" in h or "deliverable" in h), 1)
        crit_idx = next((i for i, h in enumerate(header_cells) if "criteria" in h or "acceptance" in h), 2)
        ev_idx = next((i for i, h in enumerate(header_cells) if "evidence" in h), 3)
        app_idx = next((i for i, h in enumerate(header_cells) if "approver" in h or "client" in h), 4)
        owner_idx = next((i for i, h in enumerate(header_cells) if "owner" in h), 5)
        sow_idx = next((i for i, h in enumerate(header_cells) if "sow" in h or "ref" in h or "story" in h or "stories" in h), -1)
        signoff_idx = next((i for i, h in enumerate(header_cells) if "sign-off" in h or "mechanism" in h or "review" in h), -1)

        for row in tbl.rows[1:]:
            cells = [clean_text(c.text) for c in row.cells]
            if len(cells) < 2:
                continue
            d_id = cells[id_idx] if id_idx < len(cells) else f"DEL-{len(deliverables)+1:02d}"
            d_name = cells[desc_idx] if desc_idx < len(cells) else ""
            owner = cells[owner_idx] if (owner_idx != -1 and owner_idx < len(cells)) else "Unassigned"
            criteria_str = cells[crit_idx] if (crit_idx != -1 and crit_idx < len(cells)) else None
            ev_str = cells[ev_idx] if (ev_idx != -1 and ev_idx < len(cells)) else "Test sign-off sheet / automated pipeline run output"
            app_str = cells[app_idx] if (app_idx != -1 and app_idx < len(cells)) else "[UNASSIGNED - TO BE CONFIRMED]"
            sow_ref = cells[sow_idx] if (sow_idx != -1 and sow_idx < len(cells)) else None
            signoff_str = cells[signoff_idx] if (signoff_idx != -1 and signoff_idx < len(cells)) else (cells[-1] if len(cells) >= 7 else "Formal written sign-off")

            deliverables.append(Deliverable(
                id=d_id,
                name=d_name,
                description=d_name,
                source_reference=src_ref,
                owner=owner,
                acceptance_criteria=criteria_str,
                evidence_required=ev_str,
                client_approver=app_str,
                sow_reference=sow_ref,
                review_window=signoff_str,
                rejection_rework_path="Talent PM / Team rework within 3 business days of notice"
            ))

        return deliverables

    def _parse_raid_and_decisions_tables(
        self,
        tables: List[Table],
        src_ref: SourceReference
    ) -> Tuple[List[RiskAssumption], List[DecisionItem]]:
        """Parse RAID Log and Decision Log tables (v4 B5)."""
        raid_items: List[RiskAssumption] = []
        decisions: List[DecisionItem] = []

        # RAID table
        tbl = self._find_table_by_header(tables, ["type", "description", "mitigation"])
        if not tbl:
            tbl = self._find_table_by_header(tables, ["raid", "owner", "status"])

        if tbl:
            header_cells = [c.text.strip().lower() for c in tbl.rows[0].cells]
            id_idx = next((i for i, h in enumerate(header_cells) if "item id" in h or ("id" in h and "raid" not in h)), -1)
            type_idx = next((i for i, h in enumerate(header_cells) if "type" in h), 0 if id_idx != 0 else 1)
            desc_idx = next((i for i, h in enumerate(header_cells) if "desc" in h), 1 if id_idx == -1 else 2)
            cat_idx = next((i for i, h in enumerate(header_cells) if "cat" in h), -1)
            prob_idx = next((i for i, h in enumerate(header_cells) if "prob" in h), -1)
            imp_idx = next((i for i, h in enumerate(header_cells) if "impact" in h or "severity" in h), -1)
            owner_idx = next((i for i, h in enumerate(header_cells) if "owner" in h), -1)
            mit_idx = next((i for i, h in enumerate(header_cells) if "mitigation" in h or "strategy" in h or "response" in h), -1)
            status_idx = next((i for i, h in enumerate(header_cells) if "status" in h), -1)

            for row in tbl.rows[1:]:
                cells = [clean_text(c.text) for c in row.cells]
                if len(cells) < 3:
                    continue
                r_type = cells[type_idx] if (type_idx != -1 and type_idx < len(cells)) else "Risk"
                if r_type not in ("Risk", "Assumption", "Issue", "Dependency"):
                    r_type = "Risk"
                default_prefix = "RSK" if r_type == "Risk" else ("ISS" if r_type == "Issue" else "RAID")
                r_id = cells[id_idx] if (id_idx != -1 and id_idx < len(cells)) else f"{default_prefix}-{len(raid_items)+1:02d}"
                r_desc = cells[desc_idx] if (desc_idx != -1 and desc_idx < len(cells)) else ""
                r_cat = cells[cat_idx] if (cat_idx != -1 and cat_idx < len(cells)) else "Delivery Risk"
                r_prob = cells[prob_idx] if (prob_idx != -1 and prob_idx < len(cells)) else "Medium"
                r_imp = cells[imp_idx] if (imp_idx != -1 and imp_idx < len(cells)) else "Medium"
                r_owner = cells[owner_idx] if (owner_idx != -1 and owner_idx < len(cells)) else "Delivery Manager"
                r_mit = cells[mit_idx] if (mit_idx != -1 and mit_idx < len(cells)) else "Active monitoring"
                r_status = cells[status_idx] if (status_idx != -1 and status_idx < len(cells)) else "Open"

                raid_items.append(RiskAssumption(
                    id=r_id,
                    type=r_type,
                    category=r_cat,
                    description=r_desc,
                    owner=r_owner,
                    probability=r_prob,
                    impact=r_imp,
                    mitigation_or_response=r_mit,
                    status=r_status,
                    source_reference=src_ref
                ))

        # Decision log table
        dec_tbl = self._find_table_by_header(tables, ["decision id", "decision text", "status"])
        if dec_tbl:
            for row in dec_tbl.rows[1:]:
                cells = [clean_text(c.text) for c in row.cells]
                if len(cells) >= 4:
                    decisions.append(DecisionItem(
                        id=cells[0],
                        decision_text=cells[1],
                        decision_owner=cells[2],
                        status=cells[3]
                    ))

        return raid_items, decisions

    def _parse_dependencies_assumptions_table(
        self,
        tables: List[Table],
        src_ref: SourceReference
    ) -> List[DependencyAssumptionItem]:
        """Parse Dependency and Assumption Log table."""
        items: List[DependencyAssumptionItem] = []
        tbl = None
        for t in tables:
            if not t.rows:
                continue
            headers = [c.text.strip().lower() for c in t.rows[0].cells]
            if any("mitigation" in h for h in headers) or any("probability" in h for h in headers):
                continue
            if any("type" in h for h in headers) and any("desc" in h for h in headers) and any("owner" in h for h in headers) and (any("item id" in h for h in headers) or any("id" in h for h in headers) or any("category" in h for h in headers)):
                tbl = t
                break

        if not tbl:
            tbl = self._find_table_by_header(tables, ["dependency and assumption", "owner"])
        if tbl:
            header_cells = [c.text.strip().lower() for c in tbl.rows[0].cells]
            id_idx = next((i for i, h in enumerate(header_cells) if "id" in h), 0)
            type_idx = next((i for i, h in enumerate(header_cells) if "type" in h), 1)
            desc_idx = next((i for i, h in enumerate(header_cells) if "desc" in h), 2)
            cat_idx = next((i for i, h in enumerate(header_cells) if "category" in h or "cat" in h), 3)
            owner_idx = next((i for i, h in enumerate(header_cells) if "owner" in h), 4)
            status_idx = next((i for i, h in enumerate(header_cells) if "status" in h), 5)
            for row in tbl.rows[1:]:
                cells = [clean_text(c.text) for c in row.cells]
                if len(cells) < 3:
                    continue
                d_id = cells[id_idx] if id_idx < len(cells) else f"DA-{len(items)+1:02d}"
                raw_type = cells[type_idx] if type_idx < len(cells) else "Dependency"
                d_type = "Assumption" if "ASSUMPTION" in raw_type.upper() else "Dependency"
                d_desc = cells[desc_idx] if desc_idx < len(cells) else ""
                d_cat = cells[cat_idx] if (cat_idx != -1 and cat_idx < len(cells)) else "Technical"
                d_owner = cells[owner_idx] if (owner_idx != -1 and owner_idx < len(cells)) else "Unassigned"
                d_status = cells[status_idx] if (status_idx != -1 and status_idx < len(cells)) else "Open"
                items.append(DependencyAssumptionItem(
                    id=d_id,
                    type=d_type,
                    category=d_cat,
                    description=d_desc,
                    source_reference=src_ref,
                    owner=d_owner,
                    status=d_status
                ))
        return items

    def _parse_communications_table(self, tables: List[Table]) -> List[CommunicationsPlanItem]:
        """Parse Communications and Reporting Plan table (v4 B5)."""
        comms: List[CommunicationsPlanItem] = []
        tbl = self._find_table_by_header(tables, ["report / meeting", "audience", "cadence"])
        if not tbl:
            tbl = self._find_table_by_header(tables, ["communication", "audience", "cadence"])

        if tbl:
            header_cells = [c.text.strip().lower() for c in tbl.rows[0].cells]
            has_id = ("id" in header_cells[0])
            for idx, row in enumerate(tbl.rows[1:], 1):
                cells = [c.text.strip() for c in row.cells]
                if has_id and len(cells) >= 7:
                    c_id = cells[0]
                    c_name = cells[1]
                    c_aud = cells[2]
                    c_own = cells[3]
                    c_cad = cells[4]
                    c_fmt = cells[5]
                    c_day = cells[6]
                elif len(cells) >= 6:
                    c_id = f"COM-{idx:02d}"
                    c_name = cells[0]
                    c_aud = cells[1]
                    c_own = cells[2]
                    c_cad = cells[3]
                    c_fmt = cells[4]
                    c_day = cells[5]
                else:
                    continue

                comms.append(CommunicationsPlanItem(
                    id=c_id,
                    name=c_name,
                    audience=c_aud,
                    content_owner=c_own,
                    cadence=c_cad,
                    format=c_fmt,
                    delivery_day=c_day
                ))

        return comms

    def _parse_stakeholders_and_raci(
        self,
        tables: List[Table]
    ) -> Tuple[List[Stakeholder], List[RACIItem]]:
        """Parse Stakeholders table and RACI matrix table."""
        stakeholders: List[Stakeholder] = []
        raci: List[RACIItem] = []

        # Stakeholders
        tbl = self._find_table_by_header(tables, ["stakeholder name", "role", "decision rights"])
        if tbl:
            for row in tbl.rows[1:]:
                cells = [c.text.strip() for c in row.cells]
                if len(cells) >= 5:
                    stakeholders.append(Stakeholder(
                        name=cells[0],
                        role=cells[1],
                        organization=cells[2],
                        decision_rights=cells[3],
                        escalation_responsibility=cells[4]
                    ))

        # RACI Matrix
        raci_tbl = self._find_table_by_header(tables, ["startup control", "pmo lead", "delivery manager"])
        if not raci_tbl:
            raci_tbl = self._find_table_by_header(tables, ["decision or activity", "pmo", "delivery manager"])

        if raci_tbl:
            for row in raci_tbl.rows[1:]:
                cells = [c.text.strip() for c in row.cells]
                if len(cells) >= 6:
                    raci.append(RACIItem(
                        decision_or_activity=cells[0],
                        pmo_lead=cells[1],
                        delivery_manager=cells[2],
                        talent_pm=cells[3],
                        sales_accounts=cells[4],
                        client=cells[5]
                    ))

        return stakeholders, raci

    def _parse_commercial_guardrails(
        self,
        doc: docx.Document,
        tables: List[Table],
        contract_type: str,
        src_ref: SourceReference,
        c_doc: Optional[docx.Document] = None
    ) -> CommercialGuardrail:
        """Parse Commercial Guardrails callout box or table from main doc or companion checklist doc."""
        contract_implications = f"Managed delivery under {contract_type} governance rules."
        billing_consumption = "Weekly timesheet approval and hourly/daily burn rate tracking against budget cap." if contract_type != "Fixed Bid" else "Invoicing tied strictly to formal client milestone acceptance sign-offs."
        staffing_assump = "Dedicated talent staffing mapped to contracted roles."
        commercial_exposure = "Client dependency delays must be logged immediately to prevent unfunded team standby burn."
        approved_rule = "All delivery milestones and billable hours must map directly to contracted SOW deliverables."
        non_approved_rule = "Any activity outside agreed scope requires an approved Change Order."
        work_at_risk = "Work-at-risk strictly prohibited without written PMO Lead approval and executive exception sign-off."
        change_trigger = "Material scope modifications, milestone shifts > 5 business days, or client delay."
        change_route = "PMO Lead leads -> DM aligns client -> Client approves -> Contracting issues change order"
        budget_baseline = "Established from SOW financial schedule"
        variance_ind = "Green (<5% variance)"
        margin_risk = "Low"
        escalation_thresh = "Budget variance > 10% or milestone delay > 3 days"

        all_paragraphs = list(doc.paragraphs) + (list(c_doc.paragraphs) if c_doc else [])
        all_tables = list(tables) + (list(c_doc.tables) if c_doc else [])

        for p in all_paragraphs:
            t = p.text.strip()
            if "Commercial Guardrails & Margin Protection Rules" in t or "Commercial and Margin Guardrails" in t:
                if "Contract Implications:" in t:
                    m = re.search(r'Contract Implications:\s*(.*?)(?:\n|$)', t)
                    if m and m.group(1).strip():
                        contract_implications = m.group(1).strip()
                if "Approved Work Rule:" in t:
                    m = re.search(r'Approved Work Rule:\s*(.*?)(?:\n|$)', t)
                    if m and m.group(1).strip():
                        approved_rule = m.group(1).strip()
                if "Non-Approved Work Rule:" in t:
                    m = re.search(r'Non-Approved Work Rule:\s*(.*?)(?:\n|$)', t)
                    if m and m.group(1).strip():
                        non_approved_rule = m.group(1).strip()
                if "Work-at-Risk Policy:" in t:
                    m = re.search(r'Work-at-Risk Policy:\s*(.*?)(?:\n|$)', t)
                    if m and m.group(1).strip():
                        work_at_risk = m.group(1).strip()
                if "Change Control Triggers:" in t:
                    m = re.search(r'Change Control Triggers:\s*(.*?)(?:\n|$)', t)
                    if m and m.group(1).strip():
                        change_trigger = m.group(1).strip()
                if "Change Order Route:" in t:
                    m = re.search(r'Change Order Route:\s*(.*?)(?:\n|$)', t)
                    if m and m.group(1).strip():
                        change_route = m.group(1).strip()
                if "Escalation Threshold:" in t:
                    m = re.search(r'Escalation Threshold:\s*(.*?)(?:\n|$)', t)
                    if m and m.group(1).strip():
                        escalation_thresh = m.group(1).strip()

        # Check for structured Commercial Guardrails table
        cg_tbl = self._find_table_by_header(all_tables, ["commercial guardrail area", "contract policy"])
        if not cg_tbl:
            cg_tbl = self._find_table_by_header(all_tables, ["commercial guardrail area", "governance rules"])
        if cg_tbl:
            for row in cg_tbl.rows[1:]:
                cells = [clean_text(c.text) for c in row.cells]
                if len(cells) >= 2:
                    dim, val = cells[0].lower(), cells[1]
                    if "contract" in dim and val:
                        contract_implications = val
                    elif "billing" in dim and val:
                        billing_consumption = val
                    elif "staffing" in dim and val:
                        if "• Commercial Exposure:" in val or "• Exposure Note:" in val:
                            parts = re.split(r'•\s*(?:Commercial Exposure|Exposure Note):\s*', val)
                            staffing_assump = parts[0].strip()
                            if len(parts) > 1:
                                commercial_exposure = parts[1].strip()
                        else:
                            staffing_assump = val
                    elif "approved" in dim and val:
                        if "• Approved Work:" in val or "• Non-Approved Work:" in val:
                            m_app = re.search(r'Approved Work:\s*(.*?)(?:\n|•|$)', val)
                            m_non = re.search(r'Non-Approved Work:\s*(.*?)(?:\n|•|$)', val)
                            if m_app and m_app.group(1).strip():
                                approved_rule = m_app.group(1).strip()
                            if m_non and m_non.group(1).strip():
                                non_approved_rule = m_non.group(1).strip()
                        elif "non-approved" in dim:
                            non_approved_rule = val
                        else:
                            approved_rule = val
                    elif "non-approved" in dim and val:
                        non_approved_rule = val
                    elif "work-at-risk" in dim and val:
                        work_at_risk = val
                    elif "change" in dim and val:
                        if "• Triggers:" in val or "• Route:" in val:
                            m_trig = re.search(r'Triggers:\s*(.*?)(?:\n|•|$)', val)
                            m_rt = re.search(r'Route:\s*(.*?)(?:\n|•|$)', val)
                            if m_trig and m_trig.group(1).strip():
                                change_trigger = m_trig.group(1).strip()
                            if m_rt and m_rt.group(1).strip():
                                change_route = m_rt.group(1).strip()
                        elif "triggers" in dim:
                            change_trigger = val
                        elif "route" in dim:
                            change_route = val
                        else:
                            change_trigger = val
                    elif "budget" in dim and val:
                        if "• Budget Baseline:" in val or "• Variance Baseline:" in val:
                            m_bg = re.search(r'Budget Baseline:\s*(.*?)(?:\n|•|$)', val)
                            m_vr = re.search(r'Variance Baseline:\s*(.*?)(?:\n|•|$)', val)
                            if m_bg and m_bg.group(1).strip():
                                budget_baseline = m_bg.group(1).strip()
                            if m_vr and m_vr.group(1).strip():
                                variance_ind = m_vr.group(1).strip()
                        else:
                            budget_baseline = val
                    elif "margin" in dim and val:
                        m_mar = re.search(r'Margin Risk Level:\s*([^\s—]+)', val)
                        if m_mar:
                            margin_risk = m_mar.group(1).strip()
                        else:
                            margin_risk = val
                    elif "escalation" in dim and val:
                        escalation_thresh = val

        return CommercialGuardrail(
            contract_type_implication=contract_implications,
            billing_consumption_assumption=billing_consumption,
            staffing_assumption=staffing_assump,
            commercial_exposure_note=commercial_exposure,
            approved_work_rule=approved_rule,
            non_approved_work_rule=non_approved_rule,
            work_at_risk_rule=work_at_risk,
            change_control_trigger=change_trigger,
            change_order_route=change_route,
            budget_baseline=budget_baseline,
            variance_indicator=variance_ind,
            margin_risk_indicator=margin_risk,
            escalation_threshold=escalation_thresh,
            source_reference=src_ref
        )

    def _parse_talent_onboarding(
        self,
        doc: docx.Document,
        tables: List[Table],
        pmo_lead: str,
        dm: str,
        tpm: str,
        src_ref: SourceReference
    ) -> TalentOnboardingRecord:
        """Parse Talent Onboarding record and talent roster table."""
        roster: List[TalentMember] = []

        # Check for leadership paragraph overrides in Section 3.4
        def _is_assigned(val_str: Optional[str]) -> bool:
            if not val_str:
                return False
            v = val_str.strip().upper()
            return not ("UNASSIGNED" in v or "CONFIRMATION" in v or v in ("", "NONE", "N/A", "TBD"))

        for p in doc.paragraphs:
            t = p.text.strip()
            if "Onboarding Leadership:" in t:
                m_tpm = re.search(r'Talent PM:\s*([^|]+)', t)
                m_dm = re.search(r'Delivery Manager:\s*([^|]+)', t)
                m_pmo = re.search(r'PMO Lead:\s*([^|]+)', t)
                if m_tpm and m_tpm.group(1).strip():
                    val = m_tpm.group(1).strip()
                    if _is_assigned(val) or not _is_assigned(tpm):
                        tpm = val
                if m_dm and m_dm.group(1).strip():
                    val = m_dm.group(1).strip()
                    if _is_assigned(val) or not _is_assigned(dm):
                        dm = val
                if m_pmo and m_pmo.group(1).strip():
                    val = m_pmo.group(1).strip()
                    if _is_assigned(val) or not _is_assigned(pmo_lead):
                        pmo_lead = val

        # Parse Roster table
        tbl = self._find_table_by_header(tables, ["role", "named talent", "staffing status"])
        if not tbl:
            tbl = self._find_table_by_header(tables, ["named talent", "skills", "status"])

        if tbl:
            for row in tbl.rows[1:]:
                cells = [clean_text(c.text) for c in row.cells]
                if len(cells) >= 4:
                    status_val = cells[3]
                    if status_val.lower() in ("confirmed", "active", "approved", "staffed", "ready"):
                        status_val = "Confirmed"
                    elif status_val.lower() in ("pending", "needs alignment", "unassigned", "staffing required"):
                        status_val = "Staffing Required"
                    elif status_val not in ("Confirmed", "Pending", "Needs Alignment", "Staffing Required"):
                        status_val = "Confirmed" if status_val.lower() in ("confirmed", "active", "approved", "staffed", "ready") else "Staffing Required"

                    roster.append(TalentMember(
                        role=cells[0],
                        name=cells[1],
                        required_skills=cells[2],
                        status=status_val
                    ))

        return TalentOnboardingRecord(
            pmo_lead=pmo_lead,
            delivery_manager=dm,
            talent_pm=tpm,
            delivery_talent_roster=roster,
            required_roles=[tm.role for tm in roster],
            source_reference=src_ref
        )

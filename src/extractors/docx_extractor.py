"""DOCX document extractor using python-docx."""

from pathlib import Path
from typing import List, Set, Dict, Any
import docx
from src.extractors.base import BaseDocumentExtractor
from src.core.models import ExtractedDocument, DocumentSection


class DocxExtractor(BaseDocumentExtractor):
    """Extractor for Microsoft Word (.docx) documents."""

    supported_extensions: Set[str] = {".docx"}

    def extract(self, file_path: Path) -> ExtractedDocument:
        self._validate_file(file_path)

        doc = docx.Document(str(file_path))
        sections: List[DocumentSection] = []
        full_text_parts: List[str] = []

        current_section_title = "Introduction"
        current_section_lines: List[str] = []

        def flush_section():
            nonlocal current_section_title, current_section_lines
            content = self._clean_text("\n".join(current_section_lines))
            if content:
                sections.append(
                    DocumentSection(
                        title=current_section_title,
                        content=content,
                        metadata={"section_name": current_section_title}
                    )
                )
                full_text_parts.append(f"=== {current_section_title} ===\n{content}")
            current_section_lines = []

        for element in doc.element.body:
            tag = element.tag.split("}")[-1] if "}" in element.tag else element.tag

            if tag == "p":
                # Find matching paragraph
                p_text = element.text or ""
                # Check style name if possible
                style_name = ""
                for p in doc.paragraphs:
                    if p._element == element:
                        p_text = p.text
                        style_name = p.style.name if p.style else ""
                        break

                p_text_clean = p_text.strip()
                if not p_text_clean:
                    continue

                if "Heading" in style_name or style_name.startswith("Title"):
                    flush_section()
                    current_section_title = p_text_clean
                else:
                    current_section_lines.append(p_text_clean)

            elif tag == "tbl":
                # Find matching table
                for tbl in doc.tables:
                    if tbl._element == element:
                        table_rows = []
                        for row in tbl.rows:
                            row_cells = [self._clean_text(cell.text).replace("\n", " ") for cell in row.cells]
                            table_rows.append(" | ".join(row_cells))
                        if table_rows:
                            tbl_str = "[Table]\n" + "\n".join(table_rows)
                            current_section_lines.append(tbl_str)
                        break

        flush_section()

        # Fallback if no sections were flushed (e.g., if elements weren't caught in body iteration)
        if not sections:
            all_paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            for tbl in doc.tables:
                table_rows = []
                for row in tbl.rows:
                    row_cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                    table_rows.append(" | ".join(row_cells))
                if table_rows:
                    all_paras.append("[Table]\n" + "\n".join(table_rows))
            all_text = "\n\n".join(all_paras)
            sections.append(
                DocumentSection(
                    title="Main Content",
                    content=all_text,
                    metadata={"section_name": "Main Content"}
                )
            )
            full_text_parts.append(all_text)

        metadata: Dict[str, Any] = {
            "title": doc.core_properties.title or "",
            "author": doc.core_properties.author or "",
            "paragraphs_count": len(doc.paragraphs),
            "tables_count": len(doc.tables),
        }

        combined_text = "\n\n".join(full_text_parts)

        return ExtractedDocument(
            file_name=file_path.name,
            file_type="docx",
            file_path=file_path,
            text_content=combined_text,
            sections=sections,
            metadata=metadata
        )

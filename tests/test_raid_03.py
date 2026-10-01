"""Tests for RAID-03 contract reference formatting and description citation stripping."""

import pytest
from datetime import date
from src.generators.pmo_workbook.builder import extract_contract_reference, strip_citation_from_clause, build_workbook_model
from src.tools.check_artifacts import check_artifacts_directory
import json


def test_raid_03_contract_reference_formatting():
    """RAID-03: SOW refs: label, section lists, .txt and .md citations."""
    # 1. SOW refs: label for story IDs and Deliverable N.N
    ref1, _ = extract_contract_reference("sow.txt, Section 2 and Section 3 (Deliverable 3.2): Phases and Milestones")
    assert ref1 == "sow.txt | SOW refs: Deliverable 3.2 | Sections: 2, 3"

    ref2, _ = extract_contract_reference("notes.md, Section 4: Task 12 and WBS 1.2 details")
    assert "notes.md" in ref2
    assert "SOW refs: Task 12, WBS 1.2" in ref2
    assert "Sections: 4" in ref2

    # 2. Section lists (never trailing comma, never empty parts)
    ref3, _ = extract_contract_reference("Section 2 and Section 3 and Section 5: Scope details")
    assert ref3 == "Sections: 2, 3, 5"
    assert not ref3.endswith(",")
    assert "||" not in ref3


def test_raid_03_description_citation_stripping():
    """RAID-03: Remove citation from description wherever it appears after category prefix."""
    raw1 = "Scope Contradiction: sow.txt, Section 2 and Section 3 (Deliverable 3.2): Phases and Milestones: P3 is named \"Reconciliation and Final Cutover\"."
    cleaned1 = strip_citation_from_clause(raw1, category="Scope Contradiction")
    assert cleaned1 == "Phases and Milestones: P3 is named \"Reconciliation and Final Cutover\"."

    raw2 = "Ambiguous Acceptance: sow.txt, Section 3 (Deliverable 1.2): Contracted Deliverables: 1.2 is \"Security and Encryption Controls\"."
    cleaned2 = strip_citation_from_clause(raw2, category="Ambiguous Acceptance")
    assert cleaned2 == "Contracted Deliverables: 1.2 is \"Security and Encryption Controls\"."

    raw3 = "Unclear SLA: sow.txt, Section 3 (Deliverable 2.1): Contracted Deliverables: 2.1 is \"Event-driven Kafka ingestion pipelines\"."
    cleaned3 = strip_citation_from_clause(raw3, category="Unclear SLA")
    assert cleaned3 == "Contracted Deliverables: 2.1 is \"Event-driven Kafka ingestion pipelines\"."


def test_raid_03_numbered_deliverables_raid_rows():
    """RAID-03: Verify numbered_deliverables cases RAID-09, RAID-12, RAID-13 in model."""
    from src.core.models import StartupKitBaseline
    from src.llm.validation import validate_and_repair_baseline
    with open("tests/fixtures/sow/numbered_deliverables/baseline.json", "r", encoding="utf-8") as f:
        data = json.load(f)
    baseline = StartupKitBaseline.model_validate(data)
    validate_and_repair_baseline(baseline)

    model = build_workbook_model(baseline, start_date=date(2026, 10, 5))
    raid_map = {r.raid_id: r for r in model.raid_rows}

    # RAID-09 (AMB-06)
    r9 = raid_map["RAID-09"]
    assert r9.contract_reference == "sow.txt | SOW refs: Deliverable 3.2 | Sections: 2, 3"
    assert not r9.description.startswith("Scope Contradiction: sow.txt")
    assert r9.description.startswith("Scope Contradiction: Phases and Milestones:")

    # RAID-12 (AMB-09)
    r12 = raid_map["RAID-12"]
    assert r12.contract_reference == "sow.txt | SOW refs: Deliverable 1.2 | Sections: 3"
    assert not r12.description.startswith("Ambiguous Acceptance: sow.txt")
    assert r12.description.startswith("Ambiguous Acceptance: Contracted Deliverables: 1.2")

    # RAID-13 (AMB-10)
    r13 = raid_map["RAID-13"]
    assert r13.contract_reference == "sow.txt | SOW refs: Deliverable 2.1 | Sections: 3"
    assert not r13.description.startswith("Unclear SLA: sow.txt")
    assert r13.description.startswith("Unclear SLA: Contracted Deliverables: 2.1")


def test_inv_13_and_inv_23_catch_old_formats(arc_baseline, tmp_path):
    """RAID-03: Confirm INV-13 and INV-23 catch old formats."""
    import openpyxl
    from src.tools.check_artifacts import check_artifacts_directory
    from src.generators.pmo_workbook.writer import write_workbook
    from src.generators.docx_generator import DocxGenerator
    from src.llm.validation import validate_and_repair_baseline

    b = arc_baseline.model_copy(deep=True)
    validate_and_repair_baseline(b)
    gen = DocxGenerator()
    gen.write_kit_docx(b, tmp_path)
    gen.write_checklist_docx(b, tmp_path)
    model = build_workbook_model(b, start_date=date(2026, 10, 5))
    wb_file = tmp_path / "ARC_Genomics_Platform_Project_Delivery_Workbook.xlsx"
    write_workbook(model, wb_file)

    # 1. INV-13 catches old 'Stories:' label in Contract Reference
    wb = openpyxl.load_workbook(wb_file)
    ws = wb["RAID Log"]
    ws.cell(row=6, column=4, value="Stories: HS-4762 | Sections: 2, 3")
    wb.save(wb_file)
    violations = check_artifacts_directory(tmp_path)
    assert any(v.inv_id == "INV-13" for v in violations)

    # 2. INV-13 catches description starting with citation
    wb = openpyxl.load_workbook(wb_file)
    ws = wb["RAID Log"]
    ws.cell(row=6, column=4, value="SOW refs: HS-4762 | Sections: 2, 3")
    ws.cell(row=6, column=3, value="sow.txt, Section 2: Scope description")
    wb.save(wb_file)
    violations = check_artifacts_directory(tmp_path)
    assert any(v.inv_id == "INV-13" for v in violations)

    # 3. INV-23 catches invalid Contract Reference lacking valid document / section / sow ref
    wb = openpyxl.load_workbook(wb_file)
    ws = wb["RAID Log"]
    ws.cell(row=6, column=3, value="Scope description without citation")
    ws.cell(row=6, column=4, value="schedule impact")
    wb.save(wb_file)
    violations = check_artifacts_directory(tmp_path)
    assert any(v.inv_id == "INV-23" for v in violations)

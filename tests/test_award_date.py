"""Tests for VAL-05, DT-01, and KIT-09 award date provenance and start date calculations."""

import pytest
from datetime import date
from src.core.models import (
    StartupKitBaseline,
    ProjectStartupCharter,
    Milestone,
)
from src.llm.validation import validate_and_repair_baseline
from src.generators.pmo_workbook.builder import calculate_start_date


def test_val_05_and_dt_01_award_date_fallback():
    """Verify that absent award date does not use today - 1 day, and start date falls back to generation date (VAL-05, DT-01)."""
    baseline = StartupKitBaseline(
        project_name="Award Date Test",
        sow_awarded_date=None,
        open_questions=[]
    )
    report = validate_and_repair_baseline(baseline)

    # Invariant INV-18: warning on absent award date
    assert any(f.invariant_id == "INV-18" for f in report.warnings)
    assert any("award" in q.lower() for q in baseline.open_questions)

    # Start date calculation with today=2026-10-01 (Thursday)
    today = date(2026, 10, 1)
    calc_start, basis = calculate_start_date(baseline, user_start_date=None, today=today)
    assert calc_start == date(2026, 10, 5)  # First Monday on or after generation date
    assert "generation" in basis.lower()


def test_val_11_provided_requires_stated_provenance():
    """VAL-11: Only an award date carrying VAL-11 'stated' provenance is labelled 'Provided'."""
    today = date(2026, 10, 2)

    # Award date with no provenance (e.g. the retired date.today()-1 fallback) is treated as no date
    unprovenanced = StartupKitBaseline(project_name="Legacy Fixture", sow_awarded_date=date(2026, 9, 30))
    assert unprovenanced.award_date_source is None
    calc_start, basis = calculate_start_date(unprovenanced, user_start_date=None, today=today)
    assert calc_start == date(2026, 10, 5)
    assert basis == "Assumed - first Monday after generation date; confirm"
    assert "Provided" not in basis

    # Award date stated in the SOW and extracted by VAL-11
    stated = StartupKitBaseline(project_name="Stated", sow_awarded_date=date(2026, 9, 30), award_date_source="stated")
    calc_start, basis = calculate_start_date(stated, user_start_date=None, today=today)
    assert calc_start == date(2026, 9, 30)
    assert basis == "Provided"

    # Provenance flag survives a JSON round trip; legacy baseline.json without the key loads as unprovenanced
    assert StartupKitBaseline.model_validate(stated.model_dump(mode="json")).award_date_source == "stated"
    legacy = unprovenanced.model_dump(mode="json")
    legacy.pop("award_date_source")
    assert StartupKitBaseline.model_validate(legacy).award_date_source is None


def test_val_11_aggregator_sets_provenance_only_with_a_date():
    """VAL-11: The aggregator records 'stated' provenance only alongside a stated award date."""
    from src.core.models import CharterExtraction, DeliverablesExtraction, MilestonesExtraction, RAIDExtraction
    from src.llm.aggregator import BaselineAggregator

    def run(**kwargs):
        return BaselineAggregator().aggregate(
            charter=CharterExtraction(project_name="Agg"),
            deliverables_ext=DeliverablesExtraction(),
            milestones_ext=MilestonesExtraction(),
            raid_ext=RAIDExtraction(),
            kit_drafted_date=date(2026, 10, 2),
            **kwargs,
        )

    stated = run(sow_awarded_date=date(2026, 10, 1), award_date_source="stated")
    assert stated.sow_awarded_date == date(2026, 10, 1) and stated.award_date_source == "stated"

    no_source = run(sow_awarded_date=date(2026, 10, 1))
    assert no_source.award_date_source is None

    no_date = run(sow_awarded_date=None, award_date_source="stated")
    assert no_date.sow_awarded_date is None and no_date.award_date_source is None


def _render_sla_text(baseline, tmp_path):
    """Render Kit and Checklist and return (kit header SLA, kit charter SLA, checklist summary, G01-01 evidence)."""
    import docx
    from src.generators.docx_generator import DocxGenerator

    validate_and_repair_baseline(baseline)
    writer = DocxGenerator(generated_date="2026-10-01")
    kit = docx.Document(str(writer.write_kit_docx(baseline, tmp_path)))
    chk = docx.Document(str(writer.write_checklist_docx(baseline, tmp_path)))

    def cell_after(doc, label):
        for t in doc.tables:
            for r in t.rows:
                cells = [c.text for c in r.cells]
                for i, txt in enumerate(cells[:-1]):
                    if txt.strip() == label:
                        return cells[i + 1]
        return None

    summary = next(t.cell(0, 0).text for t in chk.tables if "1-Day Creation SLA Status" in t.cell(0, 0).text)
    g01_01 = next(i.evidence for i in baseline.readiness_checklist if i.item_id == "G01-01")
    return cell_after(kit, "1-Day SLA Status"), cell_after(kit, "Turnaround SLA & PMO Authorization"), summary, g01_01


def _load_fixture_baseline(name):
    import json
    from pathlib import Path
    with open(Path("tests/fixtures/sow") / name / "baseline.json", encoding="utf-8") as f:
        return StartupKitBaseline.model_validate(json.load(f))


def test_val_05_sla_status_requires_stated_provenance(tmp_path):
    """VAL-05: Kit/Checklist SLA status is determinate only for an award date with VAL-11 'stated' provenance."""
    baseline = _load_fixture_baseline("mock_sow")
    # Recorded with the retired date.today()-1 fallback: a date, but no provenance
    assert baseline.sow_awarded_date == date(2026, 9, 30) and baseline.award_date_source is None

    kit_header, kit_charter, chk_summary, g01_01 = _render_sla_text(baseline, tmp_path)
    assert kit_header == "Not determinable - award date not stated"
    assert kit_charter.startswith("Not determinable - award date not stated")
    assert "1-Day Creation SLA Status: Not determinable - award date not stated" in chk_summary
    assert "Project awarded" not in g01_01 and "SLA Met" not in g01_01
    assert "Award date not stated in SOW [CONFIRMATION REQUIRED]" in g01_01
    g01 = next(i for i in baseline.readiness_checklist if i.item_id == "G01-01")
    assert g01.status == "Confirmation Required"


def test_val_05_sla_status_determinate_with_stated_provenance(tmp_path):
    """VAL-05: A VAL-11 stated award date still yields a determinate SLA status."""
    baseline = _load_fixture_baseline("mock_sow").model_copy(update={"award_date_source": "stated"})

    kit_header, kit_charter, chk_summary, g01_01 = _render_sla_text(baseline, tmp_path / "stated")
    assert kit_header == "Met (Drafted <= 1 day)"
    assert kit_charter.startswith("SLA Met (Yes)")
    assert "1-Day Creation SLA Status: Met (Drafted within 1 business day)" in chk_summary
    assert "(Project awarded 2026-09-30). SLA Met." in g01_01


def test_val_05_kit_reingestion_preserves_kit_sla_verdict(tmp_path):
    """VAL-05 stopgap: re-ingestion keeps the Kit's own SLA verdict and G01-01 status, never upgrading or recomputing it."""
    from src.extractors.startup_kit_docx_parser import StartupKitDocxParser
    from src.generators.docx_generator import DocxGenerator

    def roundtrip(baseline, folder):
        validate_and_repair_baseline(baseline)
        kit_path = DocxGenerator(generated_date="2026-10-01").write_kit_docx(baseline, folder)
        DocxGenerator(generated_date="2026-10-01").write_checklist_docx(baseline, folder)
        return baseline, StartupKitDocxParser().parse_startup_kit_docx(kit_path)

    # Kit generated from a stated award date shows 'Met'; re-ingestion keeps it without inventing 'stated' provenance
    src, parsed = roundtrip(_load_fixture_baseline("mock_sow").model_copy(update={"award_date_source": "stated"}), tmp_path / "met")
    src_g01 = next(i for i in src.readiness_checklist if i.item_id == "G01-01")
    parsed_g01 = next(i for i in parsed.readiness_checklist if i.item_id == "G01-01")
    assert parsed.award_date_source is None and parsed.sla_verdict_source == "kit_reingested"
    assert parsed.sla_met is True
    assert parsed_g01.status == src_g01.status == "Complete"
    assert parsed_g01.evidence == src_g01.evidence

    # Kit generated from an unprovenanced date shows 'Not determinable'; re-ingestion must not upgrade it
    src, parsed = roundtrip(_load_fixture_baseline("mock_sow"), tmp_path / "nd")
    parsed_g01 = next(i for i in parsed.readiness_checklist if i.item_id == "G01-01")
    assert parsed.award_date_source is None and parsed.sla_verdict_source is None
    assert parsed.sla_met is False
    assert parsed_g01.status == "Confirmation Required"
    assert "Award date not stated in SOW" in parsed_g01.evidence

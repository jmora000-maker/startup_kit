"""Tests for MAP-08: SOW reference format normalization in traceability check."""

from datetime import date
import pytest
from src.config import normalize_sow_reference, extract_sow_references
from src.core.models import StartupKitBaseline, Deliverable, Milestone, SOWWorkItem, WorkPackageSeed
from src.generators.pmo_workbook.builder import build_workbook_model
from src.llm.validation import validate_and_repair_baseline
from src.review_ui.facts import build_fact_categories


def test_normalize_sow_reference_canonical_forms():
    """MAP-08: Verify canonical normalization across various reference formats."""
    # 1. Work Output variations
    assert normalize_sow_reference("WO-01") == "WO:1"
    assert normalize_sow_reference("WO-1") == "WO:1"
    assert normalize_sow_reference("Work Output 1") == "WO:1"
    assert normalize_sow_reference("Work Output 01") == "WO:1"
    assert normalize_sow_reference("Work Output #1") == "WO:1"
    assert normalize_sow_reference("WO 1") == "WO:1"
    assert normalize_sow_reference("WO-02") == "WO:2"
    assert normalize_sow_reference("Work Output 2") == "WO:2"

    # Distinct numbers must remain distinct
    assert normalize_sow_reference("WO-01") != normalize_sow_reference("Work Output 2")

    # 2. Deliverable variations
    assert normalize_sow_reference("Deliverable 1") == "DEL:1"
    assert normalize_sow_reference("Deliverable 01") == "DEL:1"
    assert normalize_sow_reference("Deliverable 1.1") == "DEL:1.1"
    assert normalize_sow_reference("D1") == "DEL:1"
    assert normalize_sow_reference("D01") == "DEL:1"
    assert normalize_sow_reference("D-01") == "DEL:1"
    assert normalize_sow_reference("DEL-01") == "DEL:1"
    assert normalize_sow_reference("DEL-1") == "DEL:1"

    # 3. Story variations (real SOW stories)
    assert normalize_sow_reference("Story 1") == "STORY:1"
    assert normalize_sow_reference("Story 01") == "STORY:1"
    assert normalize_sow_reference("Story-01") == "STORY:1"
    assert normalize_sow_reference("STORY-12") == "STORY:12"

    # 4. Synthetic SOW references (never conflated with real Story references)
    assert normalize_sow_reference("SOW-01") == "SYNTH:1"
    assert normalize_sow_reference("SOW-1") == "SYNTH:1"
    assert normalize_sow_reference("SOW-01") != normalize_sow_reference("Story 1")

    # 5. Phased synthetic references
    assert normalize_sow_reference("SOW-P1-01") == "SYNTH:P1:1"
    assert normalize_sow_reference("SOW-P1-1") == "SYNTH:P1:1"
    assert normalize_sow_reference("SOW-P2a-02") == "SYNTH:P2A:2"

    # 6. Task / WBS variations
    assert normalize_sow_reference("Task 1") == "TASK:1"
    assert normalize_sow_reference("Task 01") == "TASK:1"
    assert normalize_sow_reference("TASK-01") == "TASK:1"
    assert normalize_sow_reference("T-01") == "TASK:1"
    assert normalize_sow_reference("WBS 1.1") == "WBS:1.1"

    # 7. Section / Clause variations
    assert normalize_sow_reference("Section 3.1") == "SECTION:3.1"
    assert normalize_sow_reference("Clause 4") == "SECTION:4"
    assert normalize_sow_reference("§ 2.1") == "SECTION:2.1"


def test_map_08_traceability_normalization_gitlab_baseline():
    """MAP-08: Deliverables tagged WO-01/WO-02 match catalogue tagged Work Output 1/Work Output 2."""
    delivs = [
        Deliverable(id="DEL-01", name="Discovery & Onboarding Summary", description="Discovery phase", sow_reference="WO-01"),
        Deliverable(id="DEL-02", name="Task-Level Project Plan", description="Project plan", sow_reference="WO-02"),
    ]
    milestones = [
        Milestone(id="M1", description="Project Kickoff"),
        Milestone(id="M2", description="Project Sign-off"),
    ]
    catalogue = [
        SOWWorkItem(reference="Work Output 1", reference_kind="Deliverable number", title="Discovery & Onboarding Summary", phase="", owner="Toptal", type="Documentation"),
        SOWWorkItem(reference="Work Output 2", reference_kind="Deliverable number", title="Task-Level Project Plan", phase="", owner="Toptal", type="Documentation"),
    ]
    baseline = StartupKitBaseline(
        project_name="GitLab Modernization & AI Enablement Discovery",
        deliverables=delivs,
        milestones=milestones,
        sow_stories_catalogue=catalogue,
        backlog_seed=[]
    )

    validate_and_repair_baseline(baseline)
    wb = build_workbook_model(baseline, today=date(2026, 10, 3))

    # SOW References in traceability dictionary
    sow_trace = wb.traceability["SOW References"]
    assert sow_trace["missing_ids"] == []

    # Validation findings should not have any TR-01 warning for missing SOW references
    tr_findings = [f for f in baseline.validation_report.findings if f.invariant_id == "TR-01" and "Unmatched SOW References" in f.message]
    assert tr_findings == []


def test_map_08_traceability_normalization_reverse_tagging():
    """MAP-08: Deliverables tagged Work Output 1/Work Output 2 match catalogue tagged WO-01/WO-02."""
    delivs = [
        Deliverable(id="DEL-01", name="Discovery & Onboarding Summary", description="Discovery phase", sow_reference="Work Output 1"),
        Deliverable(id="DEL-02", name="Task-Level Project Plan", description="Project plan", sow_reference="Work Output 2"),
    ]
    milestones = [
        Milestone(id="M1", description="Project Kickoff"),
        Milestone(id="M2", description="Project Sign-off"),
    ]
    catalogue = [
        SOWWorkItem(reference="WO-01", reference_kind="Story ID", title="Discovery & Onboarding Summary", phase="", owner="Toptal", type="Documentation"),
        SOWWorkItem(reference="WO-02", reference_kind="Story ID", title="Task-Level Project Plan", phase="", owner="Toptal", type="Documentation"),
    ]
    baseline = StartupKitBaseline(
        project_name="GitLab Modernization & AI Enablement Discovery",
        deliverables=delivs,
        milestones=milestones,
        sow_stories_catalogue=catalogue,
        backlog_seed=[]
    )

    validate_and_repair_baseline(baseline)
    wb = build_workbook_model(baseline, today=date(2026, 10, 3))

    sow_trace = wb.traceability["SOW References"]
    assert sow_trace["missing_ids"] == []

    tr_findings = [f for f in baseline.validation_report.findings if f.invariant_id == "TR-01" and "Unmatched SOW References" in f.message]
    assert tr_findings == []


def test_map_08_genuine_mismatch_detected():
    """MAP-08: Genuine mismatch (different numbers: WO-01 vs Work Output 2) is still caught."""
    delivs = [
        Deliverable(id="DEL-01", name="Discovery & Onboarding Summary", description="Discovery phase", sow_reference="WO-01"),
    ]
    milestones = [
        Milestone(id="M1", description="Project Kickoff"),
    ]
    catalogue = [
        SOWWorkItem(reference="WO-01", reference_kind="Deliverable number", title="Discovery & Onboarding Summary", phase="", owner="Toptal", type="Documentation"),
        SOWWorkItem(reference="Work Output 2", reference_kind="Deliverable number", title="Task-Level Project Plan", phase="", owner="Toptal", type="Documentation"),
    ]
    backlog = [
        WorkPackageSeed(id="WP-01", parent_deliverable_id="DEL-01", title="Discovery Task", sow_reference="WO-01", owner="Toptal")
    ]
    baseline = StartupKitBaseline(
        project_name="GitLab Modernization & AI Enablement Discovery",
        deliverables=delivs,
        milestones=milestones,
        sow_stories_catalogue=catalogue,
        backlog_seed=backlog
    )

    wb = build_workbook_model(baseline, today=date(2026, 10, 3))

    sow_trace = wb.traceability["SOW References"]
    assert "Work Output 2" in sow_trace["missing_ids"]

    tr_findings = [f for f in baseline.validation_report.findings if f.invariant_id == "TR-01" and "Unmatched SOW References" in f.message]
    assert len(tr_findings) == 1
    assert "Work Output 2" in tr_findings[0].message


def test_map_08_fact_review_clean_when_matching():
    """MAP-08: Fact Review displays no TR-01 finding when references match via canonical normalization."""
    delivs = [
        Deliverable(id="DEL-01", name="Discovery & Onboarding Summary", description="Discovery phase", sow_reference="WO-01"),
        Deliverable(id="DEL-02", name="Task-Level Project Plan", description="Project plan", sow_reference="WO-02"),
    ]
    milestones = [
        Milestone(id="M1", description="Project Kickoff"),
        Milestone(id="M2", description="Project Sign-off"),
    ]
    catalogue = [
        SOWWorkItem(reference="Work Output 1", reference_kind="Deliverable number", title="Discovery & Onboarding Summary", phase="", owner="Toptal", type="Documentation"),
        SOWWorkItem(reference="Work Output 2", reference_kind="Deliverable number", title="Task-Level Project Plan", phase="", owner="Toptal", type="Documentation"),
    ]
    baseline = StartupKitBaseline(
        project_name="GitLab Modernization & AI Enablement Discovery",
        deliverables=delivs,
        milestones=milestones,
        sow_stories_catalogue=catalogue,
        backlog_seed=[]
    )

    validate_and_repair_baseline(baseline)
    categories = {c.key: c for c in build_fact_categories(baseline)}

    if "validation_findings" in categories:
        val_cat = categories["validation_findings"]
        labels_and_values = [(f.label, f.value) for f in val_cat.fields]
        tr_fields = [
            (lbl, val) for lbl, val in labels_and_values
            if "TR-01" in lbl and "Unmatched SOW References" in val
        ]
        assert tr_fields == []


def test_map_08_synthetic_and_real_story_conflation_prevention():
    """MAP-08: A synthetic SOW-01 placeholder on a deliverable must NEVER falsely match a real 'Story 1' catalogue item."""
    # Deliverable has a synthetic SOW-01 reference
    delivs = [
        Deliverable(id="DEL-01", name="Generic Platform Work", description="Synthetic deliverable", sow_reference="SOW-01"),
    ]
    milestones = [
        Milestone(id="M1", description="Project Kickoff"),
    ]
    # SOW catalogue has an UNRELATED real Story 1 from the contract
    catalogue = [
        SOWWorkItem(
            reference="Story 1",
            reference_kind="Story ID",
            title="Real Contract Story: User Authentication Flow",
            phase="",
            owner="Toptal",
            type="Build"
        ),
    ]
    baseline = StartupKitBaseline(
        project_name="Platform Modernization",
        deliverables=delivs,
        milestones=milestones,
        sow_stories_catalogue=catalogue,
        backlog_seed=[]
    )

    wb = build_workbook_model(baseline, today=date(2026, 10, 3))

    # Traceability check must detect that 'Story 1' is missing (not falsely satisfied by synthetic 'SOW-01')
    sow_trace = wb.traceability["SOW References"]
    assert "Story 1" in sow_trace["missing_ids"]

    # Must also emit a TR-01 ValidationFinding warning
    assert baseline.validation_report is not None
    tr_findings = [f for f in baseline.validation_report.findings if f.invariant_id == "TR-01" and "Unmatched SOW References" in f.message]
    assert len(tr_findings) == 1
    assert "Story 1" in tr_findings[0].message


def test_map_12_hypercare_capacity_non_numeric_synthetic_reference_matches():
    """MAP-12: Non-numeric synthetic reference 'Hypercare Capacity' matches between deliverable/backlog and SOW catalogue without false warnings."""
    deliv = Deliverable(
        id="DEL-20",
        name="Four-Week Hypercare Support",
        description="Post-launch hypercare support",
        sow_reference="Hypercare Capacity",
    )
    milestone = Milestone(
        id="M4",
        description="P3 Launch accepted: System live and hypercare support delivered.",
    )
    catalogue_item = SOWWorkItem(
        reference="Hypercare Capacity",
        reference_kind="Synthetic",
        title="Hypercare Capacity: up to 160 hours",
        phase="P3 Launch",
        owner="Toptal Delivery Team",
        type="Build",
    )
    wp = WorkPackageSeed(
        id="WP-36",
        parent_deliverable_id="DEL-20",
        title="Hypercare Capacity: up to 160 hours of developer and QA support",
        sow_reference="Hypercare Capacity",
        owner="Toptal Delivery Team",
    )
    baseline = StartupKitBaseline(
        project_name="ARC Genomics Platform Implementation",
        deliverables=[deliv],
        milestones=[milestone],
        sow_stories_catalogue=[catalogue_item],
        backlog_seed=[wp],
    )

    wb = build_workbook_model(baseline, today=date(2026, 10, 10))

    sow_trace = wb.traceability["SOW References"]
    assert sow_trace["missing_ids"] == []
    assert sow_trace["in_baseline"] == 1
    assert sow_trace["in_workbook"] == 1
    assert sow_trace["synthetic"] is True

    # No TR-01 warning should be emitted for SOW references
    if baseline.validation_report:
        tr_findings = [f for f in baseline.validation_report.findings if f.invariant_id == "TR-01" and "Unmatched SOW References" in f.message]
        assert tr_findings == []


def test_map_12_genuinely_missing_non_numeric_reference_produces_warning():
    """MAP-12: A genuinely missing non-numeric reference in catalogue produces clear Unmatched SOW References warning."""
    deliv = Deliverable(
        id="DEL-01",
        name="Discovery & Planning",
        description="Discovery phase",
        sow_reference="WO-01",
    )
    milestone = Milestone(
        id="M1",
        description="Project Kickoff",
    )
    catalogue = [
        SOWWorkItem(
            reference="WO-01",
            reference_kind="Deliverable number",
            title="Discovery & Planning",
            phase="",
            owner="Toptal",
            type="Documentation",
        ),
        SOWWorkItem(
            reference="Hypercare Capacity",
            reference_kind="Synthetic",
            title="Hypercare Capacity: up to 160 hours",
            phase="",
            owner="Toptal",
            type="Build",
        ),
    ]
    wp = WorkPackageSeed(
        id="WP-01",
        parent_deliverable_id="DEL-01",
        title="Discovery Task",
        sow_reference="WO-01",
        owner="Toptal",
    )
    baseline = StartupKitBaseline(
        project_name="Platform Implementation",
        deliverables=[deliv],
        milestones=[milestone],
        sow_stories_catalogue=catalogue,
        backlog_seed=[wp],
    )

    wb = build_workbook_model(baseline, today=date(2026, 10, 10))

    sow_trace = wb.traceability["SOW References"]
    assert "Hypercare Capacity" in sow_trace["missing_ids"]

    assert baseline.validation_report is not None
    tr_findings = [f for f in baseline.validation_report.findings if f.invariant_id == "TR-01" and "Unmatched SOW References" in f.message]
    assert len(tr_findings) == 1
    assert "Hypercare Capacity" in tr_findings[0].message


def test_map_12_comma_separated_references_including_synthetic_match():
    """MAP-12: Deliverable with comma-separated references ('WO-01, Hypercare Capacity') matches both catalogue items."""
    deliv = Deliverable(
        id="DEL-01",
        name="Hybrid Scope Package",
        description="Delivery and hypercare scope",
        sow_reference="WO-01, Hypercare Capacity",
    )
    milestone = Milestone(
        id="M1",
        description="Gate 1",
    )
    catalogue = [
        SOWWorkItem(
            reference="Work Output 1",
            reference_kind="Deliverable number",
            title="Discovery & Planning",
            phase="",
            owner="Toptal",
            type="Documentation",
        ),
        SOWWorkItem(
            reference="Hypercare Capacity",
            reference_kind="Synthetic",
            title="Hypercare Capacity",
            phase="",
            owner="Toptal",
            type="Build",
        ),
    ]
    wp = WorkPackageSeed(
        id="WP-01",
        parent_deliverable_id="DEL-01",
        title="Delivery Task",
        sow_reference="WO-01, Hypercare Capacity",
        owner="Toptal",
    )
    baseline = StartupKitBaseline(
        project_name="Platform Implementation",
        deliverables=[deliv],
        milestones=[milestone],
        sow_stories_catalogue=catalogue,
        backlog_seed=[wp],
    )

    wb = build_workbook_model(baseline, today=date(2026, 10, 10))

    sow_trace = wb.traceability["SOW References"]
    assert sow_trace["missing_ids"] == []
    assert sow_trace["in_baseline"] == 3
    assert sow_trace["in_workbook"] == 2
    if baseline.validation_report:
        tr_findings = [f for f in baseline.validation_report.findings if f.invariant_id == "TR-01" and "Unmatched SOW References" in f.message]
        assert tr_findings == []


def test_map_14_synthetic_milestone_prefixed_references_milestone_first_ordering():
    """MAP-14: Synthetic SOW references with milestone-first ordering (M1-WO-01..M4-WO-08) match perfectly across catalogue and backlog."""
    # Undivide Project SOW shape: 4 Milestones, 19 Deliverables, 27 bulleted Work Outputs across milestones
    milestones = [
        Milestone(id="M1", description="Phase 1 - Data and Backend Foundation"),
        Milestone(id="M2", description="Phase 2 - LLM Enrichment Layer"),
        Milestone(id="M3", description="Phase 3 - Access Control, Observability, and Pilot Readiness"),
        Milestone(id="M4", description="Phase 4 - Launch and Handoff"),
    ]
    deliv_names = [
        "PostgreSQL Core Schema", "Ingestion Engine", "Data Migration Scripts", "Transformation Worker", "Validation Rules",
        "LLM API Gateway", "Inference Router", "Prompt Pipeline", "Context Enrichment", "Vector Store",
        "RBAC Module", "Audit Logger", "Observability Dashboard", "Alerting Policies", "Load Testing Report",
        "Security Review", "Operator Runbook", "UAT Sign-off Report", "Pilot Deployment Package"
    ]
    delivs = [
        Deliverable(id=f"DEL-{i+1:02d}", name=deliv_names[i], description=f"Scope for {deliv_names[i]}", sow_reference="")
        for i in range(19)
    ]
    catalogue_items = []
    # 27 catalogue items distributed across 19 deliverables
    # M1: 7 work items (M1-WO-01 .. M1-WO-07) mapped to DEL-01..DEL-05
    for i in range(1, 8):
        deliv_idx = min(i, 5)
        catalogue_items.append(SOWWorkItem(
            reference=f"M1-WO-{i:02d}",
            reference_kind="Synthetic",
            title=f"Database schema and migrations component {i}",
            phase="Phase 1",
            owner="Toptal Delivery Team",
            type="Build",
            deliverable_id=f"DEL-{deliv_idx:02d}"
        ))
    # M2: 6 work items (M2-WO-01 .. M2-WO-06) mapped to DEL-06..DEL-10
    for i in range(1, 7):
        deliv_idx = 5 + min(i, 5)
        catalogue_items.append(SOWWorkItem(
            reference=f"M2-WO-{i:02d}",
            reference_kind="Synthetic",
            title=f"LLM gateway component {i}",
            phase="Phase 2",
            owner="Toptal Delivery Team",
            type="Build",
            deliverable_id=f"DEL-{deliv_idx:02d}"
        ))
    # M3: 6 work items (M3-WO-01 .. M3-WO-06) mapped to DEL-11..DEL-15
    for i in range(1, 7):
        deliv_idx = 10 + min(i, 5)
        catalogue_items.append(SOWWorkItem(
            reference=f"M3-WO-{i:02d}",
            reference_kind="Synthetic",
            title=f"Security component {i}",
            phase="Phase 3",
            owner="Toptal Delivery Team",
            type="Build",
            deliverable_id=f"DEL-{deliv_idx:02d}"
        ))
    # M4: 8 work items (M4-WO-01 .. M4-WO-08) mapped to DEL-16..DEL-19
    for i in range(1, 9):
        deliv_idx = 15 + min(i, 4)
        catalogue_items.append(SOWWorkItem(
            reference=f"M4-WO-{i:02d}",
            reference_kind="Synthetic",
            title=f"Deployment component {i}",
            phase="Phase 4",
            owner="Toptal Delivery Team",
            type="Build",
            deliverable_id=f"DEL-{deliv_idx:02d}"
        ))

    assert len(catalogue_items) == 27

    baseline = StartupKitBaseline(
        project_name="The Undivide Project - Pilot-Ready Product Implementation",
        deliverables=delivs,
        milestones=milestones,
        sow_stories_catalogue=catalogue_items,
        backlog_seed=[]
    )

    validate_and_repair_baseline(baseline)

    # Confirm backlog seed was rebuilt preserving full milestone-prefixed references
    assert len(baseline.backlog_seed) == 27
    assert baseline.backlog_seed[0].sow_reference == "M1-WO-01"
    assert baseline.backlog_seed[26].sow_reference == "M4-WO-08"

    wb = build_workbook_model(baseline, today=date(2026, 10, 16))

    # Zero unmatched findings in workbook traceability
    sow_trace = wb.traceability["SOW References"]
    assert sow_trace["missing_ids"] == [], f"Expected 0 missing IDs, got {len(sow_trace['missing_ids'])}: {sow_trace['missing_ids']}"
    assert sow_trace["in_baseline"] == 27
    assert sow_trace["in_workbook"] == 27

    # Zero TR-01 findings in validation report
    if baseline.validation_report:
        tr_findings = [f for f in baseline.validation_report.findings if f.invariant_id == "TR-01" and "Unmatched SOW References" in f.message]
        assert tr_findings == []


def test_map_14_synthetic_milestone_prefixed_references_wo_first_ordering():
    """MAP-14 (reopened): Synthetic SOW references with WO-first ordering (WO-M1-01..WO-M4-08) match perfectly across catalogue and backlog."""
    # Run 2 SOW shape: WO-prefix-first catalogue references (WO-M1-01 .. WO-M4-08)
    milestones = [
        Milestone(id="M1", description="Phase 1 - Data and Backend Foundation"),
        Milestone(id="M2", description="Phase 2 - LLM Enrichment Layer"),
        Milestone(id="M3", description="Phase 3 - Access Control, Observability, and Pilot Readiness"),
        Milestone(id="M4", description="Phase 4 - Launch and Handoff"),
    ]
    deliv_names = [
        "Postgres Schema and Facility Data Migration",
        "Normalization Layer and Ingestion Framework",
        "Infrastructure-as-Code for AWS Data Tier",
        "Launch Source Ingestion Connectors with Data-Quality Checks",
        "Scraper Interface Specification and News-Site Scraper Connector",
        "Authenticated API and Backend Test Suite",
        "LLM Enrichment Service",
        "Confidence-Score Capture and Review-Queue Routing",
        "Review-Queue Admin User Interface",
        "Prompt Versioning, Output Caching and Cost Controls",
        "Ingestion Normalizer Test Suite",
        "Role-Based Access Control for Four Personas",
        "Role-Based Views and Admin Panel",
        "Observability Stack",
        "CI/CD Pipeline and Integrated Staging Deployment",
        "Production Deployment and Pilot Partner Onboarding Support",
        "Architecture, Runbook and Data Structure Documentation",
        "Mentoring Runbook",
        "Design Handoff, Knowledge Transfer and Handover Checklist"
    ]
    delivs = [
        Deliverable(id=f"DEL-{i+1:02d}", name=deliv_names[i], description=f"Scope for {deliv_names[i]}", sow_reference="")
        for i in range(19)
    ]
    catalogue_items = []
    # 27 catalogue items distributed across 19 deliverables with WO-M{x}-{yy} ordering
    for i in range(1, 8):
        deliv_idx = min(i, 5)
        catalogue_items.append(SOWWorkItem(
            reference=f"WO-M1-{i:02d}",
            reference_kind="Synthetic",
            title=f"Database schema and migrations component {i}",
            phase="Phase 1",
            owner="Toptal Delivery Team",
            type="Build",
            deliverable_id=f"DEL-{deliv_idx:02d}"
        ))
    for i in range(1, 7):
        deliv_idx = 5 + min(i, 5)
        catalogue_items.append(SOWWorkItem(
            reference=f"WO-M2-{i:02d}",
            reference_kind="Synthetic",
            title=f"LLM gateway component {i}",
            phase="Phase 2",
            owner="Toptal Delivery Team",
            type="Build",
            deliverable_id=f"DEL-{deliv_idx:02d}"
        ))
    for i in range(1, 7):
        deliv_idx = 10 + min(i, 5)
        catalogue_items.append(SOWWorkItem(
            reference=f"WO-M3-{i:02d}",
            reference_kind="Synthetic",
            title=f"Security component {i}",
            phase="Phase 3",
            owner="Toptal Delivery Team",
            type="Build",
            deliverable_id=f"DEL-{deliv_idx:02d}"
        ))
    for i in range(1, 9):
        deliv_idx = 15 + min(i, 4)
        catalogue_items.append(SOWWorkItem(
            reference=f"WO-M4-{i:02d}",
            reference_kind="Synthetic",
            title=f"Deployment component {i}",
            phase="Phase 4",
            owner="Toptal Delivery Team",
            type="Build",
            deliverable_id=f"DEL-{deliv_idx:02d}"
        ))

    assert len(catalogue_items) == 27

    baseline = StartupKitBaseline(
        project_name="The Undivide Project - Pilot-Ready Product Implementation",
        deliverables=delivs,
        milestones=milestones,
        sow_stories_catalogue=catalogue_items,
        backlog_seed=[]
    )

    validate_and_repair_baseline(baseline)

    # Confirm backlog seed preserves the full WO-M1-01 string without dropping WO token
    assert len(baseline.backlog_seed) == 27
    assert baseline.backlog_seed[0].sow_reference == "WO-M1-01"
    assert baseline.backlog_seed[26].sow_reference == "WO-M4-08"

    wb = build_workbook_model(baseline, today=date(2026, 10, 16))

    # Zero unmatched findings in workbook traceability
    sow_trace = wb.traceability["SOW References"]
    assert sow_trace["missing_ids"] == [], f"Expected 0 missing IDs, got {len(sow_trace['missing_ids'])}: {sow_trace['missing_ids']}"
    assert sow_trace["in_baseline"] == 27
    assert sow_trace["in_workbook"] == 27

    # Zero TR-01 findings in validation report
    if baseline.validation_report:
        tr_findings = [f for f in baseline.validation_report.findings if f.invariant_id == "TR-01" and "Unmatched SOW References" in f.message]
        assert tr_findings == []


def test_map_14_order_independent_synthetic_token_normalization():
    """MAP-14: Test order-independent canonicalization across diverse token permutations."""
    # Milestone + WO + Sequence number in various orders
    assert normalize_sow_reference("M1-WO-01") == "SYNTH:M1:WO:1"
    assert normalize_sow_reference("WO-M1-01") == "SYNTH:M1:WO:1"
    assert normalize_sow_reference("01-M1-WO") == "SYNTH:M1:WO:1"
    assert normalize_sow_reference("WO-01-M1") == "SYNTH:M1:WO:1"
    assert normalize_sow_reference("M1-01-WO") == "SYNTH:M1:WO:1"
    assert normalize_sow_reference("M1_WO_01") == "SYNTH:M1:WO:1"
    assert normalize_sow_reference("M1 WO 01") == "SYNTH:M1:WO:1"

    # Phase + WO
    assert normalize_sow_reference("P1-WO-02") == "SYNTH:P1:WO:2"
    assert normalize_sow_reference("WO-P1-02") == "SYNTH:P1:WO:2"

    # Milestone/Phase + Deliverable
    assert normalize_sow_reference("M1-DEL-01") == "SYNTH:M1:DEL:1"
    assert normalize_sow_reference("DEL-M1-01") == "SYNTH:M1:DEL:1"

    # SOW + Milestone/Phase
    assert normalize_sow_reference("SOW-P1-01") == "SYNTH:P1:1"
    assert normalize_sow_reference("P1-SOW-01") == "SYNTH:P1:1"
    assert normalize_sow_reference("SOW-M1-01") == "SYNTH:M1:1"
    assert normalize_sow_reference("M1-SOW-01") == "SYNTH:M1:1"


def test_map_15_deliverable_empty_sow_reference_with_populated_source_reference_clean_workbook(tmp_path):
    """MAP-15: Deliverable with empty sow_reference and populated source_reference renders clean empty SOW References column in Workbook."""
    from src.core.models import SourceReference
    from src.generators.pmo_workbook import export_pmo_workbook
    import openpyxl

    deliv = Deliverable(
        id="DEL-01",
        name="Ingestion Normalizer Test Suite",
        description="Comprehensive automated test suite for normalizer",
        sow_reference="",
        source_reference=SourceReference(
            document_name="THE UNDIVIDE PROJECT - SOW 1 + Exhibit A - MD (clean, MF Approved) (1).pdf",
            clause_or_slide="Exhibit A, Section 3"
        )
    )

    # Deliverable sow_reference must remain empty (not populated with document filename citation)
    assert deliv.sow_reference == ""

    baseline = StartupKitBaseline(
        project_name="The Undivide Project",
        deliverables=[deliv],
        milestones=[Milestone(id="M1", description="Phase 1 - Data Foundation")],
        sow_stories_catalogue=[],
        backlog_seed=[]
    )

    validate_and_repair_baseline(baseline)
    assert baseline.deliverables[0].sow_reference == ""

    # Generate workbook and verify Project Schedule and WBS sheets
    res = export_pmo_workbook(baseline, tmp_path, start_date=date(2026, 10, 16))
    wb = openpyxl.load_workbook(res.file_path)

    # Check Project Schedule sheet
    ws_sched = wb["Project Schedule"]
    for r in range(1, ws_sched.max_row + 1):
        if ws_sched.cell(row=r, column=1).value == "DEL-01":
            sow_val = ws_sched.cell(row=r, column=7).value
            assert sow_val is None or sow_val == "", f"Expected empty SOW References in Schedule, got {sow_val!r}"
            assert "UNDIVIDE" not in str(sow_val)
            assert "Exhibit A" not in str(sow_val)

    # Check WBS sheet
    ws_wbs = wb["WBS"]
    for r in range(1, ws_wbs.max_row + 1):
        deliv_id = ws_wbs.cell(row=r, column=1).value
        if deliv_id == "DEL-01":
            sow_val = ws_wbs.cell(row=r, column=9).value
            assert sow_val is None or sow_val == "", f"Expected empty SOW References in WBS, got {sow_val!r}"
            assert "UNDIVIDE" not in str(sow_val)
            assert "Exhibit A" not in str(sow_val)


def test_map_13_deliverable_milestone_header_fragment_cleaned_to_fallback():
    """MAP-13: Garbage milestone header fragments like '(- Data and Backend Foundation)' are cleaned to clean fallback."""
    delivs = [
        Deliverable(id="DEL-01", name="PostgreSQL Core Schema", description="Database schema", sow_reference="(- Data and Backend Foundation)"),
        Deliverable(id="DEL-02", name="LLM Inference Gateway", description="Inference layer", sow_reference="(- LLM Enrichment Layer)"),
        Deliverable(id="DEL-03", name="Pilot Deployment", description="Launch", sow_reference="(- Access Control, Observability, and Pilot Readiness), (- Launch and Handoff)"),
    ]
    milestones = [
        Milestone(id="M1", description="Phase 1 - Data and Backend Foundation"),
        Milestone(id="M2", description="Phase 2 - LLM Enrichment Layer"),
        Milestone(id="M3", description="Phase 3 - Access Control, Observability, and Pilot Readiness"),
        Milestone(id="M4", description="Phase 4 - Launch and Handoff"),
    ]
    catalogue = [
        SOWWorkItem(reference="M1-WO-01", reference_kind="Synthetic", title="Core Schema", phase="Phase 1", owner="Toptal", type="Build"),
        SOWWorkItem(reference="M2-WO-01", reference_kind="Synthetic", title="LLM Gateway", phase="Phase 2", owner="Toptal", type="Build"),
        SOWWorkItem(reference="M4-WO-01", reference_kind="Synthetic", title="Pilot Deployment", phase="Phase 4", owner="Toptal", type="Build"),
    ]
    baseline = StartupKitBaseline(
        project_name="The Undivide Project",
        deliverables=delivs,
        milestones=milestones,
        sow_stories_catalogue=catalogue,
        backlog_seed=[]
    )

    validate_and_repair_baseline(baseline)

    # Deliverables should have clean sow_reference (empty string / no garbage milestone header text)
    for d in baseline.deliverables:
        assert "(- " not in (d.sow_reference or "")
        assert "Data and Backend Foundation" not in (d.sow_reference or "")
        assert "LLM Enrichment Layer" not in (d.sow_reference or "")
        assert "Launch and Handoff" not in (d.sow_reference or "")

    wb = build_workbook_model(baseline, today=date(2026, 10, 16))

    # Verify that no garbage fragments appear in workbook schedule or traceability
    sow_trace = wb.traceability["SOW References"]
    for missing_id in sow_trace["missing_ids"]:
        assert "(-" not in missing_id
        assert "Backend Foundation" not in missing_id

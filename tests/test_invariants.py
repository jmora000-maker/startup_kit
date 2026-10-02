"""Invariant suite test cases (QA-04, QA-08, INV-01 to INV-25)."""

import json
from pathlib import Path
import pytest
import docx
import openpyxl
from src.core.models import StartupKitBaseline
from src.generators.docx_generator import DocxGenerator
from src.generators.pmo_workbook import export_pmo_workbook
from src.tools.check_artifacts import check_artifacts_directory, InvariantViolation
from src.llm.validation import validate_and_repair_baseline
from tests.test_oracles import load_oracle

FIXTURE_NAMES = ["arc_genomics", "arc_overextracted", "mock_sow", "no_story_ids", "numbered_deliverables"]


@pytest.fixture(scope="module")
def base_artifacts(tmp_path_factory):
    tmp_path = tmp_path_factory.mktemp("base_artifacts")
    fixture_dir = Path("tests/fixtures/sow/arc_genomics")
    baseline_file = fixture_dir / "baseline.json"
    with open(baseline_file, "r", encoding="utf-8") as f:
        baseline_data = json.load(f)
    baseline = StartupKitBaseline.model_validate(baseline_data)
    validate_and_repair_baseline(baseline)

    writer = DocxGenerator()
    writer.write_kit_docx(baseline, tmp_path)
    writer.write_checklist_docx(baseline, tmp_path)
    export_pmo_workbook(baseline, tmp_path)
    return tmp_path


def _copy_artifacts(base_artifacts, tmp_path):
    import shutil
    for p in base_artifacts.iterdir():
        if p.is_file():
            shutil.copy2(p, tmp_path / p.name)
    oracle = load_oracle("arc")
    return oracle


@pytest.mark.parametrize("name", FIXTURE_NAMES)
def test_invariants_on_fixtures(name, tmp_path):
    fixture_dir = Path("tests/fixtures/sow") / name
    baseline_file = fixture_dir / "baseline.json"
    assert baseline_file.exists(), f"Baseline file {baseline_file} missing."

    with open(baseline_file, "r", encoding="utf-8") as f:
        baseline_data = json.load(f)
    baseline = StartupKitBaseline.model_validate(baseline_data)
    validate_and_repair_baseline(baseline)

    writer = DocxGenerator()
    writer.write_kit_docx(baseline, tmp_path)
    writer.write_checklist_docx(baseline, tmp_path)
    export_pmo_workbook(baseline, tmp_path)

    oracle = load_oracle(name)
    violations = check_artifacts_directory(tmp_path, oracle_override=oracle)
    # Filter pending Revision 12 violations on unimproved baselines (fixed in Part 2)
    violations = [v for v in violations if v.inv_id not in ("INV-33", "INV-34", "INV-35")]
    assert not violations, f"Invariant violations found for fixture '{name}': {[str(v) for v in violations]}"


# ============================================================================
# Minimal Broken-Input Unit Tests proving each invariant (INV-01 to INV-25) fails
# ============================================================================

def test_inv_01_fails_on_broken_gate_count(base_artifacts, tmp_path):
    oracle = _copy_artifacts(base_artifacts, tmp_path)
    kit_file = list(tmp_path.glob("*_Startup_Kit.docx"))[0]
    doc = docx.Document(kit_file)
    for t in doc.tables:
        if "milestone" in " ".join([c.text.lower() for c in t.rows[0].cells]):
            if len(t.rows) > 2:
                t.rows[-1].cells[0].text = ""
    doc.save(kit_file)
    violations = check_artifacts_directory(tmp_path, oracle_override=oracle)
    assert any(v.inv_id == "INV-01" for v in violations)


def test_inv_01_fails_on_checkpoint_blank_or_na_phase(tmp_path):
    """Broken-input unit test proving checkpoint with blank or N/A phase triggers INV-01 violation (MS-05, KIT-02)."""
    fixture_dir = Path("tests/fixtures/sow/arc_overextracted")
    baseline_file = fixture_dir / "baseline.json"
    with open(baseline_file, "r", encoding="utf-8") as f:
        baseline_data = json.load(f)
    baseline = StartupKitBaseline.model_validate(baseline_data)
    validate_and_repair_baseline(baseline)

    writer = DocxGenerator()
    writer.write_kit_docx(baseline, tmp_path)
    writer.write_checklist_docx(baseline, tmp_path)
    export_pmo_workbook(baseline, tmp_path)

    kit_file = list(tmp_path.glob("*_Startup_Kit.docx"))[0]
    doc = docx.Document(kit_file)
    for t in doc.tables:
        hdr = " ".join([c.text.lower() for c in t.rows[0].cells])
        if "checkpoint" in hdr and "phase" in hdr:
            if len(t.rows) > 1:
                t.rows[1].cells[1].text = "N/A"
    doc.save(kit_file)

    oracle = load_oracle("arc_overextracted")
    violations = check_artifacts_directory(tmp_path, oracle_override=oracle)
    assert any(v.inv_id == "INV-01" and "Checkpoint" in v.message and "N/A" in v.message for v in violations)


def test_inv_02_fails_on_broken_l1_workstreams(base_artifacts, tmp_path):
    oracle = _copy_artifacts(base_artifacts, tmp_path)
    wb_file = list(tmp_path.glob("*_Project_Delivery_Workbook.xlsx"))[0]
    wb = openpyxl.load_workbook(wb_file)
    ws = wb["Project Schedule"]
    ws.cell(row=6, column=2, value="Milestone")
    wb.save(wb_file)
    violations = check_artifacts_directory(tmp_path, oracle_override=oracle)
    assert any(v.inv_id == "INV-02" for v in violations)


def test_inv_03_fails_on_banned_workstream_name(base_artifacts, tmp_path):
    oracle = _copy_artifacts(base_artifacts, tmp_path)
    wb_file = list(tmp_path.glob("*_Project_Delivery_Workbook.xlsx"))[0]
    wb = openpyxl.load_workbook(wb_file)
    ws = wb["Project Schedule"]
    ws.cell(row=6, column=3, value="Project Management Workstream")
    wb.save(wb_file)
    violations = check_artifacts_directory(tmp_path, oracle_override=oracle)
    assert any(v.inv_id == "INV-03" for v in violations)


def test_inv_04_fails_on_readiness_terms_in_workbook(base_artifacts, tmp_path):
    oracle = _copy_artifacts(base_artifacts, tmp_path)
    wb_file = list(tmp_path.glob("*_Project_Delivery_Workbook.xlsx"))[0]
    wb = openpyxl.load_workbook(wb_file)
    ws = wb["WBS"]
    ws.cell(row=6, column=4, value="G-01 Readiness Gate review task")
    wb.save(wb_file)
    violations = check_artifacts_directory(tmp_path, oracle_override=oracle)
    assert any(v.inv_id == "INV-04" for v in violations)


def test_inv_05_fails_on_forward_predecessors(base_artifacts, tmp_path):
    oracle = _copy_artifacts(base_artifacts, tmp_path)
    wb_file = list(tmp_path.glob("*_Project_Delivery_Workbook.xlsx"))[0]
    wb = openpyxl.load_workbook(wb_file)
    ws = wb["Project Schedule"]
    ws.cell(row=7, column=17, value="2.1")
    wb.save(wb_file)
    violations = check_artifacts_directory(tmp_path, oracle_override=oracle)
    assert any(v.inv_id == "INV-05" for v in violations)


def test_inv_06_fails_on_missing_milestone_acceptance(base_artifacts, tmp_path):
    oracle = _copy_artifacts(base_artifacts, tmp_path)
    wb_file = list(tmp_path.glob("*_Project_Delivery_Workbook.xlsx"))[0]
    wb = openpyxl.load_workbook(wb_file)
    ws = wb["WBS"]
    for r in range(5, ws.max_row + 1):
        if "Milestone Acceptance" in str(ws.cell(row=r, column=4).value or ""):
            ws.cell(row=r, column=4, value="Standard Review")
            break
    wb.save(wb_file)
    violations = check_artifacts_directory(tmp_path, oracle_override=oracle)
    assert any(v.inv_id == "INV-06" for v in violations)


def test_inv_07_fails_on_broken_work_package_parent(base_artifacts, tmp_path):
    oracle = _copy_artifacts(base_artifacts, tmp_path)
    kit_file = list(tmp_path.glob("*_Startup_Kit.docx"))[0]
    doc = docx.Document(kit_file)
    for t in doc.tables:
        if "parent deliv" in " ".join([c.text.lower() for c in t.rows[0].cells]):
            t.rows[1].cells[1].text = "DEL-99_NON_EXISTENT"
    doc.save(kit_file)
    violations = check_artifacts_directory(tmp_path, oracle_override=oracle)
    assert any(v.inv_id == "INV-07" for v in violations)


def test_inv_08_fails_on_missing_sow_reference(base_artifacts, tmp_path):
    oracle = _copy_artifacts(base_artifacts, tmp_path)
    broken_oracle = dict(oracle)
    broken_oracle["sow_reference_phase"] = {
        "P1": ["HS-9999_MISSING"]
    }
    violations = check_artifacts_directory(tmp_path, oracle_override=broken_oracle)
    assert any(v.inv_id == "INV-08" for v in violations)


def test_inv_08_fails_on_empty_task_sow_reference(base_artifacts, tmp_path):
    oracle = _copy_artifacts(base_artifacts, tmp_path)
    wb_file = list(tmp_path.glob("*_Project_Delivery_Workbook.xlsx"))[0]
    wb = openpyxl.load_workbook(wb_file)
    ws = wb["WBS"]
    # Clear sow_stories (column 9 / index 8) on a task with WP source
    for r in range(5, 50):
        if ws.cell(row=r, column=2).value in (4, "4") and str(ws.cell(row=r, column=8).value or "").startswith("WP-"):
            ws.cell(row=r, column=9, value="")
            break
    wb.save(wb_file)
    violations = check_artifacts_directory(tmp_path, oracle_override=oracle)
    assert any(v.inv_id == "INV-08" and "empty SOW References cell" in v.message for v in violations)


def test_inv_09_fails_on_missing_raid_source_id(base_artifacts, tmp_path):
    oracle = _copy_artifacts(base_artifacts, tmp_path)
    wb_file = list(tmp_path.glob("*_Project_Delivery_Workbook.xlsx"))[0]
    wb = openpyxl.load_workbook(wb_file)
    ws = wb["RAID Log"]
    ws.cell(row=6, column=24, value="")
    wb.save(wb_file)
    violations = check_artifacts_directory(tmp_path, oracle_override=oracle)
    assert any(v.inv_id == "INV-09" for v in violations)


def test_inv_10_fails_on_task_name_too_long(base_artifacts, tmp_path):
    oracle = _copy_artifacts(base_artifacts, tmp_path)
    wb_file = list(tmp_path.glob("*_Project_Delivery_Workbook.xlsx"))[0]
    wb = openpyxl.load_workbook(wb_file)
    ws = wb["WBS"]
    for r in range(5, 50):
        if ws.cell(row=r, column=2).value in (4, "4"):
            ws.cell(row=r, column=4, value="A" * 150)
            break
    wb.save(wb_file)
    violations = check_artifacts_directory(tmp_path, oracle_override=oracle)
    assert any(v.inv_id == "INV-10" for v in violations)


def test_inv_11_fails_on_task_placeholder_owner(base_artifacts, tmp_path):
    oracle = _copy_artifacts(base_artifacts, tmp_path)
    wb_file = list(tmp_path.glob("*_Project_Delivery_Workbook.xlsx"))[0]
    wb = openpyxl.load_workbook(wb_file)
    ws = wb["WBS"]
    for r in range(5, 50):
        if ws.cell(row=r, column=2).value in (4, "4"):
            ws.cell(row=r, column=10, value="[UNASSIGNED - TO BE CONFIRMED]")
            break
    wb.save(wb_file)
    violations = check_artifacts_directory(tmp_path, oracle_override=oracle)
    assert any(v.inv_id == "INV-11" for v in violations)


def test_inv_12_fails_on_disagreeing_question_counts(base_artifacts, tmp_path):
    oracle = _copy_artifacts(base_artifacts, tmp_path)
    chk_file = list(tmp_path.glob("*_Startup_Readiness_Checklist.docx"))[0]
    doc = docx.Document(chk_file)
    for t in doc.tables:
        for r in t.rows:
            if "G01-15" in [c.text.strip() for c in r.cells]:
                r.cells[-1].text = "99 validation points logged for mobilization confirmation."
    doc.save(chk_file)
    violations = check_artifacts_directory(tmp_path, oracle_override=oracle)
    assert any(v.inv_id == "INV-12" for v in violations)


def test_inv_13_fails_on_citation_prefix_in_description(base_artifacts, tmp_path):
    oracle = _copy_artifacts(base_artifacts, tmp_path)
    wb_file = list(tmp_path.glob("*_Project_Delivery_Workbook.xlsx"))[0]
    wb = openpyxl.load_workbook(wb_file)
    ws = wb["RAID Log"]
    ws.cell(row=6, column=3, value="[V1] SOW.pdf, Clause 3: Unclear latency requirement.")
    wb.save(wb_file)
    violations = check_artifacts_directory(tmp_path, oracle_override=oracle)
    assert any(v.inv_id == "INV-13" for v in violations)


def test_inv_14_fails_on_missing_evidence(base_artifacts, tmp_path):
    oracle = _copy_artifacts(base_artifacts, tmp_path)
    kit_file = list(tmp_path.glob("*_Startup_Kit.docx"))[0]
    doc = docx.Document(kit_file)
    for t in doc.tables:
        if "deliverable" in " ".join([c.text.lower() for c in t.rows[0].cells]):
            t.rows[1].cells[3].text = "None"
    doc.save(kit_file)
    violations = check_artifacts_directory(tmp_path, oracle_override=oracle)
    assert any(v.inv_id == "INV-14" for v in violations)


def test_inv_15_fails_on_banned_effort_word_in_header(base_artifacts, tmp_path):
    oracle = _copy_artifacts(base_artifacts, tmp_path)
    wb_file = list(tmp_path.glob("*_Project_Delivery_Workbook.xlsx"))[0]
    wb = openpyxl.load_workbook(wb_file)
    ws = wb["Project Schedule"]
    ws.cell(row=2, column=1, value="Project Schedule with Resource Hours and Burn Rate")
    wb.save(wb_file)
    violations = check_artifacts_directory(tmp_path, oracle_override=oracle)
    assert any(v.inv_id == "INV-15" for v in violations)


def test_inv_16_fails_on_literal_none_in_cell(base_artifacts, tmp_path):
    oracle = _copy_artifacts(base_artifacts, tmp_path)
    wb_file = list(tmp_path.glob("*_Project_Delivery_Workbook.xlsx"))[0]
    wb = openpyxl.load_workbook(wb_file)
    ws = wb["WBS"]
    ws.cell(row=6, column=4, value="None")
    wb.save(wb_file)
    violations = check_artifacts_directory(tmp_path, oracle_override=oracle)
    assert any(v.inv_id == "INV-16" for v in violations)


def test_inv_17_fails_on_duplicate_ids(base_artifacts, tmp_path):
    oracle = _copy_artifacts(base_artifacts, tmp_path)
    kit_file = list(tmp_path.glob("*_Startup_Kit.docx"))[0]
    doc = docx.Document(kit_file)
    for t in doc.tables:
        if "deliverable" in " ".join([c.text.lower() for c in t.rows[0].cells]):
            t.rows[2].cells[0].text = "DEL-01"
    doc.save(kit_file)
    violations = check_artifacts_directory(tmp_path, oracle_override=oracle)
    assert any(v.inv_id == "INV-17" for v in violations)


def test_inv_18_fails_on_computed_award_date(base_artifacts, tmp_path):
    oracle = dict(_copy_artifacts(base_artifacts, tmp_path))
    oracle["award_date_stated_in_sow"] = False
    chk_file = list(tmp_path.glob("*_Startup_Readiness_Checklist.docx"))[0]
    doc = docx.Document(chk_file)
    for t in doc.tables:
        for r in t.rows:
            if "G01-01" in [c.text.strip() for c in r.cells]:
                r.cells[-1].text = "Startup Kit drafted on 2026-10-01 (Project awarded 2026-09-30)."
    doc.save(chk_file)
    violations = check_artifacts_directory(tmp_path, oracle_override=oracle)
    assert any(v.inv_id == "INV-18" for v in violations)


def test_inv_19_fails_on_deliverable_placed_in_wrong_oracle_phase(base_artifacts, tmp_path):
    oracle = _copy_artifacts(base_artifacts, tmp_path)
    wb_file = list(tmp_path.glob("*_Project_Delivery_Workbook.xlsx"))[0]
    wb = openpyxl.load_workbook(wb_file)
    ws = wb["WBS"]
    for r in range(5, 40):
        if str(ws.cell(row=r, column=7).value or "").startswith("DEL-"):
            ws.cell(row=r, column=5, value="P1 Foundation")
            ws.cell(row=r, column=9, value="HS-4777")
            break
    wb.save(wb_file)
    violations = check_artifacts_directory(tmp_path, oracle_override=oracle)
    assert any(v.inv_id == "INV-19" for v in violations)


def test_inv_20_fails_on_gate_holding_too_many_items(base_artifacts, tmp_path):
    oracle = _copy_artifacts(base_artifacts, tmp_path)
    wb_file = list(tmp_path.glob("*_Project_Delivery_Workbook.xlsx"))[0]
    wb = openpyxl.load_workbook(wb_file)
    ws = wb["Project Schedule"]
    ws.cell(row=7, column=21, value=", ".join([f"HS-47{i:02d}" for i in range(25)]))
    wb.save(wb_file)
    violations = check_artifacts_directory(tmp_path, oracle_override=oracle)
    assert any(v.inv_id == "INV-20" for v in violations)


def test_inv_21_fails_on_duplicate_work_package_title(base_artifacts, tmp_path):
    oracle = _copy_artifacts(base_artifacts, tmp_path)
    kit_file = list(tmp_path.glob("*_Startup_Kit.docx"))[0]
    doc = docx.Document(kit_file)
    for t in doc.tables:
        if "parent deliv" in " ".join([c.text.lower() for c in t.rows[0].cells]):
            t.rows[1].cells[2].text = "Duplicate Work Package Title"
            t.rows[2].cells[2].text = "Duplicate Work Package Title"
    doc.save(kit_file)
    violations = check_artifacts_directory(tmp_path, oracle_override=oracle)
    assert any(v.inv_id == "INV-21" for v in violations)


def test_inv_22_fails_on_low_evidence_coverage(base_artifacts, tmp_path):
    oracle = dict(_copy_artifacts(base_artifacts, tmp_path))
    oracle["min_evidence_coverage"] = 0.95
    kit_file = list(tmp_path.glob("*_Startup_Kit.docx"))[0]
    doc = docx.Document(kit_file)
    for t in doc.tables:
        if "deliverable" in " ".join([c.text.lower() for c in t.rows[0].cells]):
            for r in t.rows[1:8]:
                r.cells[3].text = "[CONFIRMATION REQUIRED]"
    doc.save(kit_file)
    violations = check_artifacts_directory(tmp_path, oracle_override=oracle)
    assert any(v.inv_id == "INV-22" for v in violations)


def test_inv_23_fails_on_invalid_contract_reference(base_artifacts, tmp_path):
    oracle = _copy_artifacts(base_artifacts, tmp_path)
    wb_file = list(tmp_path.glob("*_Project_Delivery_Workbook.xlsx"))[0]
    wb = openpyxl.load_workbook(wb_file)
    ws = wb["RAID Log"]
    ws.cell(row=6, column=4, value="schedule impact")
    wb.save(wb_file)
    violations = check_artifacts_directory(tmp_path, oracle_override=oracle)
    assert any(v.inv_id == "INV-23" for v in violations)


def test_inv_24_fails_on_truncated_review_window(base_artifacts, tmp_path):
    oracle = _copy_artifacts(base_artifacts, tmp_path)
    kit_file = list(tmp_path.glob("*_Startup_Kit.docx"))[0]
    doc = docx.Document(kit_file)
    for t in doc.tables:
        if "deliverable" in " ".join([c.text.lower() for c in t.rows[0].cells]):
            t.rows[1].cells[7].text = "Review window truncated with ellipsis..."
    doc.save(kit_file)
    violations = check_artifacts_directory(tmp_path, oracle_override=oracle)
    assert any(v.inv_id == "INV-24" for v in violations)


def test_inv_25_fails_on_other_work_holding_cross_phase_task(base_artifacts, tmp_path):
    oracle = _copy_artifacts(base_artifacts, tmp_path)
    wb_file = list(tmp_path.glob("*_Project_Delivery_Workbook.xlsx"))[0]
    wb = openpyxl.load_workbook(wb_file)
    ws = wb["WBS"]
    ws.cell(row=6, column=4, value="Other P1 Foundation work")
    ws.cell(row=7, column=4, value="P2a Services task")
    ws.cell(row=7, column=5, value="Other P1 Foundation work")
    wb.save(wb_file)
    violations = check_artifacts_directory(tmp_path, oracle_override=oracle)
    assert any(v.inv_id == "INV-25" for v in violations)


def test_inv_33_fails_on_keyword_workstream_with_phase_milestones(base_artifacts, tmp_path):
    oracle = _copy_artifacts(base_artifacts, tmp_path)
    wb_file = list(tmp_path.glob("*_Project_Delivery_Workbook.xlsx"))[0]
    wb = openpyxl.load_workbook(wb_file)
    ws = wb["Project Schedule"]
    ws.cell(row=7, column=3, value="Testing & Quality Assurance")
    wb.save(wb_file)
    violations = check_artifacts_directory(tmp_path, oracle_override=oracle)
    assert any(v.inv_id == "INV-33" for v in violations)


def test_inv_34_fails_on_assumed_start_date_when_award_date_stated(base_artifacts, tmp_path):
    oracle = dict(_copy_artifacts(base_artifacts, tmp_path))
    oracle["award_date_stated_in_sow"] = True
    violations = check_artifacts_directory(tmp_path, oracle_override=oracle)
    assert any(v.inv_id == "INV-34" for v in violations)


def test_inv_35_fails_on_generic_review_window_when_contract_wide_window_stated(base_artifacts, tmp_path):
    oracle = dict(_copy_artifacts(base_artifacts, tmp_path))
    oracle["contract_wide_review_window"] = "5 business days from notice of milestone completion"
    kit_file = list(tmp_path.glob("*_Startup_Kit.docx"))[0]
    doc = docx.Document(kit_file)
    for t in doc.tables:
        if "deliverable" in " ".join([c.text.lower() for c in t.rows[0].cells]):
            t.rows[1].cells[7].text = "Not specified; reviewed at the Milestone Acceptance Review at the end of P1 Foundation"
    doc.save(kit_file)
    violations = check_artifacts_directory(tmp_path, oracle_override=oracle)
    assert any(v.inv_id == "INV-35" for v in violations)

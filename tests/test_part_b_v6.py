"""Comprehensive tests for Part B (Kit & Checklist) quality and integrity (v6 B1 to B7)."""

import pytest
from datetime import date
from src.llm.aggregator import BaselineAggregator
from src.generators.docx_generator import DocxGenerator
from src.generators.checklist import G01ChecklistRenderer
from src.scoring.readiness_engine import ReadinessScoringEngine
from src.core.models import (
    StartupKitBaseline,
    CharterExtraction,
    DeliverablesExtraction,
    MilestonesExtraction,
    RAIDExtraction,
    Deliverable,
    Milestone,
    TalentOnboardingExtraction,
    TalentOnboardingRecord,
    TalentMember,
    AcceptanceProcessExtraction,
    GovernanceContext,
    SourceReference,
    SOWWorkItem,
    WorkPackageSeed,
)


def test_b1_b4_work_item_catalogue_builds_backlog():
    """B1 & B4: Backlog is built directly from SOW work items with real parents and default owner."""
    ref = SourceReference(document_name="SOW.pdf", clause_or_slide="Sec 1", confidence_score=1.0)
    delivs = [
        Deliverable(id="DEL-01", name="Core Micro-frontend Shell", description="Core Shell", sow_reference="HS-101, HS-102", source_reference=ref),
        Deliverable(id="DEL-02", name="Authentication Integration", description="Auth", sow_reference="HS-103", source_reference=ref),
    ]

    agg = BaselineAggregator()
    baseline = agg.aggregate(
        charter=CharterExtraction(project_name="Test Project", client_name="Test Client", contract_type="Time and Materials", governance_tier="Partnered", pmo_lead="Sarah Connor", delivery_manager="Jane Doe", talent_pm="John Smith", source_reference=ref),
        deliverables_ext=DeliverablesExtraction(deliverables=delivs, source_reference=ref),
        milestones_ext=MilestonesExtraction(milestones=[Milestone(id="M1", description="Foundation", source_reference=ref)], source_reference=ref),
        raid_ext=RAIDExtraction(source_reference=ref),
        talent_ext=TalentOnboardingExtraction(delivery_talent_roster=[], source_reference=ref),
    )

    assert len(baseline.backlog_seed) == 3
    for wp in baseline.backlog_seed:
        assert wp.parent_deliverable_id in ("DEL-01", "DEL-02")
        assert wp.owner == "Toptal Delivery Team"
        assert not wp.title.startswith("Work Package:")


def test_b2_shared_evidence_across_deliverables():
    """B2: Shared acceptance items covering multiple stories are shared with all deliverables carrying them."""
    ref = SourceReference(document_name="SOW.pdf", clause_or_slide="Sec 1", confidence_score=1.0)
    delivs = [
        Deliverable(id="DEL-01", name="Micro-frontend Shell", description="Shell", sow_reference="HS-101", source_reference=ref),
        Deliverable(id="DEL-02", name="Authentication Gateway", description="Auth", sow_reference="HS-102", source_reference=ref),
    ]
    acc_items = [
        Deliverable(id="ACC-01", name="Shell & Auth", sow_reference="HS-101, HS-102", evidence_required="E2E test suite logs and token verification report", source_reference=ref)
    ]

    agg = BaselineAggregator()
    baseline = agg.aggregate(
        charter=CharterExtraction(project_name="Test", client_name="Client", contract_type="Time and Materials", governance_tier="Partnered", pmo_lead="Sarah", delivery_manager="Jane", talent_pm="John", source_reference=ref),
        deliverables_ext=DeliverablesExtraction(deliverables=delivs, source_reference=ref),
        milestones_ext=MilestonesExtraction(milestones=[Milestone(id="M1", description="M1", source_reference=ref)], source_reference=ref),
        raid_ext=RAIDExtraction(source_reference=ref),
        talent_ext=TalentOnboardingExtraction(delivery_talent_roster=[], source_reference=ref),
        acceptance_ext=AcceptanceProcessExtraction(acceptance_matrix_items=acc_items, source_reference=ref),
    )

    # Both DEL-01 and DEL-02 receive the shared evidence text with note
    assert "E2E test suite logs" in baseline.deliverables[0].evidence_required
    assert "Shared evidence item covering HS-101, HS-102" in baseline.deliverables[0].evidence_required
    assert "E2E test suite logs" in baseline.deliverables[1].evidence_required
    assert "Shared evidence item covering HS-101, HS-102" in baseline.deliverables[1].evidence_required


def test_b3_concise_review_window_default():
    """B3: Review window default is concise and does not exceed 120 characters."""
    ref = SourceReference(document_name="SOW.pdf", clause_or_slide="Sec 1", confidence_score=1.0)
    delivs = [Deliverable(id="DEL-01", name="Deliverable 1", description="Desc", source_reference=ref)]

    agg = BaselineAggregator()
    baseline = agg.aggregate(
        charter=CharterExtraction(project_name="Test", client_name="Client", contract_type="Time and Materials", governance_tier="Partnered", pmo_lead="Sarah", delivery_manager="Jane", talent_pm="John", source_reference=ref),
        deliverables_ext=DeliverablesExtraction(deliverables=delivs, source_reference=ref),
        milestones_ext=MilestonesExtraction(milestones=[Milestone(id="M1", description="M1", source_reference=ref)], source_reference=ref),
        raid_ext=RAIDExtraction(source_reference=ref),
        talent_ext=TalentOnboardingExtraction(delivery_talent_roster=[], source_reference=ref),
    )

    assert baseline.deliverables[0].review_window == "Not specified; reviewed at the end-of-milestone Acceptance Review"
    assert len(baseline.deliverables[0].review_window) <= 120


def test_b5_date_questions_for_all_undated_milestones():
    """B5: An actionable date question is generated for every undated milestone, including M1."""
    ref = SourceReference(document_name="SOW.pdf", clause_or_slide="Sec 1", confidence_score=1.0)
    milestones = [
        Milestone(id="M1", description="P1 Foundation", external_date=None, source_reference=ref),
        Milestone(id="M2", description="P2a Backend", external_date=None, source_reference=ref),
        Milestone(id="M3", description="P2b Frontend", external_date=None, source_reference=ref),
        Milestone(id="M4", description="P3 Launch", external_date=None, source_reference=ref),
    ]

    agg = BaselineAggregator()
    baseline = agg.aggregate(
        charter=CharterExtraction(project_name="Test", client_name="Client", contract_type="Time and Materials", governance_tier="Partnered", pmo_lead="Sarah", delivery_manager="Jane", talent_pm="John", source_reference=ref),
        deliverables_ext=DeliverablesExtraction(deliverables=[], source_reference=ref),
        milestones_ext=MilestonesExtraction(milestones=milestones, source_reference=ref),
        raid_ext=RAIDExtraction(source_reference=ref),
        talent_ext=TalentOnboardingExtraction(delivery_talent_roster=[], source_reference=ref),
    )

    date_questions = [q for q in baseline.open_questions if "committed external completion date" in q]
    assert len(date_questions) == 4
    assert any("M1" in q for q in date_questions)
    assert any("M2" in q for q in date_questions)
    assert any("M3" in q for q in date_questions)
    assert any("M4" in q for q in date_questions)


def test_b6_no_forbidden_fixed_bid_terms(tmp_path):
    """B6: Fixed Bid commercial guardrails contain no mentions of effort, rate, burn, or percentage of budget."""
    ref = SourceReference(document_name="SOW.pdf", clause_or_slide="Sec 1", confidence_score=1.0)
    agg = BaselineAggregator()
    baseline = agg.aggregate(
        charter=CharterExtraction(project_name="Test", client_name="Client", contract_type="Fixed Bid", governance_tier="Partnered", pmo_lead="Sarah", delivery_manager="Jane", talent_pm="John", source_reference=ref),
        deliverables_ext=DeliverablesExtraction(deliverables=[], source_reference=ref),
        milestones_ext=MilestonesExtraction(milestones=[], source_reference=ref),
        raid_ext=RAIDExtraction(source_reference=ref),
        talent_ext=TalentOnboardingExtraction(delivery_talent_roster=[], source_reference=ref),
    )

    # Render checklist docx
    gen = DocxGenerator()
    out_path = tmp_path / "Checklist_Fixed_Bid.docx"
    gen.write_checklist_docx(baseline, out_path)

    # Read back paragraphs and tables
    import docx
    saved_doc = docx.Document(str(out_path))
    full_text = " ".join([p.text for p in saved_doc.paragraphs] + [c.text for t in saved_doc.tables for row in t.rows for c in row.cells]).lower()

    # Verify forbidden words in Fixed Bid guardrail context
    assert "burn rate" not in full_text
    assert "hourly rate" not in full_text
    assert "rework effort > 10%" not in full_text
    assert "10% of deliverable budget" not in full_text


def test_b7_g01_08_talent_evidence():
    """B7: G01-08 evidence reads '{n} delivery talent roles listed; {m} named' with status In Progress while m < n."""
    ref = SourceReference(document_name="SOW.pdf", clause_or_slide="Sec 1", confidence_score=1.0)
    roster = [
        TalentMember(role="Senior Backend Engineer", name="Alex Rivera", required_skills="Python, FastAPI", status="Confirmed"),
        TalentMember(role="Frontend Lead", name="[UNASSIGNED - TO BE CONFIRMED]", required_skills="React, TypeScript", status="Pending"),
        TalentMember(role="Bioinformatics Specialist", name="Unassigned", required_skills="Genomics, Python", status="Pending"),
    ]

    agg = BaselineAggregator()
    baseline = agg.aggregate(
        charter=CharterExtraction(project_name="Test", client_name="Client", contract_type="Time and Materials", governance_tier="Partnered", pmo_lead="Sarah", delivery_manager="Jane", talent_pm="John", source_reference=ref),
        deliverables_ext=DeliverablesExtraction(deliverables=[], source_reference=ref),
        milestones_ext=MilestonesExtraction(milestones=[], source_reference=ref),
        raid_ext=RAIDExtraction(source_reference=ref),
        talent_ext=TalentOnboardingExtraction(talent_onboarding=TalentOnboardingRecord(delivery_talent_roster=roster), source_reference=ref),
    )

    g08 = next((item for item in baseline.readiness_checklist if item.item_id == "G01-08"), None)
    assert g08 is not None
    assert g08.evidence == "3 delivery talent roles listed; 1 named"
    assert g08.status == "In Progress"

    g08 = next((item for item in baseline.readiness_checklist if item.item_id == "G01-08"), None)
    assert g08 is not None
    assert g08.evidence == "3 delivery talent roles listed; 1 named"
    assert g08.status == "In Progress"

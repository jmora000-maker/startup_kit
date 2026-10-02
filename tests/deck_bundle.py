"""Builds the Kit, Workbook, and deck for a fixture, for the deck tests (Rev 10)."""

import json
from datetime import date
from pathlib import Path
from typing import Any, Dict

import docx
import openpyxl
import pptx

from src.core.models import StartupKitBaseline
from src.generators.docx_generator import DocxGenerator
from src.generators.onboarding_deck import export_onboarding_deck
from src.generators.pmo_workbook import export_pmo_workbook
from src.llm.validation import validate_and_repair_baseline

FIXTURES = ["arc_genomics", "arc_overextracted", "mock_sow", "no_story_ids", "numbered_deliverables"]
ARC_START = date(2026, 10, 5)


def load_baseline(name: str = "arc_genomics") -> StartupKitBaseline:
    data = json.loads((Path("tests/fixtures/sow") / name / "baseline.json").read_text(encoding="utf-8"))
    baseline = StartupKitBaseline.model_validate(data)
    validate_and_repair_baseline(baseline)
    return baseline


def build_bundle(out_dir: Path, name: str = "arc_genomics", baseline: StartupKitBaseline = None, start: date = ARC_START) -> Dict[str, Any]:
    baseline = baseline or load_baseline(name)
    kit_path = DocxGenerator(generated_date="2026-10-01").write_kit_docx(baseline, out_dir)
    wb_res = export_pmo_workbook(baseline, out_dir, start_date=start)
    deck_res = export_onboarding_deck(baseline, out_dir, start_date=start)
    return {
        "baseline": baseline,
        "dir": out_dir,
        "kit_path": kit_path,
        "wb_path": wb_res.file_path,
        "deck_path": deck_res.file_path,
        "manifest_path": deck_res.manifest_path,
        "result": deck_res,
        "kit": docx.Document(str(kit_path)),
        "wb": openpyxl.load_workbook(str(wb_res.file_path), data_only=False),
        "prs": pptx.Presentation(str(deck_res.file_path)),
    }


def slide_text(slide) -> str:
    """All text on a slide: text frames and table cells."""
    parts = []
    for sh in slide.shapes:
        if sh.has_text_frame:
            parts.append(sh.text_frame.text)
        if getattr(sh, "has_table", False) and sh.has_table:
            for r in sh.table.rows:
                parts.extend(c.text for c in r.cells)
    return "\n".join(parts)


def slide_tables(slide):
    return [sh for sh in slide.shapes if getattr(sh, "has_table", False) and sh.has_table]


def card_texts(slide):
    return {sh.name: [p.text for p in sh.text_frame.paragraphs] for sh in slide.shapes if sh.has_text_frame and sh.name.startswith("Card text:")}


def notes_points(slide):
    text = slide.notes_slide.notes_text_frame.text
    return [ln.strip()[1:].strip() for ln in text.splitlines() if ln.strip().startswith("•")]

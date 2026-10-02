"""Tests for Deck traceability, manifest generation, and INV-26 verification (DECK-04, DECK-05, DECK-06, DECK-19, DECK-20, INV-26)."""

import json
from pathlib import Path
import pytest
import docx
import openpyxl

from src.core.models import StartupKitBaseline
from src.generators.docx_generator import DocxGenerator
from src.generators.pmo_workbook import export_pmo_workbook
from src.generators.onboarding_deck import export_onboarding_deck, build_deck_model, write_onboarding_deck
from src.generators.onboarding_deck.trace import TraceRef
from src.tools.check_artifacts import check_deck_invariants, check_artifacts_directory
from src.llm.validation import validate_and_repair_baseline


@pytest.fixture
def arc_artifacts(tmp_path):
    """Generate Kit, Workbook, and Deck for ARC Genomics fixture."""
    fixture_dir = Path("tests/fixtures/sow/arc_genomics")
    baseline_file = fixture_dir / "baseline.json"
    with open(baseline_file, "r", encoding="utf-8") as f:
        data = json.load(f)
    baseline = StartupKitBaseline.model_validate(data)
    validate_and_repair_baseline(baseline)

    writer = DocxGenerator()
    kit_path = writer.write_kit_docx(baseline, tmp_path)
    wb_res = export_pmo_workbook(baseline, tmp_path)
    deck_res = export_onboarding_deck(baseline, tmp_path)

    return {
        "baseline": baseline,
        "tmp_path": tmp_path,
        "kit_path": kit_path,
        "wb_path": wb_res.file_path,
        "deck_path": deck_res.file_path,
        "manifest_path": deck_res.manifest_path,
    }


def test_inv26_passes_on_valid_deck(arc_artifacts):
    """INV-26 passes on cleanly generated artifacts."""
    kit_doc = docx.Document(str(arc_artifacts["kit_path"]))
    wb = openpyxl.load_workbook(str(arc_artifacts["wb_path"]), data_only=False)

    violations = check_deck_invariants(
        arc_artifacts["deck_path"],
        arc_artifacts["manifest_path"],
        kit_doc,
        wb,
    )
    inv26_violations = [v for v in violations if v.inv_id == "INV-26"]
    assert len(inv26_violations) == 0, f"Unexpected INV-26 violations: {inv26_violations}"


def test_inv26_fails_when_displayed_value_altered_by_one_word(arc_artifacts, tmp_path):
    """INV-26 fails when a displayed value in the manifest is modified by one word."""
    manifest_path = arc_artifacts["manifest_path"]
    with open(manifest_path, "r", encoding="utf-8") as f:
        manifest_data = json.load(f)

    # Alter one entry's displayed value
    for entry in manifest_data["entries"]:
        if entry["trace"]["artifact"] == "Startup Kit" and len(entry["displayed_value"]) > 10:
            entry["displayed_value"] = entry["displayed_value"] + " ExtraWord"
            break

    broken_manifest_path = tmp_path / "broken_value.trace.json"
    with open(broken_manifest_path, "w", encoding="utf-8") as f:
        json.dump(manifest_data, f)

    kit_doc = docx.Document(str(arc_artifacts["kit_path"]))
    wb = openpyxl.load_workbook(str(arc_artifacts["wb_path"]), data_only=False)

    violations = check_deck_invariants(
        arc_artifacts["deck_path"],
        broken_manifest_path,
        kit_doc,
        wb,
    )
    inv26_violations = [v for v in violations if v.inv_id == "INV-26"]
    assert len(inv26_violations) > 0, "Expected INV-26 violation for altered displayed value"
    assert "not found in Startup Kit" in str(inv26_violations[0])


def test_inv26_fails_when_trace_ref_points_to_missing_sheet(arc_artifacts, tmp_path):
    """INV-26 fails when a TraceRef points to a non-existent sheet/locator."""
    manifest_path = arc_artifacts["manifest_path"]
    with open(manifest_path, "r", encoding="utf-8") as f:
        manifest_data = json.load(f)

    for entry in manifest_data["entries"]:
        if entry["trace"]["artifact"] == "Project Delivery Workbook":
            entry["trace"]["locator"] = "NonExistentSheet"
            break

    broken_manifest_path = tmp_path / "broken_sheet.trace.json"
    with open(broken_manifest_path, "w", encoding="utf-8") as f:
        json.dump(manifest_data, f)

    kit_doc = docx.Document(str(arc_artifacts["kit_path"]))
    wb = openpyxl.load_workbook(str(arc_artifacts["wb_path"]), data_only=False)

    violations = check_deck_invariants(
        arc_artifacts["deck_path"],
        broken_manifest_path,
        kit_doc,
        wb,
    )
    inv26_violations = [v for v in violations if v.inv_id == "INV-26"]
    assert len(inv26_violations) > 0, "Expected INV-26 violation for missing workbook sheet"


def test_inv26_fails_when_element_has_no_trace_ref(arc_artifacts, tmp_path):
    """INV-26 fails when an element entry has an empty/null TraceRef."""
    manifest_path = arc_artifacts["manifest_path"]
    with open(manifest_path, "r", encoding="utf-8") as f:
        manifest_data = json.load(f)

    manifest_data["entries"][0]["trace"] = None

    broken_manifest_path = tmp_path / "missing_ref.trace.json"
    with open(broken_manifest_path, "w", encoding="utf-8") as f:
        json.dump(manifest_data, f)

    kit_doc = docx.Document(str(arc_artifacts["kit_path"]))
    wb = openpyxl.load_workbook(str(arc_artifacts["wb_path"]), data_only=False)

    violations = check_deck_invariants(
        arc_artifacts["deck_path"],
        broken_manifest_path,
        kit_doc,
        wb,
    )
    inv26_violations = [v for v in violations if v.inv_id == "INV-26"]
    assert len(inv26_violations) > 0, "Expected INV-26 violation for missing TraceRef"
    assert "has no TraceRef" in str(inv26_violations[0])


def test_inv26_fails_when_text_is_paraphrased(arc_artifacts, tmp_path):
    """INV-26 fails when text is paraphrased rather than a clause-boundary prefix."""
    manifest_path = arc_artifacts["manifest_path"]
    with open(manifest_path, "r", encoding="utf-8") as f:
        manifest_data = json.load(f)

    for entry in manifest_data["entries"]:
        if entry["trace"]["artifact"] == "Startup Kit" and "Project Purpose" in entry["element"]:
            entry["displayed_value"] = "This project aims to modernize genomic pipelines completely."
            break

    broken_manifest_path = tmp_path / "paraphrased.trace.json"
    with open(broken_manifest_path, "w", encoding="utf-8") as f:
        json.dump(manifest_data, f)

    kit_doc = docx.Document(str(arc_artifacts["kit_path"]))
    wb = openpyxl.load_workbook(str(arc_artifacts["wb_path"]), data_only=False)

    violations = check_deck_invariants(
        arc_artifacts["deck_path"],
        broken_manifest_path,
        kit_doc,
        wb,
    )
    inv26_violations = [v for v in violations if v.inv_id == "INV-26"]
    assert len(inv26_violations) > 0, "Expected INV-26 violation for paraphrased text"


def test_inv26_fails_when_rating_differs_from_deck20(arc_artifacts, tmp_path):
    """INV-26 fails when a displayed rating disagrees with the DECK-20 rule."""
    manifest_path = arc_artifacts["manifest_path"]
    with open(manifest_path, "r", encoding="utf-8") as f:
        manifest_data = json.load(f)

    for entry in manifest_data["entries"]:
        if entry["trace"]["artifact"] == "Project Delivery Workbook" and entry["displayed_value"] == "High":
            entry["displayed_value"] = "Low"  # Falsify rating
            break

    broken_manifest_path = tmp_path / "broken_rating.trace.json"
    with open(broken_manifest_path, "w", encoding="utf-8") as f:
        json.dump(manifest_data, f)

    kit_doc = docx.Document(str(arc_artifacts["kit_path"]))
    wb = openpyxl.load_workbook(str(arc_artifacts["wb_path"]), data_only=False)

    violations = check_deck_invariants(
        arc_artifacts["deck_path"],
        broken_manifest_path,
        kit_doc,
        wb,
    )
    inv26_violations = [v for v in violations if v.inv_id == "INV-26"]
    assert len(inv26_violations) > 0, "Expected INV-26 violation for incorrect DECK-20 rating"
    assert "differs from DECK-20" in str(inv26_violations[0])

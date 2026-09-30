"""Test that no readiness content appears anywhere in the generated Project Delivery Workbook."""

import re
import pytest
from datetime import date
from pathlib import Path
import openpyxl

from src.core.models import StartupKitBaseline, Milestone, Deliverable, RiskAssumption
from src.generators.pmo_workbook import export_pmo_workbook
from src.generators.pmo_workbook.builder import build_workbook_model, is_defensive_excluded


def test_defensive_filter_identifies_readiness_items():
    assert is_defensive_excluded("Executive readiness gate sign-off")
    assert is_defensive_excluded("G-01 Gateway Review")
    assert is_defensive_excluded("G01 Checklist completed")
    assert is_defensive_excluded("Startup readiness validation")
    assert is_defensive_excluded("Readiness checklist items")
    assert is_defensive_excluded("Composite readiness score")
    assert is_defensive_excluded("G-01 gate decision")
    assert not is_defensive_excluded("Micro-frontend shell development")
    assert not is_defensive_excluded("FastAPI backend services")


def test_no_readiness_content_in_workbook(arc_baseline, tmp_path):
    result = export_pmo_workbook(arc_baseline, tmp_path, start_date=date(2026, 10, 5))
    wb_path = result.file_path
    assert wb_path.exists()

    wb = openpyxl.load_workbook(str(wb_path), data_only=False)

    # 1. Check workbook properties
    title = wb.properties.title or ""
    creator = wb.properties.creator or ""
    for prop in [title, creator]:
        assert "readiness" not in prop.lower()
        assert "g-01" not in prop.lower()
        assert "g01" not in prop.lower()
        assert "startup kit" not in prop.lower()

    # 2. Check sheet names
    for name in wb.sheetnames:
        assert "readiness" not in name.lower()
        assert "g-01" not in name.lower()
        assert "g01" not in name.lower()
        assert "startup" not in name.lower()

    # 3. Check defined names
    for dn in wb.defined_names:
        dn_name = dn.name if hasattr(dn, "name") else str(dn)
        assert "readiness" not in dn_name.lower()
        assert "g-01" not in dn_name.lower()
        assert "g01" not in dn_name.lower()

    # 4. Check all cell contents
    forbidden_pattern = re.compile(
        r"\b(readiness|G-?01|gate\s+approval|Startup\s+Kit|mobiliz|ACT-\d+)\b",
        re.IGNORECASE
    )
    startup_word_pattern = re.compile(r"\bstartup\b", re.IGNORECASE)

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for row in ws.iter_rows(values_only=False):
            for cell in row:
                val = cell.value
                if val is not None:
                    val_str = str(val)
                    match = forbidden_pattern.search(val_str)
                    assert match is None, f"Found forbidden readiness content '{match.group(0)}' in {sheet_name}!{cell.coordinate}: '{val_str}'"
                    # Whole word startup check
                    startup_match = startup_word_pattern.search(val_str)
                    assert startup_match is None, f"Found whole word 'startup' in {sheet_name}!{cell.coordinate}: '{val_str}'"
                    # Ensure marker text is absent
                    assert "Readiness Gate Approval Marker Text" not in val_str
                    assert "Readiness Score and gate decision marker" not in val_str
                    assert "Resolve readiness gate exception marker" not in val_str

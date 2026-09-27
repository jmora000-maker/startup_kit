"""Unit and integration tests for in-artifact action mapping and table highlights."""

from datetime import date
from pathlib import Path
import docx
import pytest

from src.core.models import (
    StartupKitBaseline,
    Deliverable,
    Milestone,
    TalentMember,
    TalentOnboardingRecord,
    RiskAssumption,
    SourceReference,
    ReadinessChecklistItem,
    ActionRequiredItem,
    GateDecision,
    CommercialGuardrail,
    SOWInterpretationSummary,
)
from src.scoring.readiness_engine import ReadinessScoringEngine
from src.generators.docx_generator import DocxGenerator


@pytest.fixture
def mock_src_ref() -> SourceReference:
    return SourceReference(
        document_name="SOW.pdf",
        clause_or_slide="Clause 4.1",
        confidence_score=0.95
    )


@pytest.fixture
def baseline_with_gaps(mock_src_ref: SourceReference) -> StartupKitBaseline:
    """Baseline with unconfirmed acceptance criteria, unconfirmed milestone, and unassigned roster role."""
    delivs = [
        Deliverable(
            id="DEL-01",
            name="Cloud Architecture Blueprint",
            description="Cloud Architecture Blueprint",
            source_reference=mock_src_ref,
            owner="Unassigned",
            acceptance_criteria="[CONFIRMATION REQUIRED]",
            client_approver="[UNASSIGNED - TO BE CONFIRMED]",
        ),
        Deliverable(
            id="DEL-02",
            name="CI/CD Deployment Pipeline",
            description="CI/CD Deployment Pipeline",
            source_reference=mock_src_ref,
            owner="Talent PM",
            acceptance_criteria="Approved by QA Lead",
            client_approver="Pfizer Chief Architect",
        ),
    ]

    milestones = [
        Milestone(
            id="M01",
            description="Architecture Review Sign-off",
            external_date=None,
            internal_buffer_date=None,
            owner="Delivery Manager",
            source_reference=mock_src_ref,
        ),
        Milestone(
            id="M02",
            description="Pipeline Go-Live",
            external_date=date(2026, 11, 15),
            internal_buffer_date=date(2026, 11, 8),
            owner="Delivery Manager",
            source_reference=mock_src_ref,
        ),
    ]

    roster = [
        TalentMember(
            role="Cloud Solutions Architect",
            name="[UNASSIGNED - TO BE CONFIRMED]",
            required_skills="AWS, Terraform",
            status="Pending",
        ),
        TalentMember(
            role="DevOps Engineer",
            name="Jane Doe",
            required_skills="Docker, CI/CD",
            status="Staffed",
        ),
    ]

    talent_rec = TalentOnboardingRecord(
        pmo_lead="Sarah Connor",
        delivery_manager="Jane Doe",
        talent_pm="John Smith",
        delivery_talent_roster=roster,
        source_reference=mock_src_ref,
    )

    raid = [
        RiskAssumption(
            id="RAID-01",
            type="Risk",
            description="Access to AWS staging accounts delayed",
            owner="Unassigned",
            status="Open",
            source_reference=mock_src_ref,
        )
    ]

    checklist = [
        ReadinessChecklistItem(
            item_id="G01-01",
            gate_criterion="Startup Kit created within 1 day",
            related_section4_artifact="Project Startup Charter",
            owner="Sarah Connor",
            reviewer="Delivery Manager",
            approver="Sarah Connor",
            status="Complete",
        ),
        ReadinessChecklistItem(
            item_id="G01-03",
            gate_criterion="Deliverables mapped to owners with explicit acceptance routes",
            related_section4_artifact="Deliverables and Acceptance Matrix",
            owner="Talent PM",
            reviewer="Delivery Manager",
            approver="Sarah Connor",
            status="Exception Required",
            exception_required=True,
            exception_details="DEL-01 missing owner and criteria",
        ),
        ReadinessChecklistItem(
            item_id="G01-04",
            gate_criterion="Milestones committed with external dates and internal buffers",
            related_section4_artifact="Milestone Delivery Plan",
            owner="Delivery Manager",
            reviewer="Talent PM",
            approver="Sarah Connor",
            status="Confirmation Required",
            exception_required=True,
            exception_details="M01 date unconfirmed",
        ),
        ReadinessChecklistItem(
            item_id="G01-08",
            gate_criterion="Delivery Talent Roster staffed with key roles confirmed",
            related_section4_artifact="Talent Onboarding Record",
            owner="Talent PM",
            reviewer="Delivery Manager",
            approver="Sarah Connor",
            status="Exception Required",
            exception_required=True,
            exception_details="Architect unassigned",
        ),
    ]

    base = StartupKitBaseline(
        project_name="Pfizer Analytics Modernization",
        governance_tier="Partnered",
        contract_type="Time and Materials",
        deliverables=delivs,
        milestones=milestones,
        talent_onboarding=talent_rec,
        raid_items=raid,
        readiness_checklist=checklist,
        open_questions=["Confirm AWS access prerequisites with Pfizer Security team."],
    )

    return ReadinessScoringEngine.evaluate_and_rescore(base)


def test_deliverable_action_linking(baseline_with_gaps: StartupKitBaseline):
    """Test that deliverables with unconfirmed criteria/owners are linked to their Action ID."""
    deliv_1 = baseline_with_gaps.deliverables[0]
    deliv_2 = baseline_with_gaps.deliverables[1]

    # DEL-01 has gaps, DEL-02 does not
    assert deliv_1.linked_action_id is not None
    assert deliv_1.linked_action_id.startswith("ACT-")
    assert deliv_2.linked_action_id is None


def test_milestone_action_linking(baseline_with_gaps: StartupKitBaseline):
    """Test that milestones with unconfirmed dates are linked to their Action ID."""
    m_1 = baseline_with_gaps.milestones[0]
    m_2 = baseline_with_gaps.milestones[1]

    assert m_1.linked_action_id is not None
    assert m_1.linked_action_id.startswith("ACT-")
    assert m_2.linked_action_id is None


def test_roster_action_linking(baseline_with_gaps: StartupKitBaseline):
    """Test that unstaffed talent roster roles are linked to their Action ID."""
    roster = baseline_with_gaps.talent_onboarding.delivery_talent_roster
    tm_1 = roster[0]
    tm_2 = roster[1]

    assert tm_1.linked_action_id is not None
    assert tm_1.linked_action_id.startswith("ACT-")
    assert tm_2.linked_action_id is None


def test_in_artifact_banner_rendering(baseline_with_gaps: StartupKitBaseline, tmp_path: Path):
    """Test that generated Word document contains in-table action annotations with recovery deltas."""
    generator = DocxGenerator()
    out_file = generator.write_docx(baseline_with_gaps, tmp_path)

    doc = docx.Document(str(out_file))
    doc_text = " ".join(p.text for p in doc.paragraphs) + " " + " ".join(
        c.text for tbl in doc.tables for r in tbl.rows for c in r.cells
    )

    # Check for in-table Action Required annotations and recovery deltas in document
    assert "ACT-" in doc_text
    assert "Deliverables and Acceptance Matrix" in doc_text
    assert "Recovery" in doc_text or "Score" in doc_text


def test_in_artifact_cell_shading(baseline_with_gaps: StartupKitBaseline, tmp_path: Path):
    """Test that table cells with linked action items contain [ACT-XX] text."""
    generator = DocxGenerator()
    out_file = generator.write_docx(baseline_with_gaps, tmp_path)

    doc = docx.Document(str(out_file))
    tables = doc.tables

    # Find Deliverables table
    deliv_table = next(
        tbl for tbl in tables
        if any("acceptance criteria" in c.text.lower() for c in tbl.rows[0].cells)
    )

    # Row 1 (DEL-01) should have [ACT- in acceptance criteria or owner cell
    row_1_text = " ".join(c.text for c in deliv_table.rows[1].cells)
    assert "[ACT-" in row_1_text

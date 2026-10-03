"""Tests for deck completeness, table capacities, and INV-27 (DECK-07, INV-27)."""

import copy

import pytest
from pptx import Presentation

from src.tools.check_artifacts import check_deck_invariants
from tests.deck_bundle import FIXTURES, build_bundle, slide_tables


def _inv27(bundle, deck_path=None):
    out = check_deck_invariants(deck_path or bundle["deck_path"], bundle["manifest_path"], bundle["kit"], bundle["wb"])
    return [str(v) for v in out if v.inv_id == "INV-27"]


def test_inv27_passes_on_the_arc_deck(arc_bundle):
    assert _inv27(arc_bundle) == []


@pytest.mark.parametrize("name", FIXTURES)
def test_inv27_passes_on_every_fixture(name, tmp_path):
    assert _inv27(build_bundle(tmp_path, name)) == []


def test_inv27_fails_when_a_table_exceeds_capacity(arc_bundle, tmp_path):
    prs = Presentation(str(arc_bundle["deck_path"]))
    tbl = slide_tables(prs.slides[3])[0].table._tbl
    for _ in range(6):
        tbl.append(copy.deepcopy(tbl.tr_lst[1]))
    broken = tmp_path / "broken_capacity.pptx"
    prs.save(str(broken))
    msgs = _inv27(arc_bundle, broken)
    assert msgs and "exceeding capacity" in msgs[0]


def _replace_everywhere(slide, old, new):
    n = 0
    for sh in slide.shapes:
        frames = []
        if sh.has_text_frame:
            frames.append(sh.text_frame)
        if getattr(sh, "has_table", False) and sh.has_table:
            frames.extend(c.text_frame for r in sh.table.rows for c in r.cells)
        for tf in frames:
            for p in tf.paragraphs:
                for r in p.runs:
                    if old in r.text:
                        r.text = r.text.replace(old, new)
                        n += 1
    return n


def test_inv27_fails_when_a_deliverable_id_is_dropped_from_slide_4(arc_bundle, tmp_path):
    prs = Presentation(str(arc_bundle["deck_path"]))
    assert _replace_everywhere(prs.slides[3], "DEL-19", "DEL-xx") >= 1
    broken = tmp_path / "dropped_id.pptx"
    prs.save(str(broken))
    assert any("Slide 4 does not contain DEL-19" in m for m in _inv27(arc_bundle, broken))


def test_inv27_fails_when_a_gate_or_deliverable_is_dropped_from_slide_3(arc_bundle, tmp_path):
    prs = Presentation(str(arc_bundle["deck_path"]))
    assert _replace_everywhere(prs.slides[2], "M4", "Mx") >= 1
    assert _replace_everywhere(prs.slides[2], "DEL-16", "DEL-yy") >= 1
    broken = tmp_path / "dropped_gate.pptx"
    prs.save(str(broken))
    msgs = _inv27(arc_bundle, broken)
    assert any("Slide 3 does not contain M4" in m for m in msgs)
    assert any("Slide 3 does not contain DEL-16" in m for m in msgs)


def test_overflow_row_lists_every_omitted_id(arc_bundle):
    """DECK-07: the last row of a table that exceeds its capacity lists every omitted ID."""
    table = slide_tables(arc_bundle["prs"].slides[3])[0].table
    rows = [r.cells[0].text for r in table.rows][1:]
    shown = [r.split()[0] for r in rows[:-1]]
    last = rows[-1]
    assert last.startswith("+") and last.endswith("(see Startup Kit · Deliverables and Acceptance Matrix)")
    listed = [x.strip(" ,") for x in last.split("more:")[1].split("(see")[0].split(",") if x.strip()]
    assert int(last[1:last.index(" ")]) == len(listed)
    assert sorted(shown + listed) == [f"DEL-{n:02d}" for n in range(1, 20)]


def test_slide5_capacity_counts_overflow_row_with_eight_risks(tmp_path):
    """DECK-07, INV-27: 8 qualifying risks must produce 5 real rows + '+3 more' overflow row (6 total rows)."""
    from src.core.models import StartupKitBaseline, RiskAssumption, ProjectStartupCharter
    from src.generators.onboarding_deck.builder import build_deck_model
    from src.generators.pmo_workbook.builder import build_workbook_model
    from src.generators.onboarding_deck.writer import write_onboarding_deck

    risks = [
        RiskAssumption(id=f"RSK-0{i}", type="Risk", description=f"Risk number {i} with high impact", probability="High", impact="High", owner="Delivery Manager", mitigation=f"Mitigate {i}", workstream="P1")
        for i in range(1, 9)
    ]
    baseline = StartupKitBaseline(project_name="Capacity Test 8", charter=ProjectStartupCharter(project_name="Capacity Test 8"), raid_items=risks)
    wb = build_workbook_model(baseline)
    model = build_deck_model(baseline, wb)
    deck_path = write_onboarding_deck(model, tmp_path / "deck8.pptx")
    prs = Presentation(str(deck_path))
    tbl = slide_tables(prs.slides[4])[0].table
    body_rows = list(tbl.rows)[1:]
    assert len(body_rows) == 6, f"Expected 6 total body rows (5 real + 1 overflow), got {len(body_rows)}"
    assert body_rows[-1].cells[0].text.startswith("+3 more: ")


def test_slide5_capacity_boundary_case_with_six_risks(tmp_path):
    """DECK-07, INV-27: Exactly 6 qualifying risks must show all 6 real rows with NO overflow row."""
    from src.core.models import StartupKitBaseline, RiskAssumption, ProjectStartupCharter
    from src.generators.onboarding_deck.builder import build_deck_model
    from src.generators.pmo_workbook.builder import build_workbook_model
    from src.generators.onboarding_deck.writer import write_onboarding_deck

    risks = [
        RiskAssumption(id=f"RSK-0{i}", type="Risk", description=f"Risk number {i} with high impact", probability="High", impact="High", owner="Delivery Manager", mitigation=f"Mitigate {i}", workstream="P1")
        for i in range(1, 7)
    ]
    baseline = StartupKitBaseline(project_name="Capacity Test 6", charter=ProjectStartupCharter(project_name="Capacity Test 6"), raid_items=risks)
    wb = build_workbook_model(baseline)
    model = build_deck_model(baseline, wb)
    deck_path = write_onboarding_deck(model, tmp_path / "deck6.pptx")
    prs = Presentation(str(deck_path))
    tbl = slide_tables(prs.slides[4])[0].table
    body_rows = list(tbl.rows)[1:]
    assert len(body_rows) == 6, f"Expected 6 real body rows, got {len(body_rows)}"
    assert not body_rows[-1].cells[0].text.startswith("+"), "Boundary case (6 items) should have no overflow row"

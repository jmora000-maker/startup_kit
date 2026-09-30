"""Comprehensive unit and integration tests for Project Delivery Workbook generation (v2 spec)."""

import pytest
from datetime import date, timedelta
from pathlib import Path
from unittest.mock import patch
import openpyxl

from src.core.models import (
    StartupKitBaseline,
    ProjectStartupCharter,
    GovernanceContext,
    Deliverable,
    Milestone,
    RiskAssumption,
    DependencyAssumptionItem,
    ContractAmbiguityItem,
    WorkPackageSeed,
    ActionRequiredItem,
    GateDecision,
    TalentOnboardingRecord,
    CommercialGuardrail,
    CommunicationsPlanItem,
    SourceReference,
    OutputSelection,
    RunResult,
)
from src.generators.pmo_workbook import (
    export_pmo_workbook,
    PMOWorkbookResult,
)
from src.generators.pmo_workbook.workstreams import (
    TAXONOMY,
    classify_milestone,
    parse_milestone_phase,
)
from src.generators.pmo_workbook.task_library import (
    classify_deliverable_work_type,
    WORK_TYPES,
)
from src.generators.pmo_workbook.mapping import (
    tokenize_v2,
    compute_idf,
    compute_score,
    map_deliverables_to_milestones_v2,
    map_work_packages_to_deliverables,
    link_raid_item_v2,
    detect_default_filled_milestone,
)
from src.generators.pmo_workbook.builder import (
    build_workbook_model,
    next_working_day,
    clean_text_v2,
    truncate_task_name,
    deduplicate_notes,
    normalize_owner_v2,
)
from src.generators.pmo_workbook.writer import write_workbook
from src.orchestrator import StartupKitController
from src.extractors.service import IngestionService
from src.generators.docx_generator import DocxGenerator
from main import create_mock_llm_client


# =========================================================================
# 1. Output Selection Unit & Matrix Tests (Section 3.1)
# =========================================================================

def test_output_selection_defaults_and_flags():
    """Verify all OutputSelection flag combinations per Section 3.1 table."""
    # 1. Default (no flags) -> workbook only
    sel = OutputSelection.from_flags(all_=False, kit=False, checklist=False, export_tools=False)
    assert sel == OutputSelection(kit=False, checklist=False, workbook=True)

    # 2. --kit -> Kit only
    sel = OutputSelection.from_flags(all_=False, kit=True, checklist=False, export_tools=False)
    assert sel == OutputSelection(kit=True, checklist=False, workbook=False)

    # 3. --checklist -> Checklist only
    sel = OutputSelection.from_flags(all_=False, kit=False, checklist=True, export_tools=False)
    assert sel == OutputSelection(kit=False, checklist=True, workbook=False)

    # 4. --kit --checklist -> Kit + Checklist
    sel = OutputSelection.from_flags(all_=False, kit=True, checklist=True, export_tools=False)
    assert sel == OutputSelection(kit=True, checklist=True, workbook=False)

    # 5. --export-tools -> Workbook only
    sel = OutputSelection.from_flags(all_=False, kit=False, checklist=False, export_tools=True)
    assert sel == OutputSelection(kit=False, checklist=False, workbook=True)

    # 6. --kit --export-tools -> Kit + Workbook
    sel = OutputSelection.from_flags(all_=False, kit=True, checklist=False, export_tools=True)
    assert sel == OutputSelection(kit=True, checklist=False, workbook=True)

    # 7. --checklist --export-tools -> Checklist + Workbook
    sel = OutputSelection.from_flags(all_=False, kit=False, checklist=True, export_tools=True)
    assert sel == OutputSelection(kit=False, checklist=True, workbook=True)

    # 8. --all -> All three
    sel = OutputSelection.from_flags(all_=True, kit=False, checklist=False, export_tools=False)
    assert sel == OutputSelection(kit=True, checklist=True, workbook=True)

    # 9. --all with other flags -> All three
    sel = OutputSelection.from_flags(all_=True, kit=True, checklist=False, export_tools=False)
    assert sel == OutputSelection(kit=True, checklist=True, workbook=True)


def test_controller_run_output_matrix(populated_inputs_dir, tmp_path):
    """Verify StartupKitController.run() writes exactly what OutputSelection specifies."""
    mock_llm = create_mock_llm_client()
    controller = StartupKitController(
        ingestion_service=IngestionService(),
        llm_client=mock_llm,
        aggregator=None,
        doc_writer=DocxGenerator()
    )

    # Case A: Default (workbook only)
    out_dir_a = tmp_path / "out_a"
    res_a = controller.run(inputs_dir=populated_inputs_dir, output_dir=out_dir_a, outputs=OutputSelection())
    assert res_a.kit_path is None
    assert res_a.checklist_path is None
    assert res_a.workbook is not None
    assert len(list(out_dir_a.glob("*.docx"))) == 0
    assert len(list(out_dir_a.glob("*.xlsx"))) == 1

    # Case B: --kit only
    out_dir_b = tmp_path / "out_b"
    res_b = controller.run(inputs_dir=populated_inputs_dir, output_dir=out_dir_b, outputs=OutputSelection(kit=True, checklist=False, workbook=False))
    assert res_b.kit_path is not None
    assert res_b.checklist_path is None
    assert res_b.workbook is None
    assert len(list(out_dir_b.glob("*_Startup_Kit.docx"))) == 1
    assert len(list(out_dir_b.glob("*.docx"))) == 1
    assert len(list(out_dir_b.glob("*.xlsx"))) == 0

    # Case C: --checklist only
    out_dir_c = tmp_path / "out_c"
    res_c = controller.run(inputs_dir=populated_inputs_dir, output_dir=out_dir_c, outputs=OutputSelection(kit=False, checklist=True, workbook=False))
    assert res_c.kit_path is None
    assert res_c.checklist_path is not None
    assert res_c.workbook is None
    assert len(list(out_dir_c.glob("*_Startup_Readiness_Checklist.docx"))) == 1
    assert len(list(out_dir_c.glob("*.docx"))) == 1
    assert len(list(out_dir_c.glob("*.xlsx"))) == 0

    # Case D: --all (all three files)
    out_dir_d = tmp_path / "out_d"
    res_d = controller.run(inputs_dir=populated_inputs_dir, output_dir=out_dir_d, outputs=OutputSelection.from_flags(all_=True, kit=False, checklist=False, export_tools=False))
    assert res_d.kit_path is not None
    assert res_d.checklist_path is not None
    assert res_d.workbook is not None
    assert len(list(out_dir_d.glob("*_Startup_Kit.docx"))) == 1
    assert len(list(out_dir_d.glob("*_Startup_Readiness_Checklist.docx"))) == 1
    assert len(list(out_dir_d.glob("*.docx"))) == 2
    assert len(list(out_dir_d.glob("*.xlsx"))) == 1


# =========================================================================
# 2. Minimal Baseline Acceptance Test (Section 9)
# =========================================================================

def test_minimal_baseline_generates_valid_workbook(minimal_baseline, tmp_path):
    """Verify minimal_baseline produces valid workbook with MS-TBC, Project Management (ongoing), and empty RAID."""
    result = export_pmo_workbook(minimal_baseline, tmp_path)
    assert result.file_path.exists()
    assert result.file_path.name == "Minimal_Test_Project_Project_Delivery_Workbook.xlsx"

    wb = openpyxl.load_workbook(str(result.file_path))
    assert "Project Schedule" in wb.sheetnames
    assert "WBS" in wb.sheetnames
    assert "RAID Log" in wb.sheetnames
    assert "_Lists" in wb.sheetnames

    # Schedule checks
    ws_sched = wb["Project Schedule"]
    sched_rows = [row for row in ws_sched.iter_rows(min_row=6, values_only=True) if any(row)]
    assert any("Project Management (ongoing)" in str(r[2]) for r in sched_rows)

    # RAID checks
    ws_raid = wb["RAID Log"]
    raid_rows = [row for row in ws_raid.iter_rows(min_row=6, values_only=True) if any(row)]
    assert len(raid_rows) == 1
    assert raid_rows[0][0] == "RAID-01"
    assert "No baseline risks identified" in raid_rows[0][2]


# =========================================================================
# 3. Deliverable Work Type Classification Tests (Section 7.3)
# =========================================================================

def test_deliverable_work_type_classification():
    """Verify keyword scoring and work type assignment."""
    wt_int, _ = classify_deliverable_work_type("Azure AD / MSAL authentication integration API")
    assert wt_int == "Integration"

    wt_bld, _ = classify_deliverable_work_type("ARC micro-frontend faceted search view")
    assert wt_bld == "Build"

    wt_tst, _ = classify_deliverable_work_type("UAT execution records and defect test suite")
    assert wt_tst == "Test"

    wt_ana, _ = classify_deliverable_work_type("Performance engineering spike output report")
    assert wt_ana == "Analysis"

    wt_doc, _ = classify_deliverable_work_type("Training materials and user documentation guide")
    assert wt_doc == "Documentation"

    # Default fallback to Build when no keywords match
    wt_def, _ = classify_deliverable_work_type("General Project Artifact")
    assert wt_def == "Build"


# =========================================================================
# 4. Text Hygiene & Helper Tests (Section 9)
# =========================================================================

def test_text_hygiene_helpers():
    # clean_text_v2
    raw = "  [ACT-01: Fix item] Some description text ; - "
    cleaned = clean_text_v2(raw)
    assert cleaned == "Some description text"

    # truncate_task_name
    short_name = "Short task name"
    t_short, note_short = truncate_task_name(short_name)
    assert t_short == short_name
    assert note_short is None

    long_name = "This is an extremely long task description name that exceeds the maximum limit of one hundred and twenty characters permitted by the hygiene rules"
    t_long, note_long = truncate_task_name(long_name, max_len=50)
    assert len(t_long) <= 50
    assert t_long.endswith("...")
    assert "Full name:" in note_long

    # deduplicate_notes
    notes = ["Note A", "Note B", "Note A", None, "", "Note C"]
    assert deduplicate_notes(notes) == "Note A; Note B; Note C"

    # normalize_owner_v2
    o1, n1 = normalize_owner_v2("  jAMES sMITH  ")
    assert o1 == "James Smith"
    assert n1 is None

    o2, n2 = normalize_owner_v2("Unassigned")
    assert o2 == "[UNASSIGNED - TO BE CONFIRMED]"
    assert n2 == "Owner unassigned"

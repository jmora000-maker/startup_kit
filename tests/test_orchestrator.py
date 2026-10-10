"""End-to-end integration tests for StartupKitController and CLI."""

import pytest
from pathlib import Path
from datetime import date
import docx
import pymupdf as fitz
from pptx import Presentation

from src.orchestrator import StartupKitController
from src.extractors.service import IngestionService
from src.extractors.startup_kit_docx_parser import StartupKitDocxParser
from src.llm.client import MockLLMClient
from src.llm.aggregator import BaselineAggregator
from src.generators.docx_generator import DocxGenerator
from main import create_mock_llm_client, main
import sys


from src.core.models import OutputSelection, RunResult

MOCK_SOW_INPUTS_DIR = Path(__file__).parent / "fixtures" / "sow" / "mock_sow" / "inputs"


def test_startup_kit_controller_e2e(populated_inputs_dir, tmp_path):
    output_dir = tmp_path / "output"
    mock_llm = create_mock_llm_client()

    controller = StartupKitController(
        ingestion_service=IngestionService(),
        llm_client=mock_llm,
        aggregator=BaselineAggregator(),
        doc_writer=DocxGenerator()
    )

    result = controller.run(
        inputs_dir=populated_inputs_dir,
        output_dir=output_dir,
        tier_override="Elevated",
        contract_type_override="Fixed Bid",
        outputs=OutputSelection(kit=True, checklist=True, workbook=True)
    )

    assert isinstance(result, RunResult)
    assert result.kit_path.exists()
    assert result.kit_path.parent == output_dir
    assert result.kit_path.suffix == ".docx"
    assert result.workbook is not None
    assert result.workbook.file_path.exists()

    # Verify content in generated document
    doc = docx.Document(str(result.kit_path))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "TOPTAL PMO STARTUP KIT" in full_text
    assert "Pfizer Analytics & Cloud Modernization" in full_text


def test_controller_empty_inputs_error(tmp_path):
    empty_inputs = tmp_path / "empty_inputs"
    empty_inputs.mkdir()
    output_dir = tmp_path / "output"

    controller = StartupKitController(
        llm_client=create_mock_llm_client()
    )

    with pytest.raises(FileNotFoundError):
        controller.run(inputs_dir=empty_inputs, output_dir=output_dir)


def test_cli_main_execution_with_export_tools(populated_inputs_dir, tmp_path, monkeypatch):
    output_dir = tmp_path / "cli_tools_output"

    test_args = [
        "main.py",
        "--mock",
        "--inputs-dir", str(populated_inputs_dir),
        "--output-dir", str(output_dir),
        "--tier", "Elevated",
        "--contract-type", "Fixed Bid",
        "--all"
    ]
    monkeypatch.setattr(sys, "argv", test_args)

    exit_code = main()
    assert exit_code == 0

    created_docx = list(output_dir.glob("*_Startup_Kit.docx"))
    assert len(created_docx) == 1
    assert len(list(output_dir.glob("*.docx"))) == 2

    created_xlsx = list(output_dir.glob("*.xlsx"))
    assert len(created_xlsx) == 1


def test_concurrent_extraction_all_domain_passes(populated_inputs_dir, tmp_path):
    """Verify that concurrent multi-pass extraction properly executes all domain extractors."""
    output_dir = tmp_path / "concurrent_output"
    mock_llm = create_mock_llm_client()

    controller = StartupKitController(
        ingestion_service=IngestionService(),
        llm_client=mock_llm,
        aggregator=BaselineAggregator(),
        doc_writer=DocxGenerator()
    )

    result = controller.run(
        inputs_dir=populated_inputs_dir,
        output_dir=output_dir,
        pmo_lead="Sarah Connor",
        delivery_lead="John Connor",
        talent_pm="Kyle Reese",
        outputs=OutputSelection(kit=True)
    )

    assert result.kit_path.exists()
    doc = docx.Document(str(result.kit_path))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Sarah Connor" in full_text
    assert "John Connor" in full_text
    assert "Kyle Reese" in full_text


def test_on_progress_fires_expected_stage_sequence_for_run(populated_inputs_dir, tmp_path):
    """HTL-28: on_progress(stage, detail="") must fire, in order, at the real stage boundaries
    of run() -- "ingesting", "extracting", "validating", then "generating" once per document
    type actually produced -- for a real run through the mock client, not a simulated sequence.
    """
    output_dir = tmp_path / "progress_output"
    mock_llm = create_mock_llm_client()

    controller = StartupKitController(
        ingestion_service=IngestionService(),
        llm_client=mock_llm,
        aggregator=BaselineAggregator(),
        doc_writer=DocxGenerator()
    )

    calls = []

    def on_progress(stage, detail=""):
        calls.append((stage, detail))

    result = controller.run(
        inputs_dir=populated_inputs_dir,
        output_dir=output_dir,
        outputs=OutputSelection(kit=True, checklist=True, workbook=True, slides=True),
        on_progress=on_progress,
    )

    assert result.kit_path.exists()
    stages = [stage for stage, _ in calls]
    assert stages[0] == "ingesting"
    extracting_calls = [c for c in calls if c[0] == "extracting"]
    assert len(extracting_calls) == 16  # 1 initial + 13 concurrent increments + 1 backlog start + 1 final 14/14 increment
    assert extracting_calls[0] == ("extracting", "")
    expected_details = [f"{i}/14 complete" for i in range(1, 14)] + [
        "13/14 complete -- decomposing backlog & SOW catalogue...",
        "14/14 complete",
    ]
    assert [detail for _, detail in extracting_calls[1:]] == expected_details
    assert stages[17] == "validating"
    generating_calls = [c for c in calls if c[0] == "generating"]
    assert [detail for _, detail in generating_calls] == [
        "Startup Kit",
        "Readiness Checklist",
        "Delivery Workbook",
        "Onboarding Deck",
    ]


def test_on_progress_fires_expected_stage_sequence_for_run_reingest(sample_baseline, tmp_path):
    """HTL-28: on_progress must fire, in order, at run_reingest()'s real stage boundaries --
    "ingesting" (parsing the existing Kit docx), "validating" (the readiness recalculation; no
    "extracting" stage exists here since re-ingestion never re-runs LLM extraction), then
    "generating" once per document type actually produced.
    """
    writer = DocxGenerator()
    original_docx = writer.write_docx(sample_baseline, tmp_path / "ProgressProject_Startup_Kit.docx")

    controller = StartupKitController(
        doc_writer=writer,
        aggregator=BaselineAggregator(),
        docx_parser=StartupKitDocxParser()
    )

    calls = []

    def on_progress(stage, detail=""):
        calls.append((stage, detail))

    result = controller.run_reingest(
        docx_path=original_docx,
        outputs=OutputSelection(kit=True, checklist=True, workbook=True),
        create_backup=False,
        on_progress=on_progress,
    )

    assert result.kit_path.exists()
    stages = [stage for stage, _ in calls]
    assert stages == [
        "ingesting",
        "validating",
        "generating",  # Startup Kit
        "generating",  # Readiness Checklist
        "generating",  # Delivery Workbook
    ]
    generating_details = [detail for stage, detail in calls if stage == "generating"]
    assert generating_details == ["Startup Kit", "Readiness Checklist", "Delivery Workbook"]


def test_controller_default_mock_inputs_dir(tmp_path, monkeypatch):
    """Verify that StartupKitController defaults to config.mock_inputs_dir when inputs_dir is None and MockLLMClient is used."""
    output_dir = tmp_path / "mock_dir_output"
    mock_llm = create_mock_llm_client()

    captured_inputs = {}

    class DummyIngestionService(IngestionService):
        def ingest_directory(self, directory_path):
            captured_inputs["dir"] = Path(directory_path)
            return super().ingest_directory(directory_path)

    controller = StartupKitController(
        ingestion_service=DummyIngestionService(),
        llm_client=mock_llm,
        aggregator=BaselineAggregator(),
        doc_writer=DocxGenerator()
    )

    from src.config import config
    # Point the default at the tracked fixture via an absolute path so the test is independent of the CWD
    monkeypatch.setattr(config, "mock_inputs_dir", MOCK_SOW_INPUTS_DIR)
    result = controller.run(
        inputs_dir=None,
        output_dir=output_dir,
        outputs=OutputSelection(kit=True)
    )
    assert captured_inputs["dir"] == MOCK_SOW_INPUTS_DIR
    assert result.kit_path.exists()


def test_controller_default_mock_output_dir(tmp_path, monkeypatch):
    """Verify that StartupKitController defaults to config.mock_output_dir when output_dir is None and MockLLMClient is used."""
    mock_llm = create_mock_llm_client()

    captured_paths = {}

    class DummyDocWriter:
        def write_kit_docx(self, baseline, output_path):
            captured_paths["out"] = Path(output_path)
            return Path(output_path) / "test.docx"

        def write_docx(self, baseline, output_path):
            captured_paths["out"] = Path(output_path)
            return Path(output_path) / "test.docx"

    controller = StartupKitController(
        ingestion_service=IngestionService(),
        llm_client=mock_llm,
        aggregator=BaselineAggregator(),
        doc_writer=DummyDocWriter()
    )

    from src.config import config
    controller.run(
        inputs_dir=MOCK_SOW_INPUTS_DIR,
        output_dir=None,
        outputs=OutputSelection(kit=True)
    )
    assert captured_paths["out"] == config.mock_output_dir


def test_val_11_date_conflict_reaches_baseline_as_finding_or_open_question(tmp_path):
    """VAL-11/INV-38: a date-conflict warning from extract_stated_award_date must reach the
    baseline through orchestrator.run() as a real ValidationFinding or open question naming both
    disagreeing dates and which one was preferred -- not only a log line that nothing ever reads.
    """
    inputs_dir = tmp_path / "inputs"
    inputs_dir.mkdir()
    (inputs_dir / "sow.txt").write_text(
        "This SOW becomes effective on October 7, 2026. "
        "Section 3: Estimated Start Date October 14, 2026.",
        encoding="utf-8",
    )
    output_dir = tmp_path / "output"

    captured = {}

    class CapturingDocWriter:
        def write_kit_docx(self, baseline, output_path):
            captured["baseline"] = baseline
            return Path(output_path) / "test.docx"

    controller = StartupKitController(
        ingestion_service=IngestionService(),
        llm_client=create_mock_llm_client(),
        aggregator=BaselineAggregator(),
        doc_writer=CapturingDocWriter(),
    )
    controller.run(
        inputs_dir=inputs_dir,
        output_dir=output_dir,
        outputs=OutputSelection(kit=True),
    )

    baseline = captured["baseline"]
    # The preamble date wins per VAL-11's precedence rule; this is the resolved date, not the conflict itself.
    assert baseline.sow_awarded_date == date(2026, 10, 7)

    conflict_findings = [
        f for f in (baseline.validation_report.findings if baseline.validation_report else [])
        if "2026-10-07" in f.message and "2026-10-14" in f.message
    ]
    conflict_open_questions = [
        q for q in baseline.open_questions
        if "2026-10-07" in q and "2026-10-14" in q
    ]
    assert conflict_findings or conflict_open_questions, (
        "Expected the award-date conflict warning (naming both 2026-10-07 and 2026-10-14, and "
        "which was preferred) to reach the baseline as a validation finding or open question, "
        "but found neither."
    )


def test_orchestrator_exhausted_providers_raises_clear_stage_error(tmp_path):
    """Verify LLM-03 (4a): when extraction fails across all providers, orchestrator raises clear stage error."""
    from unittest.mock import MagicMock
    from src.llm.client import LangChainLLMClient
    from src.review_ui.generate import format_run_error

    mock_anthropic = MagicMock()
    mock_anthropic.with_structured_output.side_effect = RuntimeError("Anthropic 429 RateLimit")
    mock_anthropic.invoke.side_effect = RuntimeError("Anthropic 429 RateLimit")

    mock_openai = MagicMock()
    mock_openai.with_structured_output.side_effect = RuntimeError("OpenAI 429 QuotaExceeded")
    mock_openai.invoke.side_effect = RuntimeError("OpenAI 429 QuotaExceeded")

    client = LangChainLLMClient(
        primary_provider="anthropic",
        chat_model=mock_anthropic,
        openai_chat_model=mock_openai,
    )

    inputs_dir = tmp_path / "inputs"
    inputs_dir.mkdir()
    (inputs_dir / "sow.txt").write_text("Statement of Work content", encoding="utf-8")
    output_dir = tmp_path / "output"

    controller = StartupKitController(
        llm_client=client,
    )

    with pytest.raises(RuntimeError) as exc_info:
        controller.run(inputs_dir=inputs_dir, output_dir=output_dir)

    err_str = str(exc_info.value)
    assert err_str.startswith("extraction failed:")
    assert "OpenAI 429 QuotaExceeded" in err_str
    # Verify no raw traceback in error message
    assert "Traceback (most recent call last)" not in err_str

    # Verify that UI's format_run_error formats it cleanly without duplicate stage prefixes
    formatted = format_run_error(exc_info.value, "extracting")
    assert formatted.startswith("extraction failed:")
    assert not formatted.startswith("extraction failed: extraction failed:")


def test_orchestrator_openai_only_truncation_raises_clear_stage_error(tmp_path):
    """Verify LLM-03 (4b compound case): when OpenAI is the only configured provider and its response truncates,
    orchestrator raises clear stage error ('extraction failed: ...'), not returning truncated content and not a raw traceback.
    """
    from unittest.mock import MagicMock
    from src.llm.client import LangChainLLMClient
    from src.review_ui.generate import format_run_error

    mock_openai = MagicMock()
    mock_openai_response = MagicMock()
    mock_openai_response.response_metadata = {"finish_reason": "length"}
    mock_openai_response.content = '{"title": "Truncated Incomplete Payload'
    mock_openai.invoke.return_value = mock_openai_response
    mock_openai.with_structured_output.side_effect = RuntimeError(
        "OpenAI response was truncated due to reaching max_tokens (16384). Payload is incomplete and cannot be parsed safely."
    )

    client = LangChainLLMClient(
        primary_provider="openai",
        api_key="",
        chat_model=None,
        openai_chat_model=mock_openai,
    )

    inputs_dir = tmp_path / "inputs"
    inputs_dir.mkdir()
    (inputs_dir / "sow.txt").write_text("Statement of Work content", encoding="utf-8")
    output_dir = tmp_path / "output"

    controller = StartupKitController(
        llm_client=client,
    )

    with pytest.raises(RuntimeError) as exc_info:
        controller.run(inputs_dir=inputs_dir, output_dir=output_dir)

    err_str = str(exc_info.value)
    assert err_str.startswith("extraction failed:")
    assert "OpenAI response was truncated due to reaching max_tokens" in err_str
    assert "Payload is incomplete and cannot be parsed safely" in err_str
    # Verify no raw traceback in error message
    assert "Traceback (most recent call last)" not in err_str

    # Verify UI formatting preserves the clear stage message without duplicate prefixes
    formatted = format_run_error(exc_info.value, "extracting")
    assert formatted.startswith("extraction failed:")
    assert not formatted.startswith("extraction failed: extraction failed:")

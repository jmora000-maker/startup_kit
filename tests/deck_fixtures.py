"""Small synthetic Kit, Workbook, and deck builders for the broken-input unit tests of INV-26, INV-30 to INV-32."""

from pathlib import Path
from typing import Any, Dict, List, Optional

import docx
import openpyxl
import pptx
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

KIT_RAID_ROW = [
    "RSK-01", "Risk", "Latency targets may be missed because Client components are slow.", "Technical",
    "Medium", "High", "Delivery Manager", "Use query observability (HS-4942) to attribute misses", "Open",
]


def make_kit(path: Path, raid_row: Optional[List[str]] = None) -> docx.Document:
    d = docx.Document()
    t = d.add_table(rows=2, cols=4)
    for r, vals in enumerate([["Project Name", "Demo Project", "Client Sponsor", "Acme"], ["Governance Tier", "Partnered", "Contract Type", "Fixed Bid"]]):
        for c, val in enumerate(vals):
            t.cell(r, c).text = val
    d.add_heading("Deliverables and Acceptance Matrix", level=2)
    t = d.add_table(rows=3, cols=3)
    for r, vals in enumerate([
        ["ID", "Deliverable Name", "Acceptance Criteria"],
        ["DEL-01", "Backend Integration/Load Tests and Performance Engineering Spike", "Search P95 < 1s under load. Targets are defined, and the backlog is triaged."],
        ["DEL-02", "Production Smoke Tests and 48-Hour Defect Watch", "Smoke tests pass 100% and the defect watch is clean."],
    ]):
        for c, val in enumerate(vals):
            t.cell(r, c).text = val
    d.add_heading("RAID Log", level=2)
    t = d.add_table(rows=2, cols=9)
    hdr = ["Item ID", "Type", "Description", "Category", "Probability", "Impact", "Owner", "Mitigation / Response", "Status"]
    for c, val in enumerate(hdr):
        t.cell(0, c).text = val
    for c, val in enumerate(raid_row or KIT_RAID_ROW):
        t.cell(1, c).text = val
    d.save(str(path))
    return docx.Document(str(path))


def make_workbook(path: Path, raid_overrides: Optional[Dict[str, str]] = None) -> Any:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Project Schedule"
    ws["A1"] = "DEMO PROJECT - PROJECT SCHEDULE"
    ws["A2"] = "Client: Acme | Contract: Fixed Bid"
    ws["A3"] = "Start Date: 2026-10-05 (Provided) | Generated 2026-10-02 from the project baseline"
    ws.append([])
    ws.append(["WBS Code", "Row Type", "Workstream", "Milestone ID", "Milestone", "Planned Start", "Planned Finish"])
    ws.append(["1.1", "Milestone", "P1 Foundation", "M1", "P1 Foundation accepted", "2026-10-05", "2026-11-13"])
    ws2 = wb.create_sheet("RAID Log")
    for _ in range(4):
        ws2.append([])
    hdr = ["RAID ID", "Type", "Description", "Owner", "Probability", "Impact", "Trigger / Early Warning", "Mitigation / Response", "Status", "Source ID"]
    ws2.append(hdr)
    row = {
        "RAID ID": "RAID-01", "Type": "Risk", "Description": "Latency targets may be missed because Client components are slow.",
        "Owner": "Delivery Manager", "Probability": "Medium", "Impact": "High", "Trigger / Early Warning": "",
        "Mitigation / Response": "Use query observability (HS-4942) to attribute misses", "Status": "Open", "Source ID": "RSK-01",
    }
    row.update(raid_overrides or {})
    ws2.append([row[h] for h in hdr])
    wb.save(str(path))
    return openpyxl.load_workbook(str(path), data_only=False)


def blank_deck(slide_count: int = 2):
    prs = pptx.Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    slides = [prs.slides.add_slide(blank) for _ in range(slide_count)]
    return prs, slides


def add_card(slide, x, y, w, h):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    card.name = "Card"
    return card


def add_text(slide, x, y, w, h, paragraphs, name="Text", size=11.0, margins=0.0):
    """A text frame with word wrap on, autofit off, top anchor, and explicit run sizes."""
    from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE

    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    box.name = name
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(margins)
    for i, item in enumerate(paragraphs):
        text, psize = (item if isinstance(item, tuple) else (item, size))
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = text
        r.font.size = Pt(psize)
        r.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
    return box


def add_table(slide, x, y, col_widths, rows, size=10.0, row_h=0.3, name="Table"):
    shape = slide.shapes.add_table(len(rows), len(col_widths), Inches(x), Inches(y), Inches(sum(col_widths)), Inches(row_h * len(rows)))
    shape.name = name
    for i, w in enumerate(col_widths):
        shape.table.columns[i].width = Inches(w)
    for ri, r in enumerate(rows):
        shape.table.rows[ri].height = Inches(row_h)
        for ci, val in enumerate(r):
            cell = shape.table.cell(ri, ci)
            cell.margin_left = cell.margin_right = Inches(0.06)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            cell.text_frame.paragraphs[0].text = ""
            run = cell.text_frame.paragraphs[0].add_run()
            run.text = val
            run.font.size = Pt(size)
    return shape


def trace(artifact: str, locator: str, key: str, field: str) -> Dict[str, str]:
    return {"artifact": artifact, "locator": locator, "key": key, "field": field}


def entry(slide: int, shape: str, element: str, value: str, *traces: Dict[str, str], **extra: Any) -> Dict[str, Any]:
    e = {"slide": slide, "shape": shape, "element": element, "displayed_value": value, "traces": list(traces)}
    e.update(extra)
    return e

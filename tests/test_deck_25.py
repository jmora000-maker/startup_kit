"""Tests for DECK-25: Slide 3 Deliverables column shows deliverable names rather than bare IDs."""

import pytest
from datetime import date
from pathlib import Path
import tempfile
import pptx
import docx
import openpyxl

from src.core.models import (
    StartupKitBaseline,
    Deliverable,
    Milestone,
    SOWWorkItem,
    ProjectStartupCharter,
)
from src.generators.pmo_workbook.builder import build_workbook_model
from src.generators.pmo_workbook import export_pmo_workbook
from src.generators.onboarding_deck import build_deck_model, export_onboarding_deck
from src.generators.docx_generator import DocxGenerator
from src.llm.validation import validate_and_repair_baseline
from src.tools.check_artifacts import check_deck_invariants
from src.tools import deck_checks
from tests.test_map_16 import ALL_UNDIVIDE_WORK_OUTPUTS


def _make_undivide_baseline() -> StartupKitBaseline:
    deliverables = []
    catalogue = []
    for idx, (ms_id, wo_title, ac_text, ref) in enumerate(ALL_UNDIVIDE_WORK_OUTPUTS, start=1):
        deliv_id = f"DEL-{idx:02d}"
        deliverables.append(
            Deliverable(
                id=deliv_id,
                name=wo_title,
                description=f"Delivery and operationalization of {wo_title}.",
                acceptance_criteria=ac_text,
                sow_reference=ref,
                owner="Toptal Delivery Team",
                client_approver="Client Approver",
            )
        )
        catalogue.append(
            SOWWorkItem(
                reference=ref,
                reference_kind="Synthetic",
                title=wo_title,
                phase="P" + ms_id[1:] if ms_id.startswith("M") else ms_id,
                owner="Toptal",
                type="Build",
                deliverable_id=deliv_id,
            )
        )

    milestones = [
        Milestone(id="M1", description="Phase 1 - Data and Backend Foundation: Infrastructure and pipeline deployment."),
        Milestone(id="M2", description="Phase 2 - LLM Enrichment Layer: Enrichment service and prompt pipelines."),
        Milestone(id="M3", description="Phase 3 - Query and Retrieval Interface: Search API and caching layer."),
        Milestone(id="M4", description="Phase 4 - Access Control, Observability, and Pilot Readiness: RBAC and production cutover."),
    ]

    baseline = StartupKitBaseline(
        project_name="The Undivide Project",
        charter=ProjectStartupCharter(project_name="The Undivide Project"),
        deliverables=deliverables,
        milestones=milestones,
        sow_stories_catalogue=catalogue,
        backlog_seed=[],
    )
    validate_and_repair_baseline(baseline)
    return baseline


def test_deck_25_deliverables_column_shows_names_for_realistic_counts(tmp_path):
    """DECK-25: Slide 3 Deliverables column renders deliverable names, never bare IDs, for 27-item SOW."""
    baseline = _make_undivide_baseline()
    wb = build_workbook_model(baseline)
    deck = build_deck_model(baseline, wb)

    slide3 = deck.slides[1]  # Slide 3 is index 1 in deck.slides (slides 2 to 7)
    assert slide3.number == 3
    assert len(slide3.tables) == 1
    table_spec = slide3.tables[0]

    # Verify rows in Slide 3 table
    # Row 0 is M1, Row 1 is M2, Row 2 is overflow row
    m1_row = table_spec.rows[0]
    m1_deliv_cell = m1_row.cells[3]

    # M1 has 8 deliverables: DEL-01 to DEL-08
    assert len(m1_deliv_cell.paras) == 8
    for i, line in enumerate(m1_deliv_cell.paras, start=1):
        deliv_id = f"DEL-{i:02d}"
        expected_name = ALL_UNDIVIDE_WORK_OUTPUTS[i - 1][1]
        line_text = "".join(r.text for r in line)
        assert line_text.startswith(deliv_id), f"Expected line to start with {deliv_id}, got: {line_text}"
        assert expected_name in line_text, f"Expected {expected_name} in line, got: {line_text}"
        assert line_text != deliv_id, "Deliverable column must NOT be a bare ID"

    # M2 has 6 deliverables: DEL-09 to DEL-14
    m2_row = table_spec.rows[1]
    m2_deliv_cell = m2_row.cells[3]
    assert len(m2_deliv_cell.paras) == 6
    for i, line in enumerate(m2_deliv_cell.paras, start=9):
        deliv_id = f"DEL-{i:02d}"
        expected_name = ALL_UNDIVIDE_WORK_OUTPUTS[i - 1][1]
        line_text = "".join(r.text for r in line)
        assert line_text.startswith(deliv_id)
        assert expected_name in line_text
        assert line_text != deliv_id

    # Font size is compliant (9.0 pt for deliverables column)
    assert table_spec.deliverables_pt in (9.0, 10.0)


def test_deck_25_overflow_row_with_high_deliverable_count(tmp_path):
    """DECK-25 / INV-27: Omitted milestones and deliverables degrade cleanly into '+N more' overflow row."""
    baseline = _make_undivide_baseline()

    kit_path = DocxGenerator(generated_date="2026-10-01").write_kit_docx(baseline, tmp_path)
    wb_res = export_pmo_workbook(baseline, tmp_path)
    deck_res = export_onboarding_deck(baseline, tmp_path)

    prs = pptx.Presentation(str(deck_res.file_path))
    slide3_table = [sh for sh in prs.slides[2].shapes if sh.has_table][0].table

    # Body rows: M1 row, M2 row, +2 more overflow row
    body_rows = list(slide3_table.rows)[1:]
    assert len(body_rows) == 3

    last_row = body_rows[-1]
    overflow_text = last_row.cells[0].text
    assert overflow_text.startswith("+2 more: ")
    assert "M3" in overflow_text
    assert "M4" in overflow_text
    for did in [f"DEL-{i:02d}" for i in range(15, 28)]:
        assert did in overflow_text, f"Expected {did} in overflow text: {overflow_text}"
    assert overflow_text.endswith("(see Project Delivery Workbook · Project Schedule)")

    # Validate INV-32 and INV-27 layout and completeness
    kit_doc = docx.Document(str(kit_path))
    wb_doc = openpyxl.load_workbook(str(wb_res.file_path), data_only=False)
    violations = check_deck_invariants(deck_res.file_path, deck_res.manifest_path, kit_doc, wb_doc)

    # Filter out talking points count (not relevant to table layout) and check INV-32 & INV-27
    layout_violations = [v for v in violations if v.inv_id in ("INV-32", "INV-27")]
    assert layout_violations == []


def test_deck_25_moderate_deliverable_count_fits_at_10pt_or_9pt():
    """DECK-25: A baseline with standard deliverable counts fits with full names without bare IDs."""
    deliverables = [
        Deliverable(id="DEL-01", name="Postgres Database and Schema Setup", description="DB setup.", owner="Team"),
        Deliverable(id="DEL-02", name="API Endpoints and Authentication", description="API endpoints.", owner="Team"),
        Deliverable(id="DEL-03", name="Frontend Admin Console", description="Admin UI.", owner="Team"),
    ]
    milestones = [
        Milestone(id="M1", description="Phase 1 - Foundation and Data Setup"),
        Milestone(id="M2", description="Phase 2 - Application Features and UI"),
    ]
    catalogue = [
        SOWWorkItem(reference="WO-01", reference_kind="Synthetic", title="DB Setup", phase="P1", deliverable_id="DEL-01"),
        SOWWorkItem(reference="WO-02", reference_kind="Synthetic", title="API Endpoints", phase="P1", deliverable_id="DEL-02"),
        SOWWorkItem(reference="WO-03", reference_kind="Synthetic", title="Admin Console", phase="P2", deliverable_id="DEL-03"),
    ]
    baseline = StartupKitBaseline(
        project_name="Small Project",
        charter=ProjectStartupCharter(project_name="Small Project"),
        deliverables=deliverables,
        milestones=milestones,
        sow_stories_catalogue=catalogue,
        backlog_seed=[],
    )
    validate_and_repair_baseline(baseline)
    wb = build_workbook_model(baseline)
    deck = build_deck_model(baseline, wb)

    slide3 = deck.slides[1]
    table_spec = slide3.tables[0]
    assert len(table_spec.rows) == 2  # M1 and M2 (no overflow)
    assert table_spec.deliverables_pt == 10.0  # Fits at 10.0 pt

    for row in table_spec.rows:
        deliv_cell = row.cells[3]
        for line in deliv_cell.paras:
            line_text = "".join(r.text for r in line)
            assert "DEL-" in line_text
            assert line_text.strip() not in ("DEL-01", "DEL-02", "DEL-03")

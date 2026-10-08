"""HTL-10: Tests for resuming document generation from an approved, corrected review run."""

from datetime import date
from pathlib import Path
import docx
import pytest

from src.core.models import (
    ProjectStartupCharter,
    Stakeholder,
    StartupKitBaseline,
)
from src.review_storage.local import LocalReviewStorage
from src.review_ui import facts as review


def _create_sample_baseline(
    project_name: str = "Delta Cloud Modernization",
    delivery_manager: str = "Original Delivery Manager",
) -> StartupKitBaseline:
    return StartupKitBaseline(
        project_name=project_name,
        contract_type="Fixed Bid",
        governance_tier="Elevated",
        sow_awarded_date=date(2026, 11, 15),
        award_date_source="stated",
        charter=ProjectStartupCharter(
            project_name=project_name,
            client_name="Delta Corp",
            governance_tier="Elevated",
            contract_type="Fixed Bid",
            delivery_manager=delivery_manager,
        ),
        stakeholders=[
            Stakeholder(name="Diana Prince", role="Executive Sponsor", organization="Delta Corp"),
            Stakeholder(name=delivery_manager, role="Delivery Manager", organization="Toptal"),
        ],
    )


def test_round_trip_resume_generation_reflects_corrected_baseline_in_generated_kit(tmp_path):
    """HTL-10: Full round-trip test -- create a run, correct a fact via save_run_review, approve it,
    then trigger generation. Confirm the GENERATED Kit document's actual rendered content reflects
    the corrected value, not the original extracted one."""
    storage_dir = tmp_path / "review_queue"
    storage = LocalReviewStorage(base_dir=storage_dir)

    initial_baseline = _create_sample_baseline(
        project_name="Delta Cloud Modernization",
        delivery_manager="Arthur Pendelton (Original DM)",
    )
    run_id = storage.create_run(
        project_name="Delta Cloud Modernization",
        baseline=initial_baseline.model_dump(mode="json"),
        validation_report={"findings": []},
    )

    # 1. Review and correct a fact via save_run_review
    baseline = review.load_run_baseline(run_id, storage=storage)
    categories = review.build_fact_categories(baseline)
    submitted = review.extracted_values(categories)
    submitted["named_roles.delivery_manager"] = "Samantha Vance (Corrected DM)"
    notes = {"named_roles": "Updated DM to Samantha Vance per staffing plan"}

    audit, corrected = review.save_run_review(
        run_id=run_id,
        baseline=baseline,
        categories=categories,
        submitted=submitted,
        notes=notes,
        storage=storage,
    )

    # 2. Approve the run (HTL-08 / HTL-09)
    storage.update_status(run_id, "approved", if_state="pending_review")
    assert storage.get_run(run_id).state == "approved"

    # 3. Resume and generate documents (HTL-10)
    references, run_result = review.generate_approved_run(run_id, storage=storage)

    # 4. Verify generated references contain the Kit document and files exist on disk
    assert "Delta_Cloud_Modernization_Startup_Kit.docx" in references
    kit_file_path = Path(references["Delta_Cloud_Modernization_Startup_Kit.docx"])
    assert kit_file_path.exists()

    # 5. Open the real generated Word document and inspect actual paragraph and table text
    doc = docx.Document(str(kit_file_path))
    full_text = "\n".join(
        [p.text for p in doc.paragraphs]
        + [cell.text for tbl in doc.tables for row in tbl.rows for cell in row.cells]
    )

    # Confirm the corrected DM name appears in the generated document
    assert "Samantha Vance (Corrected DM)" in full_text
    # Confirm the original DM name does NOT appear in the generated document
    assert "Arthur Pendelton (Original DM)" not in full_text


def test_generation_blocked_when_state_is_not_approved(tmp_path):
    """HTL-10 / HTL-12: Confirm generation is blocked/unavailable when state is anything other than 'approved'."""
    storage_dir = tmp_path / "review_queue"
    storage = LocalReviewStorage(base_dir=storage_dir)

    # 1. State: pending_review
    baseline = _create_sample_baseline("Pending Project")
    run_pending = storage.create_run(
        project_name="Pending Project",
        baseline=baseline.model_dump(mode="json"),
        validation_report={},
    )
    with pytest.raises(ValueError, match=r"expected 'approved'"):
        review.generate_approved_run(run_pending, storage=storage)

    # 2. State: rejected
    baseline_rej = _create_sample_baseline("Rejected Project")
    run_rejected = storage.create_run(
        project_name="Rejected Project",
        baseline=baseline_rej.model_dump(mode="json"),
        validation_report={},
    )
    storage.update_status(run_rejected, "rejected", if_state="pending_review")
    with pytest.raises(ValueError, match=r"expected 'approved'"):
        review.generate_approved_run(run_rejected, storage=storage)

    # 3. State: already generated
    baseline_gen = _create_sample_baseline("Generated Project")
    run_gen = storage.create_run(
        project_name="Generated Project",
        baseline=baseline_gen.model_dump(mode="json"),
        validation_report={},
    )
    storage.update_status(run_gen, "approved", if_state="pending_review")
    # First generation moves state to 'generated'
    review.generate_approved_run(run_gen, storage=storage)
    assert storage.get_run(run_gen).state == "generated"

    # Second generation attempt must fail because state is 'generated'
    with pytest.raises(ValueError, match=r"expected 'approved'"):
        review.generate_approved_run(run_gen, storage=storage)


def test_run_transitions_to_generated_on_disk_verified_via_fresh_storage(tmp_path):
    """HTL-10: Confirm the run transitions to 'generated' on disk after successful generation,
    verified via a fresh LocalReviewStorage instance reading from disk."""
    storage_dir = tmp_path / "review_queue"
    storage = LocalReviewStorage(base_dir=storage_dir)

    baseline = _create_sample_baseline("Transition Project")
    run_id = storage.create_run(
        project_name="Transition Project",
        baseline=baseline.model_dump(mode="json"),
        validation_report={},
    )
    storage.update_status(run_id, "approved", if_state="pending_review")

    # Generate documents
    references, result = review.generate_approved_run(run_id, storage=storage)

    # Verify state via a FRESH LocalReviewStorage instance
    fresh_storage = LocalReviewStorage(base_dir=storage_dir)
    fresh_run = fresh_storage.get_run(run_id)
    assert fresh_run.state == "generated"

    # Confirm status.json content directly
    status_file = storage_dir / run_id / "status.json"
    assert '"state": "generated"' in status_file.read_text(encoding="utf-8")


def test_put_generated_files_called_and_returns_valid_file_references(tmp_path):
    """HTL-10 / HTL-13: Confirm put_generated_files is called and its returned references
    map generated artifact filenames to real files in the run's storage directory."""
    storage_dir = tmp_path / "review_queue"
    storage = LocalReviewStorage(base_dir=storage_dir)

    baseline = _create_sample_baseline("References Project")
    run_id = storage.create_run(
        project_name="References Project",
        baseline=baseline.model_dump(mode="json"),
        validation_report={},
    )
    storage.update_status(run_id, "approved", if_state="pending_review")

    references, result = review.generate_approved_run(run_id, storage=storage)

    # Expected artifact types: Kit, Checklist, Workbook, Deck, Deck Manifest
    expected_filenames = [
        "References_Project_Startup_Kit.docx",
        "References_Project_Startup_Readiness_Checklist.docx",
        "References_Project_Project_Delivery_Workbook.xlsx",
        "References_Project_Talent_Onboarding_Deck.pptx",
        "References_Project_Talent_Onboarding_Deck.trace.json",
    ]

    for fname in expected_filenames:
        assert fname in references, f"Missing {fname} in references: {list(references.keys())}"
        target_path = Path(references[fname])
        assert target_path.exists(), f"Referenced file {target_path} does not exist"
        assert target_path.parent == storage_dir / run_id

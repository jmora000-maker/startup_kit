"""DECK-21 (layout and alignment) and INV-32: broken-input unit tests on synthetic decks."""

import pytest
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Inches, Pt

from src.tools import deck_checks
from tests.deck_fixtures import add_card, add_table, add_text, blank_deck


def _inv32(prs):
    return [str(v) for v in deck_checks.check_inv32(prs) if v.inv_id == "INV-32"]


def _good_card_deck():
    prs, slides = blank_deck(2)
    s = slides[1]
    add_card(s, 0.83, 1.70, 3.71, 4.58)
    add_text(s, 1.03, 1.90, 3.31, 4.18, [("Client roles", 16.0), ("Client Approver: signs each milestone.", 11.0)], name="Card text")
    return prs, s


def test_inv32_passes_on_a_well_formed_card_and_table():
    prs, s = _good_card_deck()
    add_table(s, 4.81, 1.70, [0.80, 6.00, 0.89], [["ID", "Acceptance Criteria", "Gate"], ["DEL-01", "Shell load P50 < 1.5s.", "M1"]])
    assert _inv32(prs) == []


def test_inv32_fails_on_a_table_that_runs_off_the_slide():
    """L4: a long overflow row in a narrow column wraps past the slide bottom."""
    prs, s = _good_card_deck()
    rows = [["ID", "Acceptance Criteria", "Gate"]] + [["DEL-%02d" % i, "Criteria text", "M1"] for i in range(1, 14)]
    rows.append(["+6 more: DEL-14, DEL-15, DEL-16, DEL-17, DEL-18, DEL-19 (see Startup Kit · Deliverables and Acceptance Matrix)", "", ""])
    add_table(s, 4.81, 1.70, [0.80, 6.00, 0.89], rows)
    msgs = _inv32(prs)
    assert any("outside the content area" in m and "Table" in m for m in msgs), msgs


def test_inv32_fails_when_text_overflows_its_frame():
    prs, slides = blank_deck(2)
    add_card(slides[1], 0.83, 1.70, 3.71, 4.58)
    add_text(slides[1], 1.03, 1.90, 3.31, 4.18, [("Client roles", 16.0)] + [("A long bullet about delivery scope and the client approver for each milestone gate " * 3, 11.0)] * 8, name="Card text")
    assert any("fit estimate" in m for m in _inv32(prs))


@pytest.mark.parametrize("size", [9.5, 10.0])
def test_inv32_fails_on_body_text_below_the_minimum(size):
    prs, slides = blank_deck(2)
    add_card(slides[1], 0.83, 1.70, 3.71, 4.58)
    add_text(slides[1], 1.03, 1.90, 3.31, 4.18, [("Working rhythm", 16.0), ("Kickoff Call: One-time", size)], name="Card text")
    assert any("below the 10.5 pt minimum" in m for m in _inv32(prs))


def test_inv32_fails_on_table_text_below_ten_points():
    prs, slides = blank_deck(3)
    add_table(slides[1], 0.83, 1.70, [1.35, 1.96], [["Client Sponsor", "Acme"]], size=9.5)
    assert any("below the 10 pt minimum" in m for m in _inv32(prs))


def test_inv32_allows_nine_points_only_in_the_slide_3_deliverables_column():
    prs, slides = blank_deck(3)
    header = ["Workstream", "Milestone", "Dates", "Deliverables"]
    shape = add_table(slides[2], 0.83, 1.70, [2.0, 2.6, 1.9, 5.17], [header, ["P1", "M1", "2026-10-05", "DEL-01 Shell"]], size=10.0)
    shape.table.cell(1, 3).text_frame.paragraphs[0].runs[0].font.size = Pt(9)
    assert _inv32(prs) == []
    shape.table.cell(1, 1).text_frame.paragraphs[0].runs[0].font.size = Pt(9)
    assert any("R1C1" in m and "below the 10 pt minimum" in m for m in _inv32(prs))


def test_inv32_fails_on_shapes_outside_the_content_area():
    prs, slides = blank_deck(2)
    add_card(slides[1], 0.30, 1.70, 3.71, 4.58)
    add_card(slides[1], 4.81, 1.70, 3.71, 5.40)
    msgs = _inv32(prs)
    assert sum("outside the content area" in m for m in msgs) == 2


def test_inv32_fails_on_overlapping_content_shapes():
    """L1/L3: a heading text frame that covers the table inside its card."""
    prs, slides = blank_deck(2)
    s = slides[1]
    add_card(s, 0.83, 1.70, 3.71, 4.58)
    add_text(s, 1.03, 1.90, 3.31, 4.18, [("Key facts", 16.0)], name="Card text")
    add_table(s, 1.03, 2.40, [1.35, 1.96], [["Client Sponsor", "Acme"]])
    assert any("overlap" in m for m in _inv32(prs))


def test_inv32_card_text_frame_must_be_inset_and_heading_one_line():
    prs, slides = blank_deck(2)
    add_card(slides[1], 0.83, 1.70, 3.71, 4.58)
    add_text(slides[1], 0.95, 1.80, 3.47, 4.38, [("A very long card heading that cannot fit on one line", 16.0)], name="Card text")
    msgs = _inv32(prs)
    assert any("not inset" in m for m in msgs)
    assert any("wraps to more than one line" in m for m in msgs)


def test_inv32_card_text_frame_must_not_autofit():
    prs, slides = blank_deck(2)
    add_card(slides[1], 0.83, 1.70, 3.71, 4.58)
    box = add_text(slides[1], 1.03, 1.90, 3.31, 4.18, [("Heading", 16.0)], name="Card text")
    box.text_frame.auto_size = MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT
    assert any("autofit off" in m for m in _inv32(prs))


LONG_TITLE = "Pfizer Analytics and Cloud Modernization Platform"


def _cover_deck(title, title_pt=None):
    """A one-slide deck with the cover on the template's CUSTOM_1 layout, built as the writer builds it."""
    import pptx

    from src.config import DECK_TEMPLATE_PATH

    prs = pptx.Presentation(str(DECK_TEMPLATE_PATH))
    ids = prs.slides._sldIdLst
    rids = [el.rId for el in list(ids)]
    for el in list(ids):
        ids.remove(el)
    for rid in rids:
        prs.part.drop_rel(rid)
    layout = next(l for l in prs.slide_masters[0].slide_layouts if l.name == "CUSTOM_1")
    slide = prs.slides.add_slide(layout)
    slide.placeholders[0].text = title
    slide.placeholders[1].text = "Talent Team Onboarding · Acme · Start 2026-10-05"
    if title_pt is not None:
        slide.placeholders[0].text_frame.paragraphs[0].runs[0].font.size = Pt(title_pt)
    return prs, slide


def test_inv32_passes_on_a_cover_whose_title_fits_on_one_line():
    prs, _ = _cover_deck("Acme Genomics Platform")
    assert _inv32(prs) == []


def test_inv32_fails_when_the_cover_title_is_estimated_to_run_past_the_subtitle():
    """M2: a long project name wraps to two lines at the layout size and its estimated bottom passes the subtitle top."""
    prs, _ = _cover_deck(LONG_TITLE)
    assert any("cover title" in m and "below the subtitle's top" in m for m in _inv32(prs))


def test_inv32_cover_margin_catches_known_wrap():
    """Cover-only 1.15 width margin (Proxima Nova is wider than the shared 0.5 x font-size estimate): the mock project's
    38-character title is exactly one line by the shared formula but wraps in a real render (Appendix M2)."""
    prs, _ = _cover_deck("Pfizer Analytics & Cloud Modernization")
    assert any("cover title" in m and "below the subtitle's top" in m for m in _inv32(prs))


@pytest.mark.parametrize("title", ["ARC Genomics Platform", "Acme Genomics Platform", "Syngenta Crop Protection"])
def test_inv32_cover_margin_allows_short_title(title):
    prs, _ = _cover_deck(title)
    assert _inv32(prs) == []


def test_inv32_passes_when_a_long_cover_title_is_reduced_to_one_line():
    prs, _ = _cover_deck(LONG_TITLE, title_pt=28.0)
    assert _inv32(prs) == []


def test_inv32_fails_when_a_cover_placeholder_leaves_its_layout_bounds():
    prs, s = _cover_deck("Acme Genomics Platform")
    s.placeholders[1].top = s.placeholders[1].top + Inches(0.5)
    assert any("cover subtitle" in m and "outside the cover layout placeholder region" in m for m in _inv32(prs))


def test_inv32_flags_runs_without_an_explicit_font_size():
    prs, slides = blank_deck(2)
    b = slides[1].shapes.add_textbox(Inches(0.83), Inches(1.7), Inches(3), Inches(1))
    b.text_frame.text = "Text with no explicit size"
    assert any("no explicit font size" in m for m in _inv32(prs))


# ---------------------------------------------------------------------------------------------
# DECK-21 on the written ARC deck
# ---------------------------------------------------------------------------------------------
import pytest
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.ns import qn

from src.generators.onboarding_deck import layout as L
from tests.deck_bundle import FIXTURES, build_bundle, card_texts, slide_tables


def _inches(emu):
    return emu / 914400.0


def _cards(slide):
    return [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and s.name.startswith("Card:")]


@pytest.mark.parametrize("name", FIXTURES)
def test_inv32_passes_on_every_content_slide_of_every_fixture(name, tmp_path):
    bundle = build_bundle(tmp_path / name, name)
    assert [str(v) for v in deck_checks.check_inv32(bundle["prs"])] == [], name


def _cover_of(tmp_path, project_name):
    from tests.deck_bundle import load_baseline

    baseline = load_baseline("mock_sow")
    baseline.project_name = project_name
    prs = build_bundle(tmp_path, baseline=baseline)["prs"]
    slide = prs.slides[0]
    return prs, slide.placeholders[0], slide.placeholders[1]


def test_cover_title_shrinks_to_one_line_for_the_mock_project(tmp_path):
    """DECK-21 (5): the 38-character mock title is reduced below the layout's 43 pt, not below 28 pt, and stays on one line."""
    prs, title, sub = _cover_of(tmp_path, "Pfizer Analytics & Cloud Modernization")
    size = title.text_frame.paragraphs[0].runs[0].font.size.pt
    assert 28 <= size < 43
    assert deck_checks.cover_title_lines(title.text_frame.text, title.width / 914400.0, size) == 1
    assert [str(v) for v in deck_checks.check_inv32(prs)] == []


def test_cover_title_keeps_the_layout_size_when_it_fits(tmp_path):
    _, title, _ = _cover_of(tmp_path, "ARC Genomics Platform")
    assert title.text_frame.paragraphs[0].runs[0].font.size is None


def test_cover_title_that_cannot_fit_at_28_pt_wraps_and_pushes_the_subtitle_down(tmp_path):
    name = "Pfizer Global Analytics and Cloud Modernization Programme for Manufacturing and Supply Planning"
    prs, title, sub = _cover_of(tmp_path, name)
    assert title.text_frame.paragraphs[0].runs[0].font.size.pt == 28
    assert deck_checks.cover_title_lines(name, title.width / 914400.0, 28.0) == 2
    assert sub.top >= title.top + title.height - 1
    assert [str(v) for v in deck_checks.check_inv32(prs)] == []


def test_card_text_frames_follow_deck_21_1(arc_bundle):
    checked = 0
    for idx in range(1, 7):
        slide = arc_bundle["prs"].slides[idx]
        for card in _cards(slide):
            frame = next(s for s in slide.shapes if s.name == card.name.replace("Card:", "Card text:"))
            tf = frame.text_frame
            assert tf.word_wrap is True and tf.auto_size == MSO_AUTO_SIZE.NONE and tf.vertical_anchor == MSO_ANCHOR.TOP
            assert abs(_inches(frame.left - card.left) - 0.20) < 0.015
            assert abs(_inches(frame.top - card.top) - 0.20) < 0.015
            assert abs(_inches((card.left + card.width) - (frame.left + frame.width)) - 0.20) < 0.015
            heading = tf.paragraphs[0]
            assert L.wrap_lines(heading.text, _inches(frame.width), heading.runs[0].font.size.pt) == 1
            assert heading.space_after.pt == 8
            checked += 1
    assert checked == 9  # 3 + 1 + 3 + 2 cards on slides 2, 4, 6, and 7


def test_bullets_are_paragraph_bullets_not_typed_characters(arc_bundle):
    slide = arc_bundle["prs"].slides[5]
    bullets = 0
    for sh in slide.shapes:
        if sh.has_text_frame and sh.name.startswith("Card text:"):
            for p in sh.text_frame.paragraphs[1:]:
                assert not p.text.startswith("•"), "no typed bullet character"
                pPr = p._p.pPr
                if pPr is not None and pPr.find(qn("a:buChar")) is not None:
                    assert pPr.find(qn("a:buChar")).get("char") == "•"
                    assert pPr.get("marL") == str(round(0.17 * 914400)) and pPr.get("indent") == str(-round(0.17 * 914400))
                    bullets += 1
    assert bullets >= 10


def test_body_spacing_and_minimum_fonts_in_cards(arc_bundle):
    for idx in range(1, 7):
        for sh in arc_bundle["prs"].slides[idx].shapes:
            if sh.has_text_frame and sh.name.startswith("Card text:"):
                for p in sh.text_frame.paragraphs[1:]:
                    size = p.runs[0].font.size.pt
                    assert size >= 10.5
                    if size != 11 or p.runs[0].font.name != "Proxima Nova Semibold":
                        pass
                    if p.runs[0].font.name == "Calibri":
                        assert p.space_after.pt == 4


@pytest.mark.parametrize("idx, widths", [
    (2, [2.00, 2.60, 1.90, 5.17]),
    (3, [0.80, 6.00, 0.89]),
    (4, [1.50, 3.70, 0.85, 1.60, 3.00, 1.02]),
])
def test_table_column_widths_header_and_row_heights(arc_bundle, idx, widths):
    table = slide_tables(arc_bundle["prs"].slides[idx])[0].table
    assert [round(_inches(c.width), 2) for c in table.columns] == widths
    assert abs(_inches(table.rows[0].height) - 0.40) < 0.005
    assert all(_inches(r.height) >= 0.30 - 0.005 for r in list(table.rows)[1:])
    for cell in (table.cell(1, 0), table.cell(1, 1)):
        assert round(_inches(cell.margin_left), 2) == 0.06 and round(_inches(cell.margin_right), 2) == 0.06
        assert round(_inches(cell.margin_top), 2) == 0.04 and round(_inches(cell.margin_bottom), 2) == 0.04
        assert cell.vertical_anchor == MSO_ANCHOR.TOP


def test_key_facts_table_widths_and_key_value_style(arc_bundle):
    shape = slide_tables(arc_bundle["prs"].slides[1])[0]
    table = shape.table
    assert [round(_inches(c.width), 2) for c in table.columns] == [1.35, 1.96]
    assert table.first_row is False, "a key-value table has no header row"
    label, value = table.cell(0, 0), table.cell(0, 1)
    assert label.fill.fore_color.rgb == L_BAND and value.fill.fore_color.rgb == L_WHITE
    assert (label.text_frame.paragraphs[0].runs[0].font.name, label.text_frame.paragraphs[0].runs[0].font.size.pt, label.text_frame.paragraphs[0].runs[0].font.bold) == ("Proxima Nova", 10, True)
    assert (value.text_frame.paragraphs[0].runs[0].font.name, value.text_frame.paragraphs[0].runs[0].font.size.pt) == ("Calibri", 10)


L_BAND = __import__("pptx.dml.color", fromlist=["RGBColor"]).RGBColor(0xF8, 0xFA, 0xFC)
L_WHITE = __import__("pptx.dml.color", fromlist=["RGBColor"]).RGBColor(0xFF, 0xFF, 0xFF)


@pytest.mark.parametrize("idx", [1, 2, 3, 4])
def test_built_in_table_style_is_removed(arc_bundle, idx):
    table = slide_tables(arc_bundle["prs"].slides[idx])[0].table
    assert table._tbl.tblPr.find(qn("a:tableStyleId")) is None
    assert table.horz_banding is False


def test_overflow_row_is_one_merged_cell_across_all_columns(arc_bundle):
    table = slide_tables(arc_bundle["prs"].slides[3])[0].table
    last = list(table.rows)[-1]
    origin = last.cells[0]
    assert origin.is_merge_origin and origin.span_width == len(table.columns)
    assert origin.text.startswith("+") and "more:" in origin.text
    assert all(c.is_spanned for c in list(last.cells)[1:])


def test_slide_3_deliverables_column_font_is_ten_or_nine(arc_bundle):
    table = slide_tables(arc_bundle["prs"].slides[2])[0].table
    sizes = {r.font.size.pt for row in list(table.rows)[1:] for p in row.cells[3].text_frame.paragraphs for r in p.runs}
    assert sizes <= {9.0, 10.0}


def test_estimate_matches_the_k_capacities_on_arc(arc_bundle):
    """Slide 4: 1 sentence + 6 steps + 2 facts fit the left card; slide 6 shows 5 client roles then +N more."""
    card = card_texts(arc_bundle["prs"].slides[3])["Card text: How acceptance works"]
    assert len(card) == 1 + 1 + 6 + 2
    roles = card_texts(arc_bundle["prs"].slides[5])["Card text: Client roles"]
    assert len(roles) == 1 + 5 + 1


def test_fit_falls_back_to_plus_n_more_never_to_overflow():
    """DECK-21 (5): when text cannot fit at the minimum font, bullets move into a `+N more` line."""
    from src.generators.onboarding_deck.builder import Block, bullet_block, fit_card, para, see_more
    from src.generators.onboarding_deck.spec import BULLET, CardSpec, Run

    items = [para(BULLET, Run("A long bullet about delivery scope, the client approver, and each milestone gate " * 2)) for _ in range(12)]
    card = CardSpec("Client roles", 0.83, 1.70, 3.71, L.CARD_H)
    fit_card(card, [bullet_block([], items, 12, lambda n: see_more(n, "Startup Kit"))])
    assert card.body_pt >= 10.5
    assert card.paras[-1].text.startswith("+") and "more" in card.paras[-1].text
    inner_w, inner_h = L.card_inner(card.w, card.h)
    assert L.frame_height_pt([L.Para(card.heading, card.heading_pt, 0, 8)] + [L.Para(p.text, card.body_pt, 0, 4, 0.17) for p in card.paras], inner_w) <= inner_h * 72

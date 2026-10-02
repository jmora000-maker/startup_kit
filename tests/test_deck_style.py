"""Tests for the deck's explicit styles and typography (DECK-13, Appendix K.1)."""

import pytest
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches, Pt

from src.generators.onboarding_deck.styles import (
    COLOR_ACCENT_BLUE,
    COLOR_AMBER_BG,
    COLOR_AMBER_TEXT,
    COLOR_DEEP_NAVY,
    COLOR_GREEN_BG,
    COLOR_GREEN_TEXT,
    COLOR_RED_BG,
    COLOR_RED_TEXT,
    COLOR_WHITE,
    FONT_BODY,
    FONT_SIZE_TABLE_HEADER,
    FONT_SIZE_TITLE,
    FONT_TITLE,
    SLIDE_HEIGHT,
    SLIDE_WIDTH,
)
from tests.deck_bundle import slide_tables


def _run(shape, para=0, run=0):
    return shape.text_frame.paragraphs[para].runs[run]


def test_style_tokens():
    assert SLIDE_WIDTH == Inches(13.333) and SLIDE_HEIGHT == Inches(7.5)
    assert COLOR_DEEP_NAVY == RGBColor(0x0F, 0x17, 0x2A) and COLOR_ACCENT_BLUE == RGBColor(0x20, 0x4E, 0xCF)
    assert COLOR_WHITE == RGBColor(0xFF, 0xFF, 0xFF)
    assert (COLOR_RED_TEXT, COLOR_RED_BG) == (RGBColor(0x99, 0x1B, 0x1B), RGBColor(0xFE, 0xE2, 0xE2))
    assert (COLOR_AMBER_TEXT, COLOR_AMBER_BG) == (RGBColor(0x92, 0x40, 0x0E), RGBColor(0xFE, 0xF3, 0xC7))
    assert (COLOR_GREEN_TEXT, COLOR_GREEN_BG) == (RGBColor(0x16, 0x65, 0x34), RGBColor(0xDC, 0xFC, 0xE7))
    assert FONT_TITLE == "Proxima Nova" and FONT_BODY == "Calibri"
    assert FONT_SIZE_TITLE == Pt(28) and FONT_SIZE_TABLE_HEADER == Pt(12)


@pytest.mark.parametrize("idx", range(1, 7))
def test_title_and_kicker_are_set_explicitly_on_every_content_slide(arc_bundle, idx):
    """L2: the layout default renders the title blue and regular and the kicker large."""
    slide = arc_bundle["prs"].slides[idx]
    title, kicker = _run(slide.placeholders[0]), _run(slide.placeholders[1])
    assert (title.font.name, title.font.size, title.font.bold) == ("Proxima Nova", Pt(28), True)
    assert title.font.color.rgb == RGBColor(0x0F, 0x17, 0x2A)
    assert (kicker.font.name, kicker.font.size, kicker.font.bold) == ("Proxima Nova", Pt(11), True)
    assert kicker.font.color.rgb == RGBColor(0x64, 0x74, 0x8B)
    assert slide.placeholders[1].text == "ARC GENOMICS PLATFORM · TALENT TEAM ONBOARDING"


def test_cover_keeps_the_layouts_placeholder_styling(arc_bundle):
    slide = arc_bundle["prs"].slides[0]
    for ph in (slide.placeholders[0], slide.placeholders[1]):
        for r in ph.text_frame.paragraphs[0].runs:
            assert r.font.name is None and r.font.size is None and r.font.bold is None


def test_cards_are_white_rounded_rectangles_with_template_corner_and_outline(arc_bundle):
    for idx in range(1, 7):
        for sh in arc_bundle["prs"].slides[idx].shapes:
            if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and sh.name.startswith("Card:"):
                assert round(sh.adjustments[0] * 100000) == 6153
                assert sh.fill.fore_color.rgb == RGBColor(0xFF, 0xFF, 0xFF)
                assert sh.line.color.rgb == RGBColor(0xE2, 0xE8, 0xF0)
                assert sh.line.width == Pt(1.1)


def test_card_headings_and_subheadings(arc_bundle):
    slide = arc_bundle["prs"].slides[1]
    checked = 0
    for sh in slide.shapes:
        if sh.has_text_frame and sh.name.startswith("Card text:"):
            head = _run(sh)
            assert (head.font.name, head.font.bold, head.font.color.rgb) == ("Proxima Nova", True, RGBColor(0x0F, 0x17, 0x2A))
            assert head.font.size in (Pt(16), Pt(14))
            for p in sh.text_frame.paragraphs[1:]:
                if p.text in ("Purpose", "Delivery model", "Escalation path", "Phases", "Out of scope"):
                    r = p.runs[0]
                    assert (r.font.name, r.font.size, r.font.color.rgb) == ("Proxima Nova Semibold", Pt(11), RGBColor(0x20, 0x4E, 0xCF))
                    checked += 1
    assert checked == 5


def test_body_text_is_calibri_11_pt_475569(arc_bundle):
    sh = next(s for s in arc_bundle["prs"].slides[1].shapes if s.name == "Card text: Purpose & delivery")
    body = sh.text_frame.paragraphs[2].runs[0]
    assert body.font.name == "Calibri" and body.font.size in (Pt(11), Pt(10.5)) and body.font.color.rgb == RGBColor(0x47, 0x55, 0x69)


@pytest.mark.parametrize("idx", [2, 3, 4])
def test_tables_have_blue_header_alternating_rows_and_borders(arc_bundle, idx):
    shape = slide_tables(arc_bundle["prs"].slides[idx])[0]
    table = shape.table
    for cell in table.rows[0].cells:
        r = cell.text_frame.paragraphs[0].runs[0]
        assert cell.fill.fore_color.rgb == RGBColor(0x20, 0x4E, 0xCF)
        assert (r.font.name, r.font.size, r.font.bold, r.font.color.rgb) == ("Proxima Nova", Pt(12), True, RGBColor(0xFF, 0xFF, 0xFF))
    body_fills = [table.cell(i, 0).fill.fore_color.rgb for i in range(1, min(len(table.rows), 4))]
    if idx != 4:  # slide 5's rating column has its own fill, column 0 does not
        assert body_fills[0] == RGBColor(0xF8, 0xFA, 0xFC) and body_fills[1] == RGBColor(0xFF, 0xFF, 0xFF)
    xml = table.cell(1, 1)._tc.xml
    assert xml.count("E2E8F0") >= 4


def test_rating_cells_use_red_amber_green_fills(arc_bundle):
    table = slide_tables(arc_bundle["prs"].slides[4])[0].table
    cell = table.cell(1, 2)
    assert cell.text == "High" and cell.fill.fore_color.rgb == RGBColor(0xFE, 0xE2, 0xE2)
    assert cell.text_frame.paragraphs[0].runs[0].font.color.rgb == RGBColor(0x99, 0x1B, 0x1B)


def test_placeholders_are_italic_accent_blue(arc_bundle):
    table = slide_tables(arc_bundle["prs"].slides[1])[0].table
    r = table.cell(4, 1).text_frame.paragraphs[0].runs[0]
    assert r.text == "To be confirmed" and r.font.italic and r.font.color.rgb == RGBColor(0x20, 0x4E, 0xCF)


def test_no_icon_fonts_and_no_decorative_shapes(arc_bundle):
    for slide in arc_bundle["prs"].slides:
        xml = slide._element.xml
        assert 'typeface="Material' not in xml and "Material Icons" not in xml and "Material Symbols" not in xml
        for sh in slide.shapes:
            if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
                assert sh.name.startswith("Card:"), f"unexpected decorative shape {sh.name}"

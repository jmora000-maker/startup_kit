"""Tests for action tag column placement and ambiguity row hygiene (spec v5 B14)."""

from datetime import date
from src.generators.docx_generator import DocxGenerator
from src.core.models import (
    StartupKitBaseline,
    ProjectStartupCharter,
    Deliverable,
    Milestone,
    ActionRequiredItem,
    ContractAmbiguityItem,
    SOWInterpretationSummary,
    SourceReference,
)
import docx


def test_action_column_placement_in_docx(tmp_path):
    """Test that action tags appear only in their target column and Contract Ambiguities Logged has no action tag."""
    ref = SourceReference(document_name="SOW.pdf", clause_or_slide="Sec 1", confidence_score=0.9)

    baseline = StartupKitBaseline(
        project_name="Test Project",
        charter=ProjectStartupCharter(project_name="Test Project"),
        sow_interpretation=SOWInterpretationSummary(
            contracted_deliverables=["Core Service"],
            out_of_scope_items=[],
            customer_obligations=["Cloud Access"],
            assumptions=[],
            constraints=[],
            approval_expectations="Client review in 5 days"
        ),
        deliverables=[
            Deliverable(
                id="DEL-01",
                name="Core Service",
                acceptance_criteria="E2E test suite",
                evidence_required="Test report",
                client_approver="Client VP",
                owner="[UNASSIGNED - TO BE CONFIRMED]",
                linked_action_id="ACT-01",
                source_reference=ref
            )
        ],
        milestones=[
            Milestone(id="M1", description="Phase 1", external_date=None, internal_buffer_date=None, source_reference=ref)
        ],
        contract_ambiguities=[
            ContractAmbiguityItem(anomaly_id="AMB-01", conflicting_clauses="Conflicting clauses in section 1", source_reference=ref)
        ],
        action_required_items=[
            ActionRequiredItem(
                action_id="ACT-01",
                item_type="Open Exception",
                checklist_id="G01-03",
                related_artifact="Deliverables and Acceptance Matrix",
                target_entity_id="DEL-01",
                target_table_title="Deliverables and Acceptance Matrix",
                target_column_header="Owner",
                finding_description="Deliverable owner is unassigned",
                required_action="Assign named delivery owner for DEL-01",
                owner="PMO Lead"
            )
        ]
    )

    gen = DocxGenerator()
    out_file = tmp_path / "Test_Startup_Kit.docx"
    gen.write_kit_docx(baseline, out_file)

    doc = docx.Document(str(out_file))

    # Check SOW Interpretation Summary table for 'Contract Ambiguities Logged'
    sow_table = next(t for t in doc.tables if "Contracted Deliverables" in t.rows[0].cells[0].text or "Contracted Deliverables" in t.rows[1].cells[0].text)
    for row in sow_table.rows:
        if "Contract Ambiguities Logged" in row.cells[0].text:
            assert "[ACT-" not in row.cells[1].text, f"Action tag found in Contract Ambiguities Logged cell: {row.cells[1].text}"

    # Check Deliverables and Acceptance Matrix table
    deliv_table = next(t for t in doc.tables if "Deliverable Name" in t.rows[0].cells[1].text)
    deliv_row = deliv_table.rows[1]
    # Header: ID, Deliverable Name, Acceptance Criteria, Evidence Required, Client Approver, Owner, SOW Stories, Review Window
    # cells[2]: Acceptance Criteria -> should NOT have ACT-01
    assert "[ACT-01" not in deliv_row.cells[2].text, f"ACT-01 incorrectly attached to Acceptance Criteria: {deliv_row.cells[2].text}"
    # cells[5]: Owner -> SHOULD have ACT-01
    assert "[ACT-01" in deliv_row.cells[5].text, f"ACT-01 not attached to Owner: {deliv_row.cells[5].text}"

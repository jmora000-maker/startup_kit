"""Tests for CLI leadership roles, directory prompts, and charter defaults."""

import sys
import docx
import pytest
from pathlib import Path
from src.core.models import ProjectStartupCharter, TalentOnboardingRecord
from src.config import config
from main import prompt_role_names, prompt_directories, prompt_api_keys, main, parse_args


def test_project_startup_charter_role_defaults():
    """Verify that ProjectStartupCharter defaults pmo_lead, delivery_manager, and talent_pm to [UNASSIGNED - TO BE CONFIRMED]."""
    charter = ProjectStartupCharter()
    assert charter.pmo_lead == "[UNASSIGNED - TO BE CONFIRMED]"
    assert charter.delivery_manager == "[UNASSIGNED - TO BE CONFIRMED]"
    assert charter.talent_pm == "[UNASSIGNED - TO BE CONFIRMED]"


def test_talent_onboarding_record_role_defaults():
    """Verify that TalentOnboardingRecord defaults pmo_lead, delivery_manager, and talent_pm to [UNASSIGNED - TO BE CONFIRMED]."""
    record = TalentOnboardingRecord()
    assert record.pmo_lead == "[UNASSIGNED - TO BE CONFIRMED]"
    assert record.delivery_manager == "[UNASSIGNED - TO BE CONFIRMED]"
    assert record.talent_pm == "[UNASSIGNED - TO BE CONFIRMED]"


def test_prompt_directories_from_flags():
    """Verify that provided CLI directory arguments are used without prompting."""
    in_dir, out_dir = prompt_directories(
        inputs_dir=Path("custom/inputs"),
        output_dir=Path("custom/output"),
        interactive=False
    )
    assert in_dir == Path("custom/inputs")
    assert out_dir == Path("custom/output")


def test_prompt_directories_interactive_provided(monkeypatch):
    """Verify that interactive user inputs for directories are resolved as Paths."""
    inputs = iter(["my_inputs", "my_output"])
    monkeypatch.setattr("builtins.input", lambda prompt="": next(inputs))

    in_dir, out_dir = prompt_directories(interactive=True)
    assert in_dir == Path("my_inputs")
    assert out_dir == Path("my_output")


def test_prompt_directories_interactive_empty_defaults(monkeypatch):
    """Verify that empty inputs in interactive mode fall back to default directories."""
    inputs = iter(["", "   "])
    monkeypatch.setattr("builtins.input", lambda prompt="": next(inputs))

    in_dir, out_dir = prompt_directories(interactive=True)
    assert in_dir == config.inputs_dir
    assert out_dir == config.output_dir


def test_prompt_directories_non_interactive():
    """Verify that non-interactive mode returns default directories without error."""
    in_dir, out_dir = prompt_directories(interactive=False)
    assert in_dir == config.inputs_dir
    assert out_dir == config.output_dir


def test_prompt_directories_eof_handling(monkeypatch):
    """Verify that EOFError gracefully falls back to default directories."""
    def raise_eof(prompt=""):
        raise EOFError()

    monkeypatch.setattr("builtins.input", raise_eof)
    in_dir, out_dir = prompt_directories(interactive=True)
    assert in_dir == config.inputs_dir
    assert out_dir == config.output_dir


def test_prompt_role_names_from_flags():
    """Verify that provided CLI flag values are used without prompting."""
    pmo, dl, tpm = prompt_role_names(
        pmo_lead="Sarah Connor",
        delivery_lead="Jane Doe",
        talent_pm="John Smith",
        interactive=False
    )
    assert pmo == "Sarah Connor"
    assert dl == "Jane Doe"
    assert tpm == "John Smith"


def test_prompt_role_names_interactive_provided(monkeypatch):
    """Verify that interactive user inputs are captured when flags are not passed."""
    inputs = iter(["Alice Wonderland", "Bob Builder", "Charlie Chaplin"])
    monkeypatch.setattr("builtins.input", lambda prompt="": next(inputs))

    pmo, dl, tpm = prompt_role_names(interactive=True)
    assert pmo == "Alice Wonderland"
    assert dl == "Bob Builder"
    assert tpm == "Charlie Chaplin"


def test_prompt_role_names_inverted_casing_normalized(monkeypatch):
    """Verify that inverted caps-lock / inconsistent casing inputs are automatically normalized to title case."""
    inputs = iter(["jAMES", "tINA", "jUDY"])
    monkeypatch.setattr("builtins.input", lambda prompt="": next(inputs))

    pmo, dl, tpm = prompt_role_names(interactive=True)
    assert pmo == "James"
    assert dl == "Tina"
    assert tpm == "Judy"


def test_normalize_person_name_helper():
    """Verify normalize_person_name behavior across various formats and defaults."""
    from src.config import normalize_person_name

    # Inverted caps lock
    assert normalize_person_name("jAMES") == "James"
    assert normalize_person_name("tINA") == "Tina"
    assert normalize_person_name("jUDY") == "Judy"
    assert normalize_person_name("jAMES mORA") == "James Mora"

    # Standard casing preserved
    assert normalize_person_name("James Mora") == "James Mora"
    assert normalize_person_name("Carolina Gurvich") == "Carolina Gurvich"
    assert normalize_person_name("Antonia McCullough") == "Antonia McCullough"

    # All lowercase / uppercase
    assert normalize_person_name("james") == "James"
    assert normalize_person_name("JAMES") == "James"

    # Default unassigned handling
    assert normalize_person_name("") == "[UNASSIGNED - TO BE CONFIRMED]"
    assert normalize_person_name(None) == "[UNASSIGNED - TO BE CONFIRMED]"
    assert normalize_person_name("[UNASSIGNED - TO BE CONFIRMED]") == "[UNASSIGNED - TO BE CONFIRMED]"
    assert normalize_person_name("unassigned") == "[UNASSIGNED - TO BE CONFIRMED]"


def test_prompt_role_names_interactive_empty_defaults(monkeypatch):
    """Verify that empty inputs in interactive mode default to [UNASSIGNED - TO BE CONFIRMED]."""
    inputs = iter(["", "   ", ""])
    monkeypatch.setattr("builtins.input", lambda prompt="": next(inputs))

    pmo, dl, tpm = prompt_role_names(interactive=True)
    assert pmo == "[UNASSIGNED - TO BE CONFIRMED]"
    assert dl == "[UNASSIGNED - TO BE CONFIRMED]"
    assert tpm == "[UNASSIGNED - TO BE CONFIRMED]"


def test_prompt_role_names_non_interactive():
    """Verify that non-interactive mode defaults to [UNASSIGNED - TO BE CONFIRMED] without error."""
    pmo, dl, tpm = prompt_role_names(interactive=False)
    assert pmo == "[UNASSIGNED - TO BE CONFIRMED]"
    assert dl == "[UNASSIGNED - TO BE CONFIRMED]"
    assert tpm == "[UNASSIGNED - TO BE CONFIRMED]"


def test_prompt_role_names_eof_handling(monkeypatch):
    """Verify that EOFError gracefully defaults to [UNASSIGNED - TO BE CONFIRMED]."""
    def raise_eof(prompt=""):
        raise EOFError()

    monkeypatch.setattr("builtins.input", raise_eof)
    pmo, dl, tpm = prompt_role_names(interactive=True)
    assert pmo == "[UNASSIGNED - TO BE CONFIRMED]"
    assert dl == "[UNASSIGNED - TO BE CONFIRMED]"
    assert tpm == "[UNASSIGNED - TO BE CONFIRMED]"


def test_prompt_api_keys_from_flags():
    """Verify that CLI flags override environment/prompt."""
    ant, oai = prompt_api_keys(
        api_key="sk-ant-flag",
        openai_api_key="sk-openai-flag",
        interactive=False,
        default_anthropic_key="sk-ant-env",
        default_openai_key="sk-openai-env"
    )
    assert ant == "sk-ant-flag"
    assert oai == "sk-openai-flag"


def test_prompt_api_keys_never_prompts_input(monkeypatch):
    """Verify that API key resolution never prompts for interactive input."""
    def should_not_be_called(prompt=""):
        raise RuntimeError("input() must never be called for API key resolution")

    monkeypatch.setattr("builtins.input", should_not_be_called)
    ant, oai = prompt_api_keys(
        interactive=True,
        default_anthropic_key="sk-ant-env",
        default_openai_key="sk-openai-env"
    )
    assert ant == "sk-ant-env"
    assert oai == "sk-openai-env"


def test_prompt_api_keys_defaults_from_config(monkeypatch):
    """Verify resolution falls back to config values when arguments are None."""
    def should_not_be_called(prompt=""):
        raise RuntimeError("input() must never be called for API key resolution")

    monkeypatch.setattr("builtins.input", should_not_be_called)
    ant, oai = prompt_api_keys(
        api_key=None,
        openai_api_key=None,
        default_anthropic_key="sk-ant-default",
        default_openai_key="sk-openai-default"
    )
    assert ant == "sk-ant-default"
    assert oai == "sk-openai-default"


def test_cli_provider_flags_parsing(monkeypatch):
    """Verify CLI argument parsing for provider flags (--provider, --openai, --anthropic)."""
    # 1. Default provider is anthropic
    monkeypatch.setattr(sys, "argv", ["main.py"])
    args = parse_args()
    assert args.provider == "anthropic"
    assert not args.openai
    assert not args.anthropic

    # 2. Explicit --provider openai
    monkeypatch.setattr(sys, "argv", ["main.py", "--provider", "openai"])
    args = parse_args()
    assert args.provider == "openai"

    # 3. Explicit --openai flag
    monkeypatch.setattr(sys, "argv", ["main.py", "--openai", "--model", "gpt-4o-mini"])
    args = parse_args()
    assert args.openai
    assert args.model == "gpt-4o-mini"

    # 4. Explicit --anthropic flag
    monkeypatch.setattr(sys, "argv", ["main.py", "--anthropic", "--model", "claude-sonnet-5-5"])
    args = parse_args()
    assert args.anthropic
    assert args.model == "claude-sonnet-5-5"


def test_cli_main_with_explicit_roles(populated_inputs_dir, tmp_path, monkeypatch):
    """Verify that CLI execution with explicit role flags assigns them to output baseline."""
    output_dir = tmp_path / "cli_explicit_roles_output"

    test_args = [
        "main.py",
        "--mock",
        "--inputs-dir", str(populated_inputs_dir),
        "--output-dir", str(output_dir),
        "--pmo-lead", "Alex PMO",
        "--delivery-lead", "Dana Delivery",
        "--talent-pm", "Taylor Talent",
        "--non-interactive"
    ]
    monkeypatch.setattr(sys, "argv", test_args)

    exit_code = main()
    assert exit_code == 0

    created_docx = list(output_dir.glob("*.docx"))
    assert len(created_docx) == 1

    doc = docx.Document(str(created_docx[0]))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    # Check tables for role metadata
    table_texts = "\n".join(cell.text for t in doc.tables for row in t.rows for cell in row.cells)
    combined = full_text + "\n" + table_texts

    assert "Alex PMO" in combined
    assert "Dana Delivery" in combined
    assert "Taylor Talent" in combined


def test_cli_main_with_delivery_manager_alias(populated_inputs_dir, tmp_path, monkeypatch):
    """Verify that --delivery-manager flag works as an alias for --delivery-lead."""
    output_dir = tmp_path / "cli_alias_output"

    test_args = [
        "main.py",
        "--mock",
        "--inputs-dir", str(populated_inputs_dir),
        "--output-dir", str(output_dir),
        "--delivery-manager", "Morgan Manager",
        "--non-interactive"
    ]
    monkeypatch.setattr(sys, "argv", test_args)

    exit_code = main()
    assert exit_code == 0

    created_docx = list(output_dir.glob("*.docx"))
    assert len(created_docx) == 1

    doc = docx.Document(str(created_docx[0]))
    table_texts = "\n".join(cell.text for t in doc.tables for row in t.rows for cell in row.cells)
    assert "Morgan Manager" in table_texts


def test_cli_main_default_unassigned_roles(populated_inputs_dir, tmp_path, monkeypatch):
    """Verify that CLI execution without role flags defaults to [UNASSIGNED - TO BE CONFIRMED]."""
    output_dir = tmp_path / "cli_unassigned_output"

    test_args = [
        "main.py",
        "--mock",
        "--inputs-dir", str(populated_inputs_dir),
        "--output-dir", str(output_dir),
        "--non-interactive"
    ]
    monkeypatch.setattr(sys, "argv", test_args)

    exit_code = main()
    assert exit_code == 0

    created_docx = list(output_dir.glob("*.docx"))
    assert len(created_docx) == 1

    doc = docx.Document(str(created_docx[0]))
    table_texts = "\n".join(cell.text for t in doc.tables for row in t.rows for cell in row.cells)
    assert "[UNASSIGNED - TO BE CONFIRMED]" in table_texts


def test_cli_main_interactive_directories_and_roles(populated_inputs_dir, tmp_path, monkeypatch):
    """Verify that interactive prompting for both directories and roles properly executes end-to-end."""
    output_dir = tmp_path / "cli_interactive_dirs_output"

    # User inputs for: mode (1: Initial Generation), inputs_dir, output_dir, pmo_lead, delivery_lead, talent_pm
    inputs = iter([
        "1",
        str(populated_inputs_dir),
        str(output_dir),
        "Interactive PMO",
        "Interactive Delivery",
        "Interactive Talent"
    ])
    monkeypatch.setattr("builtins.input", lambda prompt="": next(inputs))

    test_args = [
        "main.py",
        "--mock"
    ]
    monkeypatch.setattr(sys, "argv", test_args)

    exit_code = main()
    assert exit_code == 0

    created_docx = list(output_dir.glob("*.docx"))
    assert len(created_docx) == 1

    doc = docx.Document(str(created_docx[0]))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    table_texts = "\n".join(cell.text for t in doc.tables for row in t.rows for cell in row.cells)
    combined = full_text + "\n" + table_texts

    assert "Interactive PMO" in combined
    assert "Interactive Delivery" in combined
    assert "Interactive Talent" in combined


def test_cli_main_interactive_empty_directories_fallback(monkeypatch):
    """Verify that interactive empty directory inputs fall back to config default directories."""
    # User inputs: "1" for mode, "" for inputs_dir, "" for output_dir, then unassigned roles
    inputs = iter(["1", "", "", "", "", ""])
    monkeypatch.setattr("builtins.input", lambda prompt="": next(inputs))

    # Mock controller.run to verify resolved directory arguments without actually running full extraction on root inputs
    captured_kwargs = {}

    def mock_run(self, inputs_dir=None, output_dir=None, **kwargs):
        captured_kwargs["inputs_dir"] = inputs_dir
        captured_kwargs["output_dir"] = output_dir
        return Path("output/mock_result.docx")

    from src.orchestrator import StartupKitController
    monkeypatch.setattr(StartupKitController, "run", mock_run)

    test_args = ["main.py", "--mock"]
    monkeypatch.setattr(sys, "argv", test_args)

    exit_code = main()
    assert exit_code == 0
    assert captured_kwargs["inputs_dir"] == config.inputs_dir
    assert captured_kwargs["output_dir"] == config.output_dir


def test_cli_main_openai_provider_routing(monkeypatch):
    """Verify that --openai or --provider openai initializes LangChainLLMClient for OpenAI."""
    captured = {}

    class DummyController:
        def __init__(self, ingestion_service=None, llm_client=None, aggregator=None, doc_writer=None):
            captured["client"] = llm_client
        def run(self, **kwargs):
            return Path("output/test.docx")

    monkeypatch.setattr("main.StartupKitController", DummyController)
    monkeypatch.setattr(sys, "argv", [
        "main.py",
        "--openai",
        "--openai-api-key", "sk-proj-test-openai",
        "--model", "gpt-4o-mini",
        "--non-interactive"
    ])

    exit_code = main()
    assert exit_code == 0
    client = captured.get("client")
    assert client is not None
    assert client.openai_api_key == "sk-proj-test-openai"
    assert client.openai_model_name == "gpt-4o-mini"
    assert client._chat_model is None


def test_cli_main_anthropic_default_provider_routing(monkeypatch):
    """Verify that default execution initializes LangChainLLMClient for Anthropic Claude."""
    captured = {}

    class DummyController:
        def __init__(self, ingestion_service=None, llm_client=None, aggregator=None, doc_writer=None):
            captured["client"] = llm_client
        def run(self, **kwargs):
            return Path("output/test.docx")

    monkeypatch.setattr("main.StartupKitController", DummyController)
    monkeypatch.setattr(sys, "argv", [
        "main.py",
        "--api-key", "sk-ant-test-anthropic",
        "--openai-api-key", "sk-proj-fallback-openai",
        "--non-interactive"
    ])

    exit_code = main()
    assert exit_code == 0
    client = captured.get("client")
    assert client is not None
    assert client.api_key == "sk-ant-test-anthropic"
    assert client.openai_api_key == "sk-proj-fallback-openai"
    assert client._chat_model is not None

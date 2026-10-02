"""RAID-09 / INV-30: the Workbook RAID rows carry the Kit RAID Log fields."""

import pytest

from src.tools import deck_checks
from tests.deck_fixtures import KIT_RAID_ROW, make_kit, make_workbook


def _inv30(tmp_path, kit_row=None, wb_overrides=None):
    kit = make_kit(tmp_path / "kit.docx", kit_row)
    wb = make_workbook(tmp_path / "wb.xlsx", wb_overrides)
    return [str(v) for v in deck_checks.check_inv30(kit, wb) if v.inv_id == "INV-30"]


def test_inv30_passes_when_workbook_matches_kit(tmp_path):
    assert _inv30(tmp_path) == []


@pytest.mark.parametrize("field, value", [
    ("Mitigation / Response", ""),
    ("Mitigation / Response", "Active monitoring by PMO and Delivery Manager"),
    ("Probability", "Low"),
    ("Impact", "Low"),
    ("Owner", "Client"),
    ("Status", "Closed"),
])
def test_inv30_fails_when_a_field_differs(tmp_path, field, value):
    msgs = _inv30(tmp_path, wb_overrides={field: value})
    assert any(field in m and "RSK-01" in m for m in msgs), msgs


def test_inv30_compares_after_raid06_normalization(tmp_path):
    row = list(KIT_RAID_ROW)
    row[4], row[6], row[8] = "Moderate", "[ACT-03: Assign risk owner (+2.3% Recovery)]", "resolved"
    assert _inv30(tmp_path, kit_row=row, wb_overrides={"Probability": "Medium", "Owner": "[UNASSIGNED - TO BE CONFIRMED]", "Status": "Closed"}) == []


def test_inv30_ignores_rows_not_sourced_from_kit_risks_and_issues(tmp_path):
    assert _inv30(tmp_path, wb_overrides={"Source ID": "DEP-01", "Mitigation / Response": ""}) == []


# --- Fixtures: the real builders (RAID-09) -----------------------------------------------------
import json
from pathlib import Path

import docx
import openpyxl

from src.core.models import StartupKitBaseline
from src.generators.docx_generator import DocxGenerator
from src.generators.pmo_workbook import export_pmo_workbook
from src.generators.pmo_workbook.builder import (
    DEFAULT_RAID_TRIGGER,
    _normalize_rating_v2,
    _normalize_status_v2,
    build_workbook_model,
    clean_text_v2,
    normalize_owner_v2,
)
from src.llm.validation import validate_and_repair_baseline

FIXTURES = ["arc_genomics", "arc_overextracted", "mock_sow", "no_story_ids", "numbered_deliverables"]


def _baseline(name):
    data = json.loads((Path("tests/fixtures/sow") / name / "baseline.json").read_text(encoding="utf-8"))
    baseline = StartupKitBaseline.model_validate(data)
    validate_and_repair_baseline(baseline)
    return baseline


@pytest.mark.parametrize("name", FIXTURES)
def test_workbook_model_carries_kit_raid_fields(name):
    """Every Workbook RAID row sourced from a Kit risk or issue has the Kit's six fields (RAID-09)."""
    baseline = _baseline(name)
    model = build_workbook_model(baseline)
    rows = {r.source_id: r for r in model.raid_rows}
    for item in baseline.raid_items:
        if not item.id or item.id not in rows:
            continue  # merged duplicates keep the first row's Source ID
        row = rows[item.id]
        assert row.probability == _normalize_rating_v2(item.probability)[0]
        assert row.impact == _normalize_rating_v2(item.impact)[0]
        assert row.owner == normalize_owner_v2(item.owner)[0]
        assert row.mitigation_or_response == clean_text_v2(item.mitigation_or_response)
        trigger = "" if item.trigger_or_early_warning == DEFAULT_RAID_TRIGGER else clean_text_v2(item.trigger_or_early_warning)
        assert row.trigger_or_early_warning == trigger
        assert row.status == _normalize_status_v2(item.status)[0]


@pytest.mark.parametrize("name", FIXTURES)
def test_written_workbook_passes_inv30_on_all_fixtures(name, tmp_path):
    baseline = _baseline(name)
    kit_path = DocxGenerator(generated_date="2026-10-01").write_kit_docx(baseline, tmp_path)
    wb_path = export_pmo_workbook(baseline, tmp_path).file_path
    violations = deck_checks.check_inv30(docx.Document(str(kit_path)), openpyxl.load_workbook(str(wb_path), data_only=False))
    assert [str(v) for v in violations] == []


def test_arc_workbook_raid_01_to_07_have_a_response(tmp_path):
    """ARC: Mitigation / Response was empty on RAID-01 to RAID-07 although the Kit lists a mitigation for each."""
    baseline = _baseline("arc_genomics")
    wb_path = export_pmo_workbook(baseline, tmp_path).file_path
    ws = openpyxl.load_workbook(str(wb_path))["RAID Log"]
    hdr = [c.value for c in ws[5]]
    resp = hdr.index("Mitigation / Response")
    first_rows = [r for r in ws.iter_rows(min_row=6, max_row=12, values_only=True)]
    assert [r[0] for r in first_rows] == [f"RAID-0{i}" for i in range(1, 8)]
    assert all(r[resp] for r in first_rows)
    assert first_rows[0][resp].startswith("Use query observability (HS-4942)")
